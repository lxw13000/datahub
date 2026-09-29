"""验证 ES 聚合口径、每盒 Top 30、全盲盒 Top 100 和 Excel 输出。"""

import argparse
import importlib.util
import io
import json
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook


SPEC = importlib.util.spec_from_file_location(
    'blind_box_report', Path(__file__).with_name('export_report.py'))
report = importlib.util.module_from_spec(SPEC)
import sys
sys.modules[SPEC.name] = report
SPEC.loader.exec_module(report)


def term_bucket(key, count, tokens, amount=None, **nested):
    bucket = {'key': key, 'doc_count': count, 'sum_tokens': {'value': float(tokens)}}
    if amount is not None:
        bucket['sum_amount'] = {'value': float(amount)}
    bucket.update(nested)
    return bucket


def terms(rows):
    return {'doc_count_error_upper_bound': 0, 'sum_other_doc_count': 0, 'buckets': rows}


class AggregateReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        self.consumption_path = root / 'consumption.json'
        self.coin_path = root / 'coin-rewards.json'
        self.diamond_path = root / 'diamond-rewards.json'
        self.top_path = root / 'top-users.json'
        self.consumption = terms([
            term_bucket(161, 1, 500, 1),
            term_bucket(165, 1, 500, 1),
        ])
        self.coin = terms([term_bucket(81, 1, 188)])
        self.diamond = terms([
            {'key': 2, 'doc_count': 1,
             'by_gift': terms([term_bucket(81, 1, 131600, 1)])},
            {'key': 12, 'doc_count': 1,
             'by_gift': terms([term_bucket(81, 1, 10000, 1)])},
        ])
        self.top = [
            {'key': 11000001,
             'doc_count': 1, 'sum_tokens': {'value': 500.0}},
        ]
        self.overall = [
            term_bucket(11000001, 2, 1000, by_box=terms([
                term_bucket(161, 1, 500), term_bucket(165, 1, 500),
            ])),
        ]
        self.save_inputs()

    def save(self, path, aggregations):
        path.write_text(json.dumps({
            'timed_out': False, '_shards': {'failed': 0}, 'max_score': 0.0,
            'aggregations': aggregations,
        }), encoding='utf-8')

    def save_inputs(self):
        self.save(self.consumption_path, {'by_box': self.consumption})
        self.save(self.coin_path, {'by_gift': self.coin})
        self.save(self.diamond_path, {'by_type': self.diamond})
        self.save(self.top_path, {
            'by_box': terms([
                {'key': 161, 'doc_count': 1,
                 'by_user': {'sum_other_doc_count': 0, 'buckets': self.top}},
                {'key': 165, 'doc_count': 1,
                 'by_user': {'sum_other_doc_count': 0, 'buckets': self.top}},
            ]),
            'by_user_all': {'sum_other_doc_count': 0, 'buckets': self.overall},
        })

    def test_actual_sample_amounts_and_workbook(self):
        boxes, gifts = report.read_summaries(
            self.consumption_path, self.coin_path, self.diamond_path)
        top, overall, overall_by_box = report.read_top_users(self.top_path, boxes)
        self.assertEqual(boxes[161].spend_tokens, 500)
        self.assertEqual(boxes[165].coin_reward_tokens, 188)
        self.assertEqual(boxes[161].diamond_reward_tokens, 131600)
        self.assertEqual(boxes[165].diamond_reward_tokens, 10000)
        self.assertEqual(gifts[161, 81].main_records, 1)
        self.assertEqual(gifts[165, 81].main_records, 1)
        self.assertEqual(top[161][0], (11000001, 500, 1))
        self.assertEqual(overall[0], (11000001, 1000, 2))
        self.assertEqual(overall_by_box[11000001], {161: 500, 165: 500})
        self.assertEqual(Decimal(boxes[161].spend_tokens) / report.COIN_PER_USD
                         - Decimal(boxes[161].diamond_reward_tokens) / report.DIAMOND_PER_USD,
                         Decimal('0.03684'))
        book = report.build_report('2026-09-29 00:00:00', '2026-09-29 23:59:59',
                                   boxes, gifts, top, overall, overall_by_box,
                                   report.load_box_prices(report.DEFAULT_BOX_PRICES))
        data = io.BytesIO()
        book.save(data)
        data.seek(0)
        saved = load_workbook(data, data_only=False)
        self.assertEqual(saved.sheetnames, ['分析总览', '普通礼物盲盒', '幸运礼物盲盒',
                                           '全盲盒消费Top100', '普通消费Top30',
                                           '幸运消费Top30'])
        self.assertTrue(all(ws.sheet_view.showGridLines for ws in saved))
        self.assertEqual(saved['分析总览']['B9'].value, '消费记录数')
        self.assertEqual(saved['普通礼物盲盒']['B5'].value, '消费记录数')
        self.assertEqual(saved['普通礼物盲盒']['C5'].value, '开盒次数')
        self.assertEqual(saved['普通消费Top30']['F4'].value, '消费记录数')
        self.assertEqual(saved['分析总览']['B10'].value, 2)
        self.assertEqual(saved['普通消费Top30']['F5'].value, 1)
        all_top = saved['全盲盒消费Top100']
        self.assertEqual([all_top.cell(4, column).value for column in range(1, 6)],
                         ['名次', '用户ID', '消费金币', '消费美元', '消费记录数'])
        self.assertEqual([all_top.cell(4, column).value for column in range(6, 16)],
                         ['普通消费美元', '幸运消费美元']
                         + [f'{box}消费美元' for box in range(161, 169)])
        self.assertEqual([all_top.cell(5, column).value for column in range(1, 6)],
                         [1, 11000001, 1000, 0.1, 2])
        self.assertEqual(all_top['B5'].number_format, '0')
        self.assertEqual(all_top['C5'].number_format, '#,##0')
        self.assertEqual(all_top['D5'].number_format, report.MONEY_FORMAT)
        self.assertEqual(all_top['F5'].number_format, report.MONEY_FORMAT)
        self.assertEqual(all_top['F5'].value, 0.05)
        self.assertEqual(all_top['G5'].value, 0.05)
        self.assertEqual(all_top['H5'].value, 0.05)
        self.assertEqual(all_top['L5'].value, 0.05)
        self.assertEqual(all_top['D5'].value, all_top['F5'].value + all_top['G5'].value)
        self.assertEqual(all_top['D5'].value,
                         sum(all_top.cell(5, column).value for column in range(8, 16)))
        self.assertEqual(all_top.freeze_panes, 'C5')
        self.assertEqual(all_top.auto_filter.ref, 'A4:O5')
        self.assertEqual(saved['普通礼物盲盒']['N6'].value, 500)
        self.assertEqual(saved['普通礼物盲盒']['O6'].value, 500)
        self.assertEqual(saved['普通礼物盲盒']['P6'].value, 0)
        self.assertEqual(saved['幸运礼物盲盒']['N6'].value, 500)
        self.assertEqual(saved['分析总览']['A4'].value, '统计开始（含）')
        self.assertEqual(saved['分析总览']['L10'].value, 0.067)
        self.assertEqual(saved['分析总览']['L11'].value, 0.037)
        self.assertEqual(saved['分析总览']['L12'].value, 0.03)
        self.assertEqual(saved['普通礼物盲盒']['C15'].value, 1)
        self.assertEqual(saved['幸运礼物盲盒']['F15'].value, 1)
        self.assertEqual(saved['普通消费Top30']['C5'].value, 11000001)
        self.assertEqual(saved['幸运消费Top30']['C5'].value, 11000001)
        self.assertFalse(any(cell.data_type == 'f' for ws in saved for row in ws for cell in row))

    def test_query_generation_and_report_command(self):
        root = Path(self.temp.name)
        query_dir = root / 'queries'
        report.write_queries(argparse.Namespace(
            start='2026-09-29 00:00:00', end='2026-09-29 23:59:59',
            output_dir=query_dir))
        self.assertEqual(len(list(query_dir.glob('*.json'))), 4)
        top_query = json.loads((query_dir / 'top-users.json').read_text(encoding='utf-8'))
        by_user = top_query['aggs']['by_box']['aggs']['by_user']['terms']
        self.assertEqual(by_user['size'], 30)
        self.assertEqual(by_user['shard_size'], 300)
        self.assertEqual(by_user['order'], [{'sum_tokens': 'desc'}, {'_key': 'asc'}])
        by_user_all = top_query['aggs']['by_user_all']['terms']
        self.assertEqual(by_user_all['size'], 100)
        self.assertEqual(by_user_all['shard_size'], 1000)
        self.assertEqual(by_user_all['order'], [{'sum_tokens': 'desc'}, {'_key': 'asc'}])
        self.assertEqual(top_query['aggs']['by_user_all']['aggs']['by_box']['terms']['size'], 8)
        self.assertEqual(top_query['query']['bool']['filter'],
                         json.loads((query_dir / 'consumption.json').read_text(encoding='utf-8'))['query']['bool']['filter'])
        coin_query = json.loads((query_dir / 'coin-rewards.json').read_text(encoding='utf-8'))
        self.assertEqual(coin_query['query']['bool']['filter'][0]['terms']['business_type'], [16])
        output = root / 'report.xlsx'
        report.write_report(argparse.Namespace(
            consumption=self.consumption_path, coin_rewards=self.coin_path,
            diamond_rewards=self.diamond_path, top_users=self.top_path,
            start='2026-09-29 00:00:00', end='2026-09-29 23:59:59',
            output=output))
        self.assertTrue(output.exists())
        self.assertEqual(load_workbook(output)['分析总览']['L10'].value, 0.067)

    def test_all_history_uses_fixed_end_without_start(self):
        root = Path(self.temp.name)
        query_dir = root / 'all-history-queries'
        report.write_queries(argparse.Namespace(
            start=None, all_history=True, end='2026-09-29 23:59:59',
            output_dir=query_dir))
        query = json.loads((query_dir / 'consumption.json').read_text(encoding='utf-8'))
        time_filter = query['query']['bool']['filter'][-1]['range']['create_time']
        self.assertEqual(time_filter, {'lte': '2026-09-29 23:59:59'})
        output = root / 'all-history-report.xlsx'
        report.write_report(argparse.Namespace(
            consumption=self.consumption_path, coin_rewards=self.coin_path,
            diamond_rewards=self.diamond_path, top_users=self.top_path,
            start=None, all_history=True, end='2026-09-29 23:59:59',
            output=output))
        self.assertEqual(load_workbook(output)['分析总览']['B4'].value, '全部历史')

    def test_diamond_commission_is_cost_not_another_opening(self):
        self.diamond['buckets'].append({
            'key': 3, 'doc_count': 1,
            'by_gift': terms([term_bucket(81, 1, 10000, 1)]),
        })
        self.save_inputs()
        boxes, gifts = report.read_summaries(
            self.consumption_path, self.coin_path, self.diamond_path)
        self.assertEqual(boxes[161].main_records, 1)
        self.assertEqual(gifts[161, 81].share_diamonds, 10000)
        self.assertEqual(report.summary_values(161, boxes[161])[11], 0.036)

    def test_top30_uses_only_returned_users(self):
        self.consumption = terms([term_bucket(161, 32, 32000, 32)])
        self.coin = terms([])
        self.diamond = terms([])
        self.top = [
            {'key': 10000 + n,
             'doc_count': 1, 'sum_tokens': {'value': 1000.0}}
            for n in range(1, 31)
        ]
        self.save(self.consumption_path, {'by_box': self.consumption})
        self.save(self.coin_path, {'by_gift': self.coin})
        self.save(self.diamond_path, {'by_type': self.diamond})
        self.overall = self.top + [
            {'key': 10031, 'doc_count': 1, 'sum_tokens': {'value': 1000.0}},
            {'key': 10032, 'doc_count': 1, 'sum_tokens': {'value': 1000.0}},
        ]
        self.overall = [term_bucket(row['key'], 1, 1000, by_box=terms([
            term_bucket(161, 1, 1000),
        ])) for row in self.overall]
        self.save(self.top_path, {
            'by_box': terms([
                {'key': 161, 'doc_count': 32,
                 'by_user': {'sum_other_doc_count': 2, 'buckets': self.top}},
            ]),
            'by_user_all': {'buckets': self.overall},
        })
        boxes, _ = report.read_summaries(
            self.consumption_path, self.coin_path, self.diamond_path)
        top, overall, _ = report.read_top_users(self.top_path, boxes)
        self.assertEqual(len(top[161]), 30)
        self.assertEqual(len(overall), 32)
        self.assertEqual(top[161][0][0], 10001)
        self.assertEqual(top[161][-1][0], 10030)
        self.top.append({'key': 10031, 'doc_count': 1,
                         'sum_tokens': {'value': 1000.0}})
        self.save(self.top_path, {
            'by_box': terms([
                {'key': 161, 'doc_count': 32, 'by_user': {'buckets': self.top}},
            ]),
            'by_user_all': {'buckets': self.overall},
        })
        with self.assertRaisesRegex(ValueError, '超过 30'):
            report.read_top_users(self.top_path, boxes)

    def test_overall_top100_combines_spending_across_boxes(self):
        self.save(self.consumption_path, {'by_box': terms([
            term_bucket(161, 2, 1300, 2),
            term_bucket(165, 1, 500, 1),
        ])})
        self.save(self.top_path, {
            'by_box': terms([
                {'key': 161, 'doc_count': 2, 'by_user': {'buckets': [
                    term_bucket(11000002, 1, 800),
                    term_bucket(11000001, 1, 500),
                ]}},
                {'key': 165, 'doc_count': 1, 'by_user': {'buckets': [
                    term_bucket(11000001, 1, 500),
                ]}},
            ]),
            'by_user_all': {'buckets': [
                term_bucket(11000001, 2, 1000, by_box=terms([
                    term_bucket(161, 1, 500), term_bucket(165, 1, 500),
                ])),
                term_bucket(11000002, 1, 800, by_box=terms([
                    term_bucket(161, 1, 800),
                ])),
            ]},
        })
        boxes, gifts = report.read_summaries(
            self.consumption_path, self.coin_path, self.diamond_path)
        by_box, overall, overall_by_box = report.read_top_users(self.top_path, boxes)
        self.assertEqual(overall, [(11000001, 1000, 2), (11000002, 800, 1)])
        self.assertEqual(by_box[161][0], (11000002, 800, 1))
        book = report.build_report('2026-09-29 00:00:00', '2026-09-29 23:59:59',
                                   boxes, gifts, by_box, overall, overall_by_box,
                                   report.load_box_prices(report.DEFAULT_BOX_PRICES))
        sheet = book['全盲盒消费Top100']
        self.assertEqual([sheet.cell(5, col).value for col in range(1, 6)],
                         [1, 11000001, 1000, 0.1, 2])
        self.assertEqual([sheet.cell(6, col).value for col in range(1, 6)],
                         [2, 11000002, 800, 0.08, 1])
        self.assertEqual([sheet.cell(5, col).value for col in range(6, 10)],
                         [0.05, 0.05, 0.05, 0])
        self.assertEqual([sheet.cell(6, col).value for col in range(6, 10)],
                         [0.08, 0, 0.08, 0])

        response = json.loads(self.top_path.read_text(encoding='utf-8'))
        response['aggregations']['by_user_all']['buckets'][0]['by_box']['buckets'][0]['sum_tokens']['value'] = 400.0
        self.top_path.write_text(json.dumps(response), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '之和与总消费不一致'):
            report.read_top_users(self.top_path, boxes)
        response['aggregations']['by_user_all']['buckets'][0]['by_box']['buckets'][0]['sum_tokens']['value'] = 500.0
        response['aggregations']['by_user_all']['buckets'] = [
            term_bucket(10000 + user, 1, 1) for user in range(101)
        ]
        self.top_path.write_text(json.dumps(response), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '超过 100'):
            report.read_top_users(self.top_path, boxes)

    def test_rejects_inexact_sum_and_incomplete_terms(self):
        self.consumption['buckets'][0]['sum_tokens']['value'] = float(2**53)
        self.save_inputs()
        with self.assertRaisesRegex(ValueError, '双精度'):
            report.read_summaries(self.consumption_path, self.coin_path, self.diamond_path)
        self.consumption['buckets'][0]['sum_tokens']['value'] = 500.0
        self.consumption['sum_other_doc_count'] = 1
        self.save_inputs()
        with self.assertRaisesRegex(ValueError, '未返回'):
            report.read_summaries(self.consumption_path, self.coin_path, self.diamond_path)

    def test_nonoverlapping_summary_parts_are_added_as_integers(self):
        part_dir = Path(self.temp.name) / 'consumption-parts'
        part_dir.mkdir()
        self.save(part_dir / 'part-01.json', {'by_box': terms([
            term_bucket(161, 1, 200, 1),
        ])})
        self.save(part_dir / 'part-02.json', {'by_box': terms([
            term_bucket(161, 1, 300, 1),
            term_bucket(165, 1, 500, 1),
        ])})
        boxes, _ = report.read_summaries(part_dir, self.coin_path, self.diamond_path)
        self.assertEqual(boxes[161].spend_tokens, 500)
        self.assertEqual(boxes[161].spend_records, 2)
        self.assertEqual(boxes[165].spend_tokens, 500)

    def test_price_config_and_consumption_difference(self):
        prices = report.load_box_prices(report.DEFAULT_BOX_PRICES)
        self.assertEqual(prices, {
            161: 500, 162: 1000, 163: 5000, 164: 10000,
            165: 500, 166: 1000, 167: 5000, 168: 10000,
        })
        config = Path(self.temp.name) / 'prices.json'
        config.write_text(json.dumps({'coin_tokens_per_open': {'161': 500}}), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '恰好配置'):
            report.load_box_prices(config)
        custom = {str(box): value for box, value in prices.items()}
        custom['162'] = 0
        config.write_text(json.dumps({'coin_tokens_per_open': custom}), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '正整数'):
            report.load_box_prices(config)
        custom['162'] = 1000
        custom['161'] = 600
        config.write_text(json.dumps({'coin_tokens_per_open': custom}), encoding='utf-8')
        boxes, gifts = report.read_summaries(
            self.consumption_path, self.coin_path, self.diamond_path)
        top, overall, overall_by_box = report.read_top_users(self.top_path, boxes)
        book = report.build_report('2026-09-29 00:00:00', '2026-09-29 23:59:59',
                                   boxes, gifts, top, overall, overall_by_box,
                                   report.load_box_prices(config))
        sheet = book['普通礼物盲盒']
        self.assertEqual(sheet['N6'].value, 600)
        self.assertEqual(sheet['O6'].value, 600)
        self.assertEqual(sheet['P6'].value, -100)
        self.assertEqual(sheet['P6'].font.color.rgb[-6:], 'C62828')

        # 一条金币消费记录可以开多次盒；记录数与 amount 之和不能混用。
        many_opens = {box: report.BoxTotals() for box in report.BOX_GIFTS}
        many_opens[161] = report.BoxTotals(spend_records=1, spend_amount=3,
                                          spend_tokens=1500)
        empty_gifts = {(box, gift): report.GiftTotals()
                       for box, gifts_for_box in report.BOX_GIFTS.items()
                       for gift in gifts_for_box}
        ws = Workbook().active
        report.write_box_sheet(ws, '普通礼物盲盒', report.NORMAL_BOXES,
                               many_opens, empty_gifts, prices)
        self.assertEqual(ws['B6'].value, 1)
        self.assertEqual(ws['C6'].value, 3)
        self.assertEqual(ws['O6'].value, 1500)
        self.assertEqual(ws['P6'].value, 0)


if __name__ == '__main__':
    unittest.main()
