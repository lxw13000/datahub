"""用模拟 ES 响应验证四次查询、每盒 Top 30、全盲盒 Top 100 和恢复。"""

import argparse
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import fetch_es
from export_report import write_report


def response(name, buckets, **extra_aggregations):
    aggregation = {'buckets': buckets}
    root = {
        'timed_out': False,
        '_shards': {'failed': 0},
        'aggregations': {name: aggregation, **extra_aggregations},
    }
    content = json.dumps(root, separators=(',', ':')).encode('utf-8')
    return content, root


class FakeES:
    def __init__(self, url='https://test-es:9200'):
        self.config = fetch_es.ESConfig(url, 'test_user', 'private_password', 30, None)
        self.calls = []

    def search(self, index, query):
        self.calls.append((index, query))
        name = next(iter(query['aggs']))
        if name == 'by_box':
            if 'by_user' in query['aggs']['by_box'].get('aggs', {}):
                return response(name, [{
                    'key': 161, 'doc_count': 32,
                    'by_user': {'sum_other_doc_count': 2, 'buckets': [{
                        'key': user, 'doc_count': 1,
                        'sum_tokens': {'value': 1000.0},
                    } for user in range(1, 31)]},
                }], by_user_all={'buckets': [{
                    'key': user, 'doc_count': 1,
                    'sum_tokens': {'value': 1000.0},
                    'by_box': {'buckets': [{
                        'key': 161, 'doc_count': 1,
                        'sum_tokens': {'value': 1000.0},
                    }]},
                } for user in range(1, 33)]})
            return response(name, [{
                'key': 161, 'doc_count': 32,
                'sum_tokens': {'value': 32000.0}, 'sum_amount': {'value': 32.0},
            }])
        if name in {'by_gift', 'by_type'}:
            return response(name, [])
        raise AssertionError(f'未知查询：{name}')


class FetchESTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.output_dir = self.root / 'responses'

    def args(self, *, resume=False):
        return argparse.Namespace(
            config=self.root / 'config.json', start=None, all_history=True,
            end='2026-09-29 23:59:59', output_dir=self.output_dir,
            resume=resume, report_output=None,
        )

    def test_export_four_queries_and_builds_report(self):
        client = FakeES()
        fetch_es.export_all(self.args(), client)
        self.assertEqual(len(client.calls), 4)
        self.assertEqual([item[0] for item in client.calls], [
            'sano_wallet_coin_record', 'sano_wallet_coin_record',
            'sano_wallet_diamond_record', 'sano_wallet_coin_record',
        ])
        top = json.loads((self.output_dir / 'top-users.json').read_text())
        self.assertEqual(len(top['aggregations']['by_box']['buckets'][0]['by_user']['buckets']), 30)
        self.assertEqual(len(top['aggregations']['by_user_all']['buckets']), 32)
        self.assertEqual(top['aggregations']['by_user_all']['buckets'][0]['by_box']['buckets'][0]['key'], 161)
        manifest = (self.output_dir / 'manifest.json').read_text(encoding='utf-8')
        self.assertIn('https://test-es:9200', manifest)
        self.assertNotIn('private_password', manifest)
        self.assertEqual(len(list((self.output_dir / 'queries').glob('*.json'))), 4)

        output = self.root / 'report.xlsx'
        write_report(argparse.Namespace(
            consumption=self.output_dir / 'consumption.json',
            coin_rewards=self.output_dir / 'coin-rewards.json',
            diamond_rewards=self.output_dir / 'diamond-rewards.json',
            top_users=self.output_dir / 'top-users.json',
            start=None, all_history=True, end='2026-09-29 23:59:59', output=output,
        ))
        book = load_workbook(output, read_only=True)
        self.assertEqual(len(book.sheetnames), 6)
        self.assertEqual(book['全盲盒消费Top100']['B5'].value, 1)
        self.assertEqual(book['全盲盒消费Top100']['B36'].value, 32)
        self.assertEqual(book['全盲盒消费Top100']['H5'].value, 0.1)
        self.assertEqual(book['普通消费Top30']['C5'].value, 1)
        self.assertEqual(book['普通消费Top30']['C34'].value, 30)
        book.close()

        resumed = FakeES()
        fetch_es.export_all(self.args(resume=True), resumed)
        self.assertEqual(resumed.calls, [])
        with self.assertRaisesRegex(ValueError, '不一致'):
            fetch_es.export_all(self.args(resume=True), FakeES('https://other-es:9200'))

    def test_config_validates_credentials_and_example(self):
        config = self.root / 'config.json'
        config.write_text(json.dumps({
            'es_url': 'https://your-es-host:9200',
            'username': 'your_username', 'password': 'your_password',
        }), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '示例值'):
            fetch_es.load_config(config)
        config.write_text(json.dumps({
            'es_url': 'https://user:secret@test-es:9200',
            'username': 'user', 'password': 'secret',
        }), encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '不能包含账号'):
            fetch_es.load_config(config)
        config.write_text(json.dumps({
            'es_url': 'https://test-es:9200',
            'username': 'user', 'password': 'secret', 'timeout_seconds': 120,
        }), encoding='utf-8')
        self.assertEqual(fetch_es.load_config(config).timeout_seconds, 120)

    def test_client_posts_to_correct_index_with_basic_auth(self):
        class MockResponse:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

            def read(self):
                return response('by_box', [])[0]

        client = fetch_es.ESClient(fetch_es.ESConfig(
            'https://test-es:9200', 'user', 'secret', 30, None))
        with patch.object(fetch_es, 'urlopen', return_value=MockResponse()) as mocked:
            _, result = client.search('sano_wallet_coin_record', {'size': 0})
        request = mocked.call_args.args[0]
        self.assertEqual(request.full_url,
                         'https://test-es:9200/sano_wallet_coin_record/_search')
        self.assertEqual(request.get_method(), 'POST')
        self.assertEqual(json.loads(request.data), {'size': 0})
        self.assertEqual(request.get_header('Authorization'), 'Basic dXNlcjpzZWNyZXQ=')
        self.assertIn('by_box', result['aggregations'])


if __name__ == '__main__':
    unittest.main()
