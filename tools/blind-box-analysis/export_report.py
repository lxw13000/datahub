"""读取 ES 聚合 JSON，生成盲盒消费、开奖及平台利润 Excel。"""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BOX_GIFTS = {
    161: (81, 82, 84, 86, 88),
    162: (174, 175, 176, 177, 178),
    163: (179, 180, 181, 182, 183),
    164: (169, 170, 171, 172, 173),
    165: (81, 82, 84, 86, 88),
    166: (174, 175, 176, 177, 178),
    167: (179, 180, 181, 182, 183),
    168: (169, 170, 171, 172, 173),
}
NORMAL_BOXES = frozenset((161, 162, 163, 164))
LUCKY_BOXES = frozenset((165, 166, 167, 168))
ALL_GIFTS = sorted({gift for gifts in BOX_GIFTS.values() for gift in gifts})
NORMAL_GIFT_TO_BOX = {gift: box for box in NORMAL_BOXES for gift in BOX_GIFTS[box]}
LUCKY_GIFT_TO_BOX = {gift: box for box in LUCKY_BOXES for gift in BOX_GIFTS[box]}
COIN_PER_USD = Decimal(10_000)
DIAMOND_PER_USD = Decimal(10_000_000)
MAX_EXACT_ES_INTEGER = 2**53 - 1
TOP_USER_COUNT = 30
TOP_USER_SHARD_SIZE = 300
ALL_TOP_USER_COUNT = 100
ALL_TOP_USER_SHARD_SIZE = 1000
DEFAULT_BOX_PRICES = Path(__file__).with_name('box-prices.json')
TOKEN_FORMAT = '#,##0'
ID_FORMAT = '0'
MONEY_FORMAT = '#,##0.000;[Red]-#,##0.000;0.000'
PERCENT_FORMAT = '0.00%;[Red]-0.00%;0.00%'


@dataclass
class BoxTotals:
    spend_records: int = 0
    spend_amount: int = 0
    spend_tokens: int = 0
    main_records: int = 0
    main_amount: int = 0
    coin_win_records: int = 0
    coin_reward_tokens: int = 0
    diamond_reward_tokens: int = 0

    def add(self, other: BoxTotals) -> None:
        for field in self.__dataclass_fields__:
            setattr(self, field, getattr(self, field) + getattr(other, field))


@dataclass
class GiftTotals:
    main_records: int = 0
    main_amount: int = 0
    coin_win_records: int = 0
    receiver_diamonds: int = 0
    share_diamonds: int = 0
    coin_reward_tokens: int = 0


def number(value: Any, label: str, *, minimum: int = 0, es_sum: bool = False) -> int:
    if isinstance(value, bool):
        raise ValueError(f'{label} 必须是整数')
    try:
        decimal = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise ValueError(f'{label} 缺失或不是数字：{value!r}') from exc
    if not decimal.is_finite() or decimal != decimal.to_integral_value() or decimal < minimum:
        raise ValueError(f'{label} 必须是大于等于 {minimum} 的整数：{value!r}')
    if es_sum and decimal > MAX_EXACT_ES_INTEGER:
        raise ValueError(
            f'{label}={value} 超过 ES sum 双精度整数安全范围；请缩小查询时间段后分别汇总'
        )
    return int(decimal)


def load_box_prices(path: Path) -> dict[int, int]:
    with path.open('r', encoding='utf-8-sig') as source:
        raw = json.load(source)
    entries = raw.get('coin_tokens_per_open') if isinstance(raw, dict) else None
    expected = {str(box) for box in BOX_GIFTS}
    if not isinstance(entries, dict) or set(entries) != expected:
        raise ValueError(f'{path}: coin_tokens_per_open 必须恰好配置 161–168 八个盲盒')
    prices: dict[int, int] = {}
    for box_text, value in entries.items():
        if type(value) is not int or value <= 0:
            raise ValueError(f'{path}: 盲盒 {box_text} 单价必须是正整数金币')
        prices[int(box_text)] = value
    return prices


def read_aggregation(path: Path, name: str) -> dict[str, Any]:
    with path.open('r', encoding='utf-8-sig') as source:
        root = json.load(source, parse_int=int, parse_float=Decimal)
    if not isinstance(root, dict):
        raise ValueError(f'{path}: 必须保存完整的 ES _search 响应对象')
    shards = root.get('_shards') or {}
    if root.get('timed_out') is True or shards.get('failed', 0):
        raise ValueError(f'{path}: ES 查询超时或有分片失败，聚合结果不能用于利润报表')
    if 'error' in root:
        raise ValueError(f'{path}: ES 返回错误：{root["error"]}')
    aggregations = root.get('aggregations', root.get('aggs'))
    if not isinstance(aggregations, dict) or not isinstance(aggregations.get(name), dict):
        raise ValueError(f'{path}: 缺少 aggregations.{name}')
    return aggregations[name]


def buckets(aggregation: dict[str, Any], label: str, *, terms: bool = True) -> list[dict[str, Any]]:
    rows = aggregation.get('buckets')
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise ValueError(f'{label}: buckets 必须是对象数组')
    if terms:
        if number(aggregation.get('sum_other_doc_count', 0), f'{label}.sum_other_doc_count') != 0:
            raise ValueError(f'{label}: 有未返回的 terms 桶，请检查 DSL 的 size')
        if number(aggregation.get('doc_count_error_upper_bound', 0),
                  f'{label}.doc_count_error_upper_bound') != 0:
            raise ValueError(f'{label}: terms 文档数不精确，请按 _key 排序并覆盖全部候选值')
    return rows


def metric(row: dict[str, Any], name: str, label: str) -> int:
    item = row.get(name)
    if not isinstance(item, dict):
        raise ValueError(f'{label}: 缺少 {name}')
    return number(item.get('value'), f'{label}.{name}.value', es_sum=True)


def source_files(path: Path) -> list[Path]:
    if path.is_file():
        return [path]
    if path.is_dir():
        files = sorted(path.glob('*.json'))
        if files:
            return files
    raise ValueError(f'{path}: 未找到 ES 聚合响应 JSON 文件')


def read_summaries(
    consumption_path: Path,
    coin_path: Path,
    diamond_path: Path,
) -> tuple[dict[int, BoxTotals], dict[tuple[int, int], GiftTotals]]:
    box_totals = {box: BoxTotals() for box in BOX_GIFTS}
    gift_totals = {(box, gift): GiftTotals() for box, gifts in BOX_GIFTS.items()
                   for gift in gifts}

    for path in source_files(consumption_path):
        consumption = read_aggregation(path, 'by_box')
        seen_boxes: set[int] = set()
        for row in buckets(consumption, f'{path}: by_box'):
            box = number(row.get('key'), '消费盲盒ID', minimum=1)
            if box not in box_totals or box in seen_boxes:
                raise ValueError(f'{path}: 未配置或重复的盲盒 ID {box}')
            seen_boxes.add(box)
            box_totals[box].spend_records += number(row.get('doc_count'), f'盲盒 {box} 消费记录数')
            box_totals[box].spend_tokens += metric(row, 'sum_tokens', f'盲盒 {box}')
            box_totals[box].spend_amount += metric(row, 'sum_amount', f'盲盒 {box}')

    for path in source_files(coin_path):
        coin = read_aggregation(path, 'by_gift')
        seen_coin_gifts: set[int] = set()
        for row in buckets(coin, f'{path}: by_gift'):
            gift = number(row.get('key'), '金币返奖礼物ID', minimum=1)
            if gift not in LUCKY_GIFT_TO_BOX or gift in seen_coin_gifts:
                raise ValueError(f'{path}: 未配置或重复的金币返奖礼物 ID {gift}')
            seen_coin_gifts.add(gift)
            box = LUCKY_GIFT_TO_BOX[gift]
            count = number(row.get('doc_count'), f'礼物 {gift} 金币中奖记录数')
            tokens = metric(row, 'sum_tokens', f'礼物 {gift} 金币返奖')
            box_totals[box].coin_win_records += count
            box_totals[box].coin_reward_tokens += tokens
            gift_totals[box, gift].coin_win_records += count
            gift_totals[box, gift].coin_reward_tokens += tokens

    for path in source_files(diamond_path):
        diamond = read_aggregation(path, 'by_type')
        seen_types: set[int] = set()
        for type_row in buckets(diamond, f'{path}: by_type'):
            business_type = number(type_row.get('key'), '钻石业务类型', minimum=1)
            if business_type not in {2, 3, 4, 5, 12, 13, 14} or business_type in seen_types:
                raise ValueError(f'{path}: 未配置或重复的钻石业务类型 {business_type}')
            seen_types.add(business_type)
            gift_to_box = NORMAL_GIFT_TO_BOX if business_type in {2, 3, 4, 5} else LUCKY_GIFT_TO_BOX
            by_gift = type_row.get('by_gift')
            if not isinstance(by_gift, dict):
                raise ValueError(f'{path}: 类型 {business_type} 缺少 by_gift')
            seen_gifts: set[int] = set()
            for row in buckets(by_gift, f'{path}: 类型 {business_type} by_gift'):
                gift = number(row.get('key'), '钻石返奖礼物ID', minimum=1)
                if gift not in gift_to_box or gift in seen_gifts:
                    raise ValueError(f'{path}: 类型 {business_type} 有未配置或重复的礼物 {gift}')
                seen_gifts.add(gift)
                box = gift_to_box[gift]
                count = number(row.get('doc_count'), f'类型 {business_type} 礼物 {gift} 记录数')
                tokens = metric(row, 'sum_tokens', f'类型 {business_type} 礼物 {gift}')
                box_totals[box].diamond_reward_tokens += tokens
                gift_totals[box, gift].receiver_diamonds += tokens if business_type in {2, 12} else 0
                gift_totals[box, gift].share_diamonds += tokens if business_type not in {2, 12} else 0
                if business_type in {2, 12}:
                    amount = metric(row, 'sum_amount', f'类型 {business_type} 礼物 {gift}')
                    box_totals[box].main_records += count
                    box_totals[box].main_amount += amount
                    gift_totals[box, gift].main_records += count
                    gift_totals[box, gift].main_amount += amount

    return box_totals, gift_totals


def read_top_users(
    path: Path,
    box_totals: dict[int, BoxTotals],
) -> tuple[dict[int, list[tuple[int, int, int]]], list[tuple[int, int, int]], dict[int, dict[int, int]]]:
    aggregation = read_aggregation(path, 'by_box')
    top_users: dict[int, list[tuple[int, int, int]]] = {box: [] for box in BOX_GIFTS}
    seen_boxes: set[int] = set()
    for box_row in buckets(aggregation, f'{path}: by_box'):
        box = number(box_row.get('key'), 'Top 用户盲盒ID', minimum=1)
        if box not in BOX_GIFTS or box in seen_boxes:
            raise ValueError(f'{path}: 未配置或重复的 Top 用户盲盒 ID {box}')
        seen_boxes.add(box)
        by_user = box_row.get('by_user')
        if not isinstance(by_user, dict):
            raise ValueError(f'{path}: 盲盒 {box} 缺少 by_user')
        # Top 30 的 terms 会丢弃其余用户桶；这不是完整用户汇总，不检查 sum_other_doc_count。
        rows = buckets(by_user, f'{path}: 盲盒 {box} by_user', terms=False)
        if len(rows) > TOP_USER_COUNT:
            raise ValueError(f'{path}: 盲盒 {box} 返回超过 {TOP_USER_COUNT} 个用户')
        if box_totals[box].spend_records > 0 and not rows:
            raise ValueError(f'{path}: 盲盒 {box} 有消费但没有 Top 用户桶')
        seen_users: set[int] = set()
        for row in rows:
            user = number(row.get('key'), 'Top 用户ID', minimum=1)
            if user in seen_users:
                raise ValueError(f'{path}: 盲盒 {box} 出现重复用户 {user}')
            seen_users.add(user)
            count = number(row.get('doc_count'), f'盲盒 {box} 用户 {user} 消费记录数', minimum=1)
            tokens = metric(row, 'sum_tokens', f'盲盒 {box} 用户 {user}')
            top_users[box].append((user, tokens, count))
        top_users[box].sort(key=lambda item: (-item[1], item[0]))
        if (sum(item[1] for item in top_users[box]) > box_totals[box].spend_tokens
                or sum(item[2] for item in top_users[box]) > box_totals[box].spend_records):
            raise ValueError(f'{path}: 盲盒 {box} Top 用户合计超过消费汇总，请检查查询时间和数据')
    missing = {box for box, total in box_totals.items()
               if total.spend_records > 0 and box not in seen_boxes}
    if missing:
        raise ValueError(f'{path}: 有消费却缺少 Top 用户盲盒桶：{sorted(missing)}')

    overall_aggregation = read_aggregation(path, 'by_user_all')
    overall_rows = buckets(overall_aggregation, f'{path}: by_user_all', terms=False)
    if len(overall_rows) > ALL_TOP_USER_COUNT:
        raise ValueError(f'{path}: 全盲盒返回超过 {ALL_TOP_USER_COUNT} 个用户')
    total_spend_records = sum(item.spend_records for item in box_totals.values())
    total_spend_tokens = sum(item.spend_tokens for item in box_totals.values())
    if total_spend_records > 0 and not overall_rows:
        raise ValueError(f'{path}: 有消费但没有全盲盒 Top 用户桶')
    overall_users: list[tuple[int, int, int]] = []
    overall_by_box: dict[int, dict[int, int]] = {}
    seen_overall_users: set[int] = set()
    for row in overall_rows:
        user = number(row.get('key'), '全盲盒 Top 用户ID', minimum=1)
        if user in seen_overall_users:
            raise ValueError(f'{path}: 全盲盒 Top 用户重复：{user}')
        seen_overall_users.add(user)
        count = number(row.get('doc_count'), f'全盲盒用户 {user} 消费记录数', minimum=1)
        tokens = metric(row, 'sum_tokens', f'全盲盒用户 {user}')
        by_box = row.get('by_box')
        if not isinstance(by_box, dict):
            raise ValueError(f'{path}: 全盲盒用户 {user} 缺少 by_box 消费结构')
        breakdown: dict[int, int] = {}
        for box_row in buckets(by_box, f'{path}: 全盲盒用户 {user} by_box'):
            box = number(box_row.get('key'), f'全盲盒用户 {user} 盲盒ID', minimum=1)
            if box not in BOX_GIFTS or box in breakdown:
                raise ValueError(f'{path}: 全盲盒用户 {user} 出现未配置或重复盲盒 {box}')
            breakdown[box] = metric(box_row, 'sum_tokens', f'全盲盒用户 {user} 盲盒 {box}')
        if sum(breakdown.values()) != tokens:
            raise ValueError(f'{path}: 全盲盒用户 {user} 各盲盒消费金币之和与总消费不一致')
        overall_by_box[user] = breakdown
        overall_users.append((user, tokens, count))
    overall_users.sort(key=lambda item: (-item[1], item[0]))
    if (sum(item[1] for item in overall_users) > total_spend_tokens
            or sum(item[2] for item in overall_users) > total_spend_records):
        raise ValueError(f'{path}: 全盲盒 Top 用户合计超过消费汇总，请检查查询时间和数据')
    return top_users, overall_users, overall_by_box


def search_filter(start: str | None, end: str, business_types: list[int], prop_ids: list[int], status: int) -> dict[str, Any]:
    time_range = {'lte': end}
    if start is not None:
        time_range['gte'] = start
    return {'bool': {'filter': [
        {'terms': {'business_type': business_types}},
        {'terms': {'prop_id': prop_ids}},
        {'term': {'status': status}},
        {'range': {'create_time': time_range}},
    ]}}


def term_group(field: str, size: int, *, with_amount: bool) -> dict[str, Any]:
    aggs = {'sum_tokens': {'sum': {'field': 'tokens'}}}
    if with_amount:
        aggs['sum_amount'] = {'sum': {'field': 'amount'}}
    return {'terms': {'field': field, 'size': size, 'order': {'_key': 'asc'}}, 'aggs': aggs}


def query_bodies(start: str | None, end: str) -> dict[str, dict[str, Any]]:
    consumption_filter = search_filter(start, end, [1, 11], sorted(BOX_GIFTS), -1)
    gift_filter = search_filter(start, end, [2, 3, 4, 5, 12, 13, 14], ALL_GIFTS, 1)
    coin_filter = search_filter(start, end, [16], ALL_GIFTS, 1)
    top_query = {
        'size': 0, 'track_total_hits': False, 'query': consumption_filter,
        'aggs': {
            'by_box': {
                'terms': {'field': 'prop_id', 'size': len(BOX_GIFTS), 'order': {'_key': 'asc'}},
                'aggs': {'by_user': {
                    'terms': {
                        'field': 'user_id', 'size': TOP_USER_COUNT,
                        'shard_size': TOP_USER_SHARD_SIZE,
                        'order': [{'sum_tokens': 'desc'}, {'_key': 'asc'}],
                    },
                    'aggs': {'sum_tokens': {'sum': {'field': 'tokens'}}},
                }},
            },
            'by_user_all': {
                'terms': {
                    'field': 'user_id', 'size': ALL_TOP_USER_COUNT,
                    'shard_size': ALL_TOP_USER_SHARD_SIZE,
                    'order': [{'sum_tokens': 'desc'}, {'_key': 'asc'}],
                },
                'aggs': {
                    'sum_tokens': {'sum': {'field': 'tokens'}},
                    'by_box': term_group('prop_id', len(BOX_GIFTS), with_amount=False),
                },
            },
        },
    }
    return {
        'consumption.json': {
            'size': 0, 'track_total_hits': False, 'query': consumption_filter,
            'aggs': {'by_box': term_group('prop_id', len(BOX_GIFTS), with_amount=True)},
        },
        'coin-rewards.json': {
            'size': 0, 'track_total_hits': False, 'query': coin_filter,
            'aggs': {'by_gift': term_group('prop_id', len(ALL_GIFTS), with_amount=False)},
        },
        'diamond-rewards.json': {
            'size': 0, 'track_total_hits': False, 'query': gift_filter,
            'aggs': {'by_type': {
                'terms': {'field': 'business_type', 'size': 7, 'order': {'_key': 'asc'}},
                'aggs': {'by_gift': term_group('prop_id', len(ALL_GIFTS), with_amount=True)},
            }},
        },
        'top-users.json': top_query,
    }


def write_queries(args: argparse.Namespace) -> None:
    start = None if getattr(args, 'all_history', False) else args.start
    validate_period(start, args.end)
    if args.output_dir.exists() and list(args.output_dir.glob('*.json')):
        raise ValueError(f'{args.output_dir} 已有 JSON；请使用空目录，避免混入旧查询')
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for name, body in query_bodies(start, args.end).items():
        with (args.output_dir / name).open('w', encoding='utf-8') as target:
            json.dump(body, target, ensure_ascii=False, indent=2)
            target.write('\n')
    print(f'已生成 4 份 ES 聚合查询 DSL：{args.output_dir}')


def validate_period(start: str | None, end: str) -> None:
    from datetime import datetime
    try:
        last = datetime.strptime(end, '%Y-%m-%d %H:%M:%S')
        first = datetime.strptime(start, '%Y-%m-%d %H:%M:%S') if start is not None else None
    except ValueError as exc:
        raise ValueError('时间必须使用 YYYY-MM-DD HH:MM:SS 格式') from exc
    if first is not None and first > last:
        raise ValueError('--start 不能晚于 --end')


def money(value: Decimal) -> float:
    # ES 返回的各类 tokens 汇总先以整数处理；仅写 Excel 美元列时进行三位小数舍入。
    return float(value.quantize(Decimal('0.001'), rounding=ROUND_HALF_UP))


def summarize(boxes: set[int], box_totals: dict[int, BoxTotals]) -> BoxTotals:
    result = BoxTotals()
    for box in boxes:
        result.add(box_totals[box])
    return result


HEADERS = [
    '范围', '消费记录数', '开盒次数', '开奖主记录数', '金币中奖记录数',
    '消费金币', '消费美元', '返奖金币', '金币返奖美元', '返奖钻石',
    '钻石返奖美元', '平台利润美元', '利润率',
]


def summary_values(label: str | int, totals: BoxTotals) -> list[Any]:
    spend = Decimal(totals.spend_tokens) / COIN_PER_USD
    coin = Decimal(totals.coin_reward_tokens) / COIN_PER_USD
    diamond = Decimal(totals.diamond_reward_tokens) / DIAMOND_PER_USD
    profit = spend - coin - diamond
    return [label, totals.spend_records, totals.spend_amount, totals.main_records,
            totals.coin_win_records, totals.spend_tokens, money(spend),
            totals.coin_reward_tokens, money(coin), totals.diamond_reward_tokens,
            money(diamond), money(profit), float(profit / spend) if spend else None]


def style_sheet(ws, header_row: int) -> None:
    ws.sheet_view.showGridLines = True
    ws.sheet_properties.tabColor = '24476B'
    ws.freeze_panes = f'A{header_row + 1}'
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            if cell.row == header_row:
                cell.fill = PatternFill('solid', fgColor='24476B')
                cell.font = Font(name='Aptos', size=10, bold=True, color='FFFFFF')
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif cell.font.name != 'Aptos':
                cell.font = Font(name='Aptos', size=10, color='243447')
                cell.alignment = Alignment(horizontal='right' if isinstance(cell.value, (int, float)) else 'left', vertical='center')
    ws.row_dimensions[header_row].height = 28
    for column in ws.columns:
        values = [str(cell.value) for cell in column if cell.value is not None]
        width = min(29, max(12, max((len(value) for value in values), default=10) + 2))
        ws.column_dimensions[get_column_letter(column[0].column)].width = width


def format_summary(ws, first: int, last: int) -> None:
    for cells in ws.iter_rows(min_row=first, max_row=last):
        for index in (6, 8, 10):
            cells[index - 1].number_format = TOKEN_FORMAT
        for index in (7, 9, 11, 12):
            cells[index - 1].number_format = MONEY_FORMAT
        cells[12].number_format = PERCENT_FORMAT
        profit = cells[11]
        profit.font = Font(name='Aptos', size=10, bold=True,
                           color='008A45' if profit.value > 0 else ('C62828' if profit.value < 0 else '243447'))


def write_overview(ws, start: str, end: str, box_totals: dict[int, BoxTotals]) -> None:
    ws['A2'] = '盲盒消费、开奖与平台利润'
    ws['A2'].font = Font(name='Aptos', size=15, bold=True, color='17324D')
    ws['A4'], ws['B4'] = '统计开始（含）', start
    ws['A5'], ws['B5'] = '统计结束（含）', end
    ws['D4'], ws['E4'] = '金币换算', '10,000 金币 = 1 美元'
    ws['D5'], ws['E5'] = '钻石换算', '10,000,000 钻石 = 1 美元'
    ws['A7'], ws['B7'] = '利润口径', '期间消费 - 期间金币返奖(16) - 期间钻石返奖(2–5/12–14)；不含 jackpot(17)'
    for column, header in enumerate(HEADERS, 1):
        ws.cell(9, column, header)
    groups = [('全部盲盒', set(BOX_GIFTS)),
              ('普通礼物盲盒', set(NORMAL_BOXES)),
              ('幸运礼物盲盒', set(LUCKY_BOXES))]
    for row_index, (label, boxes) in enumerate(groups, 10):
        for column, value in enumerate(summary_values(label, summarize(boxes, box_totals)), 1):
            ws.cell(row_index, column, value)
    style_sheet(ws, 9)
    format_summary(ws, 10, 12)
    ws.column_dimensions['A'].width = 23
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['L'].width = 20
    ws.auto_filter.ref = 'A9:M12'


def write_box_sheet(
    ws,
    title: str,
    box_ids: frozenset[int],
    box_totals: dict[int, BoxTotals],
    gift_totals: dict[tuple[int, int], GiftTotals],
    box_prices: dict[int, int],
) -> None:
    ws['A2'] = title
    ws['A2'].font = Font(name='Aptos', size=15, bold=True, color='17324D')
    ws['A4'] = '金币差额 = 实际消费金币 - 开盒次数 × 单价；非零需核查'
    ws['A4'].font = Font(name='Aptos', size=10, color='9C5700')
    for column, header in enumerate(HEADERS, 1):
        ws.cell(5, column, '盲盒ID' if column == 1 else header)
    for column, header in enumerate(('单价金币/次', '按单价应消费金币', '消费金币差额'), 14):
        ws.cell(5, column, header)
    for row_index, box in enumerate(sorted(box_ids), 6):
        for column, value in enumerate(summary_values(box, box_totals[box]), 1):
            ws.cell(row_index, column, value)
        expected = box_totals[box].spend_amount * box_prices[box]
        for column, value in enumerate((box_prices[box], expected,
                                        box_totals[box].spend_tokens - expected), 14):
            ws.cell(row_index, column, value)
    for column, value in enumerate(summary_values('合计', summarize(set(box_ids), box_totals)), 1):
        ws.cell(10, column, value)
    expected_total = sum(box_totals[box].spend_amount * box_prices[box] for box in box_ids)
    ws['O10'] = expected_total
    ws['P10'] = sum(box_totals[box].spend_tokens for box in box_ids) - expected_total
    ws['A13'] = '各盲盒礼物开奖与返奖'
    gift_headers = [
        '盲盒ID', '礼物ID', '开奖主记录数', '开奖占比', '主记录amount',
        '金币中奖记录数', '收礼人钻石', '钻石分账', '返奖钻石合计',
        '钻石返奖美元', '返奖金币', '金币返奖美元',
    ]
    for column, header in enumerate(gift_headers, 1):
        ws.cell(14, column, header)
    row_index = 15
    for box in sorted(box_ids):
        total_main = box_totals[box].main_records
        for gift in BOX_GIFTS[box]:
            item = gift_totals[box, gift]
            diamond_tokens = item.receiver_diamonds + item.share_diamonds
            values = [box, gift, item.main_records,
                      item.main_records / total_main if total_main else None,
                      item.main_amount, item.coin_win_records,
                      item.receiver_diamonds, item.share_diamonds, diamond_tokens,
                      money(Decimal(diamond_tokens) / DIAMOND_PER_USD),
                      item.coin_reward_tokens,
                      money(Decimal(item.coin_reward_tokens) / COIN_PER_USD)]
            for column, value in enumerate(values, 1):
                cell = ws.cell(row_index, column, value)
                if column in (7, 8, 9, 11):
                    cell.number_format = TOKEN_FORMAT
                elif column in (10, 12):
                    cell.number_format = MONEY_FORMAT
                elif column == 4:
                    cell.number_format = PERCENT_FORMAT
                elif column in (1, 2):
                    cell.number_format = ID_FORMAT
            row_index += 1
    style_sheet(ws, 5)
    format_summary(ws, 6, 10)
    for row in range(6, 11):
        for column in ('N', 'O', 'P'):
            ws[f'{column}{row}'].number_format = TOKEN_FORMAT
        difference = ws[f'P{row}']
        if difference.value != 0:
            difference.font = Font(name='Aptos', size=10, bold=True, color='C62828')
    for cell in ws[14][:len(gift_headers)]:
        cell.fill = PatternFill('solid', fgColor='426990')
        cell.font = Font(name='Aptos', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['L'].width = 20
    ws.column_dimensions['N'].width = 18
    ws.column_dimensions['O'].width = 22
    ws.column_dimensions['P'].width = 18
    ws.auto_filter.ref = 'A5:P10'


def write_top_sheet(ws, title: str, box_ids: frozenset[int], top_users: dict[int, list[tuple[int, int, int]]]) -> None:
    ws['A2'] = title
    ws['A2'].font = Font(name='Aptos', size=15, bold=True, color='17324D')
    ws['A3'] = 'ES 按用户消费金币聚合取前30；多分片结果可能存在排序误差'
    ws['A3'].font = Font(name='Aptos', size=10, color='9C5700')
    headers = ['盲盒ID', '名次', '用户ID', '消费金币', '消费美元', '消费记录数']
    for column, header in enumerate(headers, 1):
        ws.cell(4, column, header)
    row_index = 5
    for box in sorted(box_ids):
        for rank, (user, tokens, count) in enumerate(top_users[box], 1):
            values = [box, rank, user, tokens, money(Decimal(tokens) / COIN_PER_USD), count]
            for column, value in enumerate(values, 1):
                cell = ws.cell(row_index, column, value)
                cell.number_format = TOKEN_FORMAT if column == 4 else (MONEY_FORMAT if column == 5 else ID_FORMAT)
            row_index += 1
    style_sheet(ws, 4)
    ws.auto_filter.ref = f'A4:F{max(4, row_index - 1)}'
    ws.column_dimensions['C'].width = 18


def write_all_top_sheet(
    ws,
    users: list[tuple[int, int, int]],
    overall_by_box: dict[int, dict[int, int]],
) -> None:
    ws['A2'] = '全盲盒消费 Top 100'
    ws['A2'].font = Font(name='Aptos', size=15, bold=True, color='17324D')
    ws['A3'] = '普通为161–164，幸运为165–168；ES 多分片聚合排名可能存在误差'
    ws['A3'].font = Font(name='Aptos', size=10, color='9C5700')
    headers = ['名次', '用户ID', '消费金币', '消费美元', '消费记录数',
               '普通消费美元', '幸运消费美元'] + [f'{box}消费美元' for box in sorted(BOX_GIFTS)]
    for column, header in enumerate(headers, 1):
        ws.cell(4, column, header)
    for rank, (user, tokens, count) in enumerate(users, 1):
        by_box = overall_by_box[user]
        normal = sum(by_box.get(box, 0) for box in NORMAL_BOXES)
        lucky = sum(by_box.get(box, 0) for box in LUCKY_BOXES)
        values = [rank, user, tokens, money(Decimal(tokens) / COIN_PER_USD), count,
                  money(Decimal(normal) / COIN_PER_USD),
                  money(Decimal(lucky) / COIN_PER_USD)]
        values.extend(money(Decimal(by_box.get(box, 0)) / COIN_PER_USD)
                      for box in sorted(BOX_GIFTS))
        for column, value in enumerate(values, 1):
            cell = ws.cell(rank + 4, column, value)
            cell.number_format = (TOKEN_FORMAT if column == 3 else
                                  MONEY_FORMAT if column == 4 or column >= 6 else ID_FORMAT)
    style_sheet(ws, 4)
    ws.freeze_panes = 'C5'
    ws.auto_filter.ref = f'A4:O{max(4, len(users) + 4)}'
    ws.column_dimensions['B'].width = 18
    for column in range(6, 16):
        ws.column_dimensions[get_column_letter(column)].width = 18


def build_report(
    start: str,
    end: str,
    box_totals: dict[int, BoxTotals],
    gift_totals: dict[tuple[int, int], GiftTotals],
    top_users: dict[int, list[tuple[int, int, int]]],
    overall_top_users: list[tuple[int, int, int]],
    overall_by_box: dict[int, dict[int, int]],
    box_prices: dict[int, int],
) -> Workbook:
    workbook = Workbook()
    workbook.active.title = '分析总览'
    write_overview(workbook.active, start, end, box_totals)
    write_box_sheet(workbook.create_sheet('普通礼物盲盒'), '普通礼物盲盒',
                    NORMAL_BOXES, box_totals, gift_totals, box_prices)
    write_box_sheet(workbook.create_sheet('幸运礼物盲盒'), '幸运礼物盲盒',
                    LUCKY_BOXES, box_totals, gift_totals, box_prices)
    write_all_top_sheet(workbook.create_sheet('全盲盒消费Top100'),
                        overall_top_users, overall_by_box)
    write_top_sheet(workbook.create_sheet('普通消费Top30'), '普通盲盒各盒消费 Top 30',
                    NORMAL_BOXES, top_users)
    write_top_sheet(workbook.create_sheet('幸运消费Top30'), '幸运盲盒各盒消费 Top 30',
                    LUCKY_BOXES, top_users)
    return workbook


def write_report(args: argparse.Namespace) -> None:
    start = None if getattr(args, 'all_history', False) else args.start
    validate_period(start, args.end)
    box_totals, gift_totals = read_summaries(args.consumption, args.coin_rewards,
                                              args.diamond_rewards)
    box_prices = load_box_prices(getattr(args, 'box_prices', DEFAULT_BOX_PRICES))
    top_users, overall_top_users, overall_by_box = read_top_users(args.top_users, box_totals)
    workbook = build_report(start or '全部历史', args.end, box_totals, gift_totals,
                            top_users, overall_top_users, overall_by_box, box_prices)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f'已生成盲盒分析报表：{args.output}')


def main() -> None:
    parser = argparse.ArgumentParser(description='从 ES 聚合 JSON 生成盲盒利润分析 Excel')
    modes = parser.add_subparsers(dest='mode', required=True)
    queries = modes.add_parser('queries', help='生成四份 ES 聚合 DSL')
    query_scope = queries.add_mutually_exclusive_group(required=True)
    query_scope.add_argument('--start', help='统计开始时间，含')
    query_scope.add_argument('--all-history', action='store_true', help='从最早数据开始，不设开始时间')
    queries.add_argument('--end', required=True, help='统计结束时间，含')
    queries.add_argument('--output-dir', required=True, type=Path)
    queries.set_defaults(run=write_queries)
    report = modes.add_parser('report', help='读取四组 ES 聚合结果并输出 Excel')
    report.add_argument('--consumption', required=True, type=Path, help='消费汇总响应 JSON')
    report.add_argument('--coin-rewards', required=True, type=Path, help='金币返奖汇总响应 JSON')
    report.add_argument('--diamond-rewards', required=True, type=Path, help='钻石返奖汇总响应 JSON')
    report.add_argument('--top-users', required=True, type=Path, help='每盒 Top 30 和全盲盒 Top 100 聚合响应 JSON')
    report.add_argument('--box-prices', type=Path, default=DEFAULT_BOX_PRICES,
                        help='盲盒单价配置 JSON，默认使用工具目录下的 box-prices.json')
    report_scope = report.add_mutually_exclusive_group(required=True)
    report_scope.add_argument('--start')
    report_scope.add_argument('--all-history', action='store_true')
    report.add_argument('--end', required=True)
    report.add_argument('--output', required=True, type=Path)
    report.set_defaults(run=write_report)
    args = parser.parse_args()
    try:
        args.run(args)
    except (ValueError, OSError, json.JSONDecodeError) as exc:
        parser.exit(2, f'错误：{exc}\n')


if __name__ == '__main__':
    main()
