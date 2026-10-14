"""离线验证日期边界、聚合口径、续跑输出和 Excel 结构。"""

from __future__ import annotations

import json
import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from export_report import create_workbook, metrics, summarize
from fetch_data import (COIN_INDEX, DIAMOND_INDEX, ROOT, es_query, exact_nonnegative,
                        MySQLReader, export_es_job, fetch_mysql_snapshot, load_json, output_paths,
                        run_collection, save_json)


ANALYSIS = {
    "country_code": "BR",
    "register_start": "2026-09-01 00:00:00",
    "register_end": "2026-10-10 00:00:00",
    "observation_end": "2026-10-10 00:00:00",
    "es_time_zone": "+00:00",
    "coin_consumption_types": [1, 11, 21, 41, 42, 105, 106],
    "coin_reward_types": [62],
    "diamond_income_types": [2, 61],
    "diamond_reward_types": [61],
}


class FakeMySQL:
    def __init__(self):
        self.sqls = []

    def query(self, sql, parameters):
        self.sqls.append(sql)
        if "FROM sano_user u" in sql:
            yield {"user_id": 101, "nick": "=BAD()", "country_code": "BR",
                   "create_time": datetime(2026, 9, 1, 10), "real_person": 1,
                   "inviter_id": 300}
            yield {"user_id": 102, "nick": "后来的用户", "country_code": "BR",
                   "create_time": datetime(2026, 9, 20, 10), "real_person": 0,
                   "inviter_id": None}
            yield {"user_id": 103, "nick": "最新用户", "country_code": "BR",
                   "create_time": datetime(2026, 10, 9, 10), "real_person": 0,
                   "inviter_id": None}
        elif "FROM sano_anchor_real_error" in sql:
            yield {"user_id": 101, "first_success": datetime(2026, 9, 2, 8)}
        elif "FROM sano_user_active" in sql:
            yield {"user_id": 101, "day_int": 20260902}
            yield {"user_id": 101, "day_int": 20260904}
            yield {"user_id": 101, "day_int": 20260908}
            yield {"user_id": 102, "day_int": 20260921}
        elif "FROM sano_wallet_recharge" in sql:
            yield {"user_id": 101, "event_day": date(2026, 9, 1), "usd": Decimal("3.28")}
            yield {"user_id": 102, "event_day": date(2026, 9, 21), "usd": Decimal("9.99")}
        elif "FROM sano_wallet_withdraw" in sql:
            yield {"user_id": 101, "event_day": date(2026, 9, 25), "usd": Decimal("1.250")}


class FakeES:
    def search(self, index, body):
        assert body["size"] == 0
        user_id = body["query"]["bool"]["filter"][0]["terms"]["user_id"][0]
        expected_end = {101: "2026-10-01 00:00:00", 102: "2026-10-10 00:00:00",
                        103: "2026-10-10 00:00:00"}[user_id]
        assert body["query"]["bool"]["filter"][1]["range"]["create_time"]["lt"] == expected_end
        if user_id == 102:
            bucket = ({"key": {"userId": 102, "businessType": 11, "status": -1},
                       "tokens": {"value": 10000.0}, "doc_count": 1}
                      if index == COIN_INDEX else
                      {"key": {"userId": 102, "businessType": 2, "status": 1},
                       "tokens": {"value": 10000000.0}, "doc_count": 1})
            return {"_shards": {"failed": 0}, "timed_out": False,
                    "aggregations": {"groups": {"buckets": [bucket]}}}
        if user_id == 103:
            return {"_shards": {"failed": 0}, "timed_out": False,
                    "aggregations": {"groups": {"buckets": []}}}
        if index == COIN_INDEX:
            buckets = [
                {"key": {"userId": 101, "businessType": 1, "status": -1},
                 "tokens": {"value": 20000.0}, "doc_count": 2},
                {"key": {"userId": 101, "businessType": 62, "status": 1},
                 "tokens": {"value": 5000.0}, "doc_count": 1},
            ]
        elif index == DIAMOND_INDEX:
            buckets = [
                {"key": {"userId": 101, "businessType": 2, "status": 1},
                 "tokens": {"value": 10000000.0}, "doc_count": 1},
                {"key": {"userId": 101, "businessType": 61, "status": 1},
                 "tokens": {"value": 5000000.0}, "doc_count": 1},
            ]
        else:
            raise AssertionError(index)
        return {"_shards": {"failed": 0}, "timed_out": False,
                "aggregations": {"groups": {"buckets": buckets}}}


class ReportTests(unittest.TestCase):
    def test_manifest_replace_retries_transient_windows_lock(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            target = Path(directory) / "manifest.json"
            target.write_text('{"stage":"old"}', encoding="utf-8")
            original_replace = Path.replace
            attempts = 0

            def flaky_replace(source, destination):
                nonlocal attempts
                attempts += 1
                if attempts <= 2:
                    raise PermissionError(5, "模拟 Windows 短暂文件占用")
                return original_replace(source, destination)

            with patch.object(Path, "replace", flaky_replace), patch("fetch_data.time.sleep"):
                save_json(target, {"stage": "new"})
            self.assertEqual(attempts, 3)
            self.assertEqual(load_json(target)["stage"], "new")
            self.assertFalse((Path(directory) / ".manifest.json.part").exists())

    def test_mysql_sql_log_contains_expanded_select_only(self):
        class Cursor:
            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, traceback):
                return False

            def mogrify(self, sql, parameters):
                return sql.replace("%s", "'PH'", 1)

            def execute(self, sql, parameters):
                self.executed = (sql, parameters)

            def __iter__(self):
                return iter([{"n": 1}])

        class Connection:
            def cursor(self):
                return Cursor()

        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            log_path = Path(directory) / "queries.sql"
            reader = object.__new__(MySQLReader)
            reader.connection = Connection()
            reader.sql_log_path = log_path
            reader.query_count = 0
            self.assertEqual(list(reader.query("SELECT COUNT(*) AS n FROM sano_user WHERE country_code = %s",
                                               ("PH",))), [{"n": 1}])
            self.assertIn("country_code = 'PH'", log_path.read_text(encoding="utf-8"))
            self.assertIn("-- MySQL SELECT 1", log_path.read_text(encoding="utf-8"))
            with self.assertRaisesRegex(ValueError, "只允许执行单条 SELECT"):
                list(reader.query("DELETE FROM sano_user", ()))

    def test_mysql_maturity_and_select_only(self):
        fake = FakeMySQL()
        data = fetch_mysql_snapshot(fake, ANALYSIS)
        first, second, third = data["users"]
        self.assertEqual([first["d1"], first["d3"], first["d7"]], [True, True, True])
        self.assertEqual(first["rechargeUsd"], "3.28")
        self.assertEqual(first["withdrawUsd"], "1.250")
        self.assertTrue(first["full30Days"])
        self.assertEqual(second["d1"], True)
        self.assertFalse(second["full30Days"])
        self.assertEqual(second["rechargeUsd"], "9.99")
        self.assertEqual(second["withdrawUsd"], "0")
        self.assertIsNone(third["d1"])
        self.assertEqual(third["rechargeUsd"], "0")
        self.assertTrue(fake.sqls)
        self.assertTrue(any("FROM sano_user_active" in sql for sql in fake.sqls))
        self.assertFalse(any("sano_log_login" in sql for sql in fake.sqls))
        self.assertTrue(all(sql.lstrip().upper().startswith("SELECT ") and ";" not in sql
                            for sql in fake.sqls))

    def test_active_retention_crosses_month_and_excludes_unfinished_day(self):
        class MonthBoundaryMySQL:
            def __init__(self):
                self.active_parameters = None

            def query(self, sql, parameters):
                if "FROM sano_user u" in sql:
                    yield {"user_id": 201, "nick": "边界用户", "country_code": "BR",
                           "create_time": datetime(2026, 9, 30, 20), "real_person": 0,
                           "inviter_id": None}
                elif "FROM sano_user_active" in sql:
                    self.active_parameters = parameters
                    for active_day in (20261001, 20261003):
                        if parameters[-2] <= active_day <= parameters[-1]:
                            yield {"user_id": 201, "day_int": active_day}

        analysis = {**ANALYSIS, "observation_end": "2026-10-04 12:00:00"}
        fake = MonthBoundaryMySQL()
        user = fetch_mysql_snapshot(fake, analysis)["users"][0]
        self.assertEqual((user["d1"], user["d3"], user["d7"]), (True, True, None))
        self.assertEqual(fake.active_parameters[-2:], (20261001, 20261003))
        earlier = {**ANALYSIS, "observation_end": "2026-10-03 12:00:00"}
        earlier_user = fetch_mysql_snapshot(fake, earlier)["users"][0]
        self.assertEqual((earlier_user["d1"], earlier_user["d3"], earlier_user["d7"]),
                         (True, None, None))
        self.assertEqual(fake.active_parameters[-2:], (20261001, 20261002))

    def test_es_query_is_aggregated_and_read_only(self):
        body = es_query(ANALYSIS, COIN_INDEX, [101],
                        "2026-09-01 00:00:00", "2026-10-01 00:00:00")
        self.assertEqual(body["size"], 0)
        self.assertEqual(body["aggs"]["groups"]["composite"]["size"], 1000)
        self.assertEqual(body["query"]["bool"]["filter"][0]["terms"]["user_id"], [101])
        with self.assertRaises(ValueError):
            exact_nonnegative(2**53, "tokens")

    def test_legacy_snapshot_cannot_turn_unqueried_amounts_into_zero(self):
        with self.assertRaisesRegex(ValueError, "请用新版工具重新查询"):
            metrics({"rechargeUsd": None})

    def test_es_composite_pagination(self):
        class PagedES:
            def search(self, index, body):
                after = body["aggs"]["groups"]["composite"].get("after")
                business = 1 if after is None else 62
                result = {"_shards": {"failed": 0}, "timed_out": False,
                          "aggregations": {"groups": {"buckets": [
                              {"key": {"userId": 101, "businessType": business,
                                       "status": -1 if business == 1 else 1},
                               "tokens": {"value": 10000.0}, "doc_count": 1},
                          ]}}}
                if after is None:
                    result["aggregations"]["groups"]["after_key"] = {"userId": 101,
                                                                           "businessType": 1,
                                                                           "status": -1}
                return result

        job = {"name": "page", "index": COIN_INDEX, "ids": [101],
               "start": "2026-09-01 00:00:00", "end": "2026-10-01 00:00:00"}
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            destination = Path(directory) / "page.jsonl"
            self.assertEqual(export_es_job(PagedES(), ANALYSIS, job, destination), 2)
            self.assertEqual(len(destination.read_text(encoding="utf-8").splitlines()), 2)

    def test_end_to_end_mock_and_excel(self):
        config = {"mysql": {"host": "mock", "database": "sano", "port": 3306},
                  "es": {"es_url": "https://mock:9200"}, "analysis": ANALYSIS}
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            paths = output_paths("test", Path(directory))
            data_path = run_collection(config, paths, mysql_reader=FakeMySQL(), es_reader=FakeES())
            data = load_json(data_path)
            first, second, third = data["users"]
            self.assertEqual(first["consumptionCoinTokens"], 20000)
            self.assertEqual(first["incomeDiamondTokens"], 15000000)
            self.assertEqual(first["rewardCoinTokens"], 5000)
            self.assertEqual(first["rewardDiamondTokens"], 5000000)
            self.assertEqual(second["consumptionCoinTokens"], 10000)
            self.assertEqual(second["incomeDiamondTokens"], 10000000)
            self.assertEqual(third["incomeDiamondTokens"], 0)
            self.assertEqual(metrics(first)["reward"], Decimal("1.000"))
            self.assertEqual(metrics(first)["income"], Decimal("1.500"))
            total, daily = summarize(data["users"])
            self.assertEqual(total["users"], 3)
            self.assertEqual(total["mature30"], 1)
            self.assertEqual(total["d1_eligible"], 2)
            self.assertEqual(total["d1_retained"], 2)
            self.assertEqual(total["reward_usd"], Decimal("1.000"))
            self.assertEqual(total["recharge_usd"], Decimal("13.270"))
            self.assertEqual(total["recharge_users"], 2)
            self.assertEqual(total["consumption_usd"], Decimal("3.000"))
            self.assertEqual(total["income_usd"], Decimal("2.500"))
            self.assertEqual(len(daily), 3)
            self.assertEqual(run_collection(config, paths, resume=True), data_path)
            create_workbook(data, paths["report"])
            book = load_workbook(paths["report"], read_only=False, data_only=False)
            try:
                self.assertEqual(book.sheetnames, ["分析汇总", "注册日汇总", "用户明细"])
                detail = book["用户明细"]
                self.assertEqual(detail.max_row, 4)
                self.assertEqual(detail["A2"].value, 101)
                self.assertEqual(detail["B2"].data_type, "s")
                self.assertEqual(detail["B2"].value, "=BAD()")
                self.assertEqual(detail["M2"].value, 1.5)
                self.assertEqual(detail["O2"].value, 1)
                self.assertEqual(detail["K3"].value, 9.99)
                self.assertEqual(detail["K4"].value, 0)
                self.assertEqual(detail["L3"].value, 1)
                self.assertEqual(detail["M3"].value, 1)
                self.assertEqual(detail["H4"].value, "未到期")
                self.assertEqual(detail.freeze_panes, "E2")
                self.assertEqual(detail.auto_filter.ref, "A1:O4")
                self.assertTrue(detail.sheet_view.showGridLines)
                self.assertEqual(detail["A2"].border.left.style, "thin")
                self.assertTrue(book["分析汇总"].sheet_view.showGridLines)
                self.assertEqual(book["分析汇总"]["A3"].border.left.style, "thin")
                summary_values = {row[0].value: row[1].value for row in book["分析汇总"].iter_rows(
                    min_row=14, max_col=2) if row[0].value}
                self.assertIn("sano_user_active", book["分析汇总"]["B8"].value)
                self.assertEqual(summary_values["完整观察30天人数"], 1)
                self.assertEqual(summary_values["30天充值总金额($)"], 13.27)
                self.assertEqual(summary_values["30天消费总金额($)"], 3)
                self.assertTrue(book["注册日汇总"].sheet_view.showGridLines)
                self.assertEqual(book["分析汇总"]["B3"].value, "BR")
                self.assertEqual(book["注册日汇总"]["B2"].value, 1)
                self.assertEqual(book["注册日汇总"]["N3"].value, 9.99)
                self.assertEqual(book["注册日汇总"]["N4"].value, 0)
                for sheet in book:
                    for row in sheet:
                        for cell in row:
                            self.assertNotEqual(cell.data_type, "f")
                            self.assertNotIn(cell.value, ("#REF!", "#DIV/0!", "#VALUE!"))
            finally:
                book.close()
            manifest = load_json(paths["manifest"])
            self.assertEqual(manifest["stage"], "json_complete")
            self.assertEqual(len(manifest["completed_jobs"]), 6)
            self.assertNotIn("password", json.dumps(manifest))

    def test_empty_cohort(self):
        class EmptyMySQL:
            def query(self, sql, parameters):
                return iter(())

        class NoES:
            def search(self, index, body):
                raise AssertionError("空用户群不应发起 ES 查询")

        config = {"mysql": {"host": "mock", "database": "sano", "port": 3306},
                  "es": {"es_url": "https://mock:9200"}, "analysis": ANALYSIS}
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            paths = output_paths("empty", Path(directory))
            data_path = run_collection(config, paths, mysql_reader=EmptyMySQL(), es_reader=NoES())
            create_workbook(load_json(data_path), paths["report"])
            book = load_workbook(paths["report"], read_only=True)
            try:
                self.assertEqual(book["用户明细"].max_row, 1)
                self.assertEqual(book["分析汇总"]["B14"].value, 0)
            finally:
                book.close()

    def test_resume_after_es_interruption(self):
        class InterruptOnce(FakeES):
            def __init__(self):
                self.calls = []

            def search(self, index, body):
                self.calls.append(index)
                if index == DIAMOND_INDEX:
                    raise RuntimeError("模拟 ES 中断")
                return super().search(index, body)

        class ResumeES(FakeES):
            def __init__(self):
                self.calls = []

            def search(self, index, body):
                self.calls.append(index)
                return super().search(index, body)

        config = {"mysql": {"host": "mock", "database": "sano", "port": 3306},
                  "es": {"es_url": "https://mock:9200"}, "analysis": ANALYSIS}
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            paths = output_paths("interrupted", Path(directory))
            with self.assertRaisesRegex(RuntimeError, "模拟 ES 中断"):
                run_collection(config, paths, mysql_reader=FakeMySQL(), es_reader=InterruptOnce())
            manifest = load_json(paths["manifest"])
            self.assertEqual(len(manifest["completed_jobs"]), 1)
            resumed = ResumeES()
            data_path = run_collection(config, paths, resume=True,
                                       mysql_reader=FakeMySQL(), es_reader=resumed)
            self.assertEqual(resumed.calls, [DIAMOND_INDEX, COIN_INDEX, DIAMOND_INDEX,
                                             COIN_INDEX, DIAMOND_INDEX])
            self.assertTrue(data_path.is_file())


if __name__ == "__main__":
    unittest.main()
