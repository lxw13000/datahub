"""国别注册用户留存与注册后 30 天金额分析；只执行只读查询。"""

from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fetch_data import (ROOT, DATE_TIME_FORMAT, find_resume_paths, identity, load_configs,
                        load_json, mark_report_complete, output_paths, parse_time, run_collection)


USD_COIN = Decimal(10_000)
USD_DIAMOND = Decimal(10_000_000)
THREE_PLACES = Decimal("0.001")
NAVY = "17324D"
BLUE = "DCEAF5"
GREEN = "006C3B"
RED = "B42318"
LINE = Side(style="thin", color="B7C7D5")
BODY_BORDER = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
MONEY_FORMAT = '#,##0.000;[Red](#,##0.000)'
COUNT_FORMAT = '#,##0'
ID_FORMAT = '0'
PERCENT_FORMAT = '0.0%'

DETAIL_HEADERS = [
    "用户ID", "昵称", "国别", "邀请人", "注册时间", "实名状态", "实名时间",
    "D1是否留存", "D3是否留存", "D7是否留存", "30天充值金额($)",
    "30天消费金额($)", "30天收益金额($)", "30天提现成功金额($)",
    "30天奖励领取金额($)",
]
DAILY_HEADERS = [
    "注册日", "注册人数", "当前实名人数", "D1到期人数", "D1留存人数", "D1留存率",
    "D3到期人数", "D3留存人数", "D3留存率", "D7到期人数", "D7留存人数",
    "D7留存率", "30天到期人数", "30天充值($)", "30天消费($)",
    "30天收益($)", "30天成功提现($)", "30天奖励领取($)",
]


def usd_coin(tokens: int) -> Decimal:
    return (Decimal(tokens) / USD_COIN).quantize(THREE_PLACES, rounding=ROUND_HALF_UP)


def usd_diamond(tokens: int) -> Decimal:
    return (Decimal(tokens) / USD_DIAMOND).quantize(THREE_PLACES, rounding=ROUND_HALF_UP)


def money(value: str) -> Decimal:
    return Decimal(value).quantize(THREE_PLACES, rounding=ROUND_HALF_UP)


def metrics(user: dict) -> dict[str, Decimal]:
    required = ("rechargeUsd", "withdrawUsd", "consumptionCoinTokens",
                "incomeDiamondTokens", "rewardCoinTokens", "rewardDiamondTokens", "full30Days")
    if any(key not in user or user[key] is None for key in required):
        raise ValueError("聚合 JSON 缺少未满30天用户的金额数据；请用新版工具重新查询，不能把旧数据的未到期当作零")
    return {
        "recharge": money(user["rechargeUsd"]),
        "consumption": usd_coin(user["consumptionCoinTokens"]),
        "income": usd_diamond(user["incomeDiamondTokens"]),
        "withdraw": money(user["withdrawUsd"]),
        "reward": usd_coin(user["rewardCoinTokens"]) + usd_diamond(user["rewardDiamondTokens"]),
    }


def retention(value: bool | None) -> str:
    return "未到期" if value is None else ("是" if value else "否")


def set_text(cell, value: str) -> None:
    """数据库昵称是非可信文本，防止被 Excel 解释为公式。"""
    cell.value = value
    cell.data_type = "s"


def style_title(sheet, text: str) -> None:
    sheet.sheet_view.showGridLines = True
    sheet["A1"] = text
    sheet["A1"].font = Font(name="Arial", size=15, bold=True, color=NAVY)
    sheet.row_dimensions[1].height = 25
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def style_header(sheet, row: int, columns: int) -> None:
    for cell in sheet[row][:columns]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BODY_BORDER
    sheet.row_dimensions[row].height = 28


def style_body(sheet, first_row: int, last_row: int, columns: int) -> None:
    if last_row < first_row:
        return
    for row in sheet.iter_rows(min_row=first_row, max_row=last_row, max_col=columns):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="263746")
            cell.border = BODY_BORDER
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[row[0].row].height = 20


def summarize(users: list[dict]) -> tuple[dict, dict[str, dict]]:
    overall = {"users": len(users), "real": 0, "mature30": 0}
    daily: dict[str, dict] = defaultdict(lambda: {"users": 0, "real": 0, "mature30": 0})
    for source in [overall]:
        for n in (1, 3, 7):
            source[f"d{n}_eligible"] = 0
            source[f"d{n}_retained"] = 0
        for key in ("recharge", "consumption", "income", "withdraw", "reward"):
            source[f"{key}_users"] = 0
            source[f"{key}_usd"] = Decimal(0)
    for user in users:
        day = user["registerTime"][:10]
        group = daily[day]
        if "d1_eligible" not in group:
            for n in (1, 3, 7):
                group[f"d{n}_eligible"] = 0
                group[f"d{n}_retained"] = 0
            for key in ("recharge", "consumption", "income", "withdraw", "reward"):
                group[f"{key}_users"] = 0
                group[f"{key}_usd"] = Decimal(0)
        group["users"] += 1
        for target in (overall, group):
            target["real"] += int(user["realNameStatus"])
            for n in (1, 3, 7):
                value = user[f"d{n}"]
                if value is not None:
                    target[f"d{n}_eligible"] += 1
                    target[f"d{n}_retained"] += int(value)
            target["mature30"] += int(user["full30Days"])
            user_metrics = metrics(user)
            for key, value in user_metrics.items():
                target[f"{key}_usd"] += value
                target[f"{key}_users"] += int(value > 0)
    return overall, dict(sorted(daily.items()))


def add_summary(workbook: Workbook, data: dict, total: dict) -> None:
    sheet = workbook.active
    sheet.title = "分析汇总"
    sheet.sheet_properties.tabColor = NAVY
    style_title(sheet, "注册用户留存与30天金额分析")
    meta = data["meta"]
    conditions = [
        ("注册国家代码", meta["country_code"]),
        ("注册起点（含）", meta["register_start"]),
        ("注册终点（不含）", meta["register_end"]),
        ("观察截止（不含）", meta["observation_end"]),
        ("ES 时间区", meta["es_time_zone"]),
        ("留存数据源", "sano_user_active：对应自然日的 day_int 有记录即留存"),
        ("消费业务类型", ", ".join(map(str, meta["coin_consumption_types"]))),
        ("金额统计窗口", "各用户 D0–D29，最多统计至观察截止；未满30天为阶段金额"),
        ("奖励说明", "钻石奖励已包含在30天收益中；收益和奖励不可相加"),
    ]
    for row_no, (label, value) in enumerate(conditions, start=3):
        sheet.cell(row_no, 1, label)
        sheet.cell(row_no, 2, value)
        sheet.cell(row_no, 1).font = Font(name="Arial", size=10, bold=True, color=NAVY)
        sheet.cell(row_no, 2).font = Font(name="Arial", size=10, color="263746")
        sheet.cell(row_no, 1).border = BODY_BORDER
        sheet.cell(row_no, 2).border = BODY_BORDER
    header_row = 4 + len(conditions)
    sheet.cell(header_row, 1, "指标")
    sheet.cell(header_row, 2, "数值")
    sheet.cell(header_row, 3, "说明")
    style_header(sheet, header_row, 3)
    lines = [
        ("注册用户数", total["users"], "指定国家及注册区间"),
        ("当前实名人数", total["real"], "sano_user.real_person = 1"),
        ("当前实名率", total["real"] / total["users"] if total["users"] else None, "分母：注册用户数"),
    ]
    for n in (1, 3, 7):
        denominator, numerator = total[f"d{n}_eligible"], total[f"d{n}_retained"]
        lines.extend([
            (f"D{n} 已到期人数", denominator, "未到期用户不纳入分母"),
            (f"D{n} 留存人数", numerator, "活跃表中目标日有记录"),
            (f"D{n} 留存率", numerator / denominator if denominator else None, "留存人数 ÷ 已到期人数"),
        ])
    lines.append(("完整观察30天人数", total["mature30"], "D0 至 D29 均已结束；不限制下方金额统计"))
    names = {"recharge": "充值", "consumption": "消费", "income": "收益",
             "withdraw": "成功提现", "reward": "奖励领取"}
    for key, label in names.items():
        lines.append((f"30天{label}用户数", total[f"{key}_users"],
                      "已发生金额大于 0 的用户，含未满30天用户"))
        lines.append((f"30天{label}总金额($)", total[f"{key}_usd"],
                      "D0–D29 内已发生金额，截至观察截止"))
    for row_no, (label, value, note) in enumerate(lines, start=header_row + 1):
        sheet.cell(row_no, 1, label)
        sheet.cell(row_no, 2, value if value is not None else "未到期")
        sheet.cell(row_no, 3, note)
        if isinstance(value, Decimal):
            sheet.cell(row_no, 2).number_format = MONEY_FORMAT
        elif isinstance(value, float):
            sheet.cell(row_no, 2).number_format = PERCENT_FORMAT
        elif isinstance(value, int):
            sheet.cell(row_no, 2).number_format = COUNT_FORMAT
    style_body(sheet, header_row + 1, header_row + len(lines), 3)
    sheet.column_dimensions["A"].width = 29
    sheet.column_dimensions["B"].width = 66
    sheet.column_dimensions["C"].width = 44


def add_daily(workbook: Workbook, daily: dict[str, dict]) -> None:
    sheet = workbook.create_sheet("注册日汇总")
    sheet.sheet_view.showGridLines = True
    sheet.append(DAILY_HEADERS)
    style_header(sheet, 1, len(DAILY_HEADERS))
    for day, values in daily.items():
        row = [day, values["users"], values["real"]]
        for n in (1, 3, 7):
            denominator = values[f"d{n}_eligible"]
            numerator = values[f"d{n}_retained"]
            row.extend([denominator, numerator, numerator / denominator if denominator else None])
        row.append(values["mature30"])
        row.extend(values[f"{name}_usd"] for name in
                   ("recharge", "consumption", "income", "withdraw", "reward"))
        sheet.append(row)
    style_body(sheet, 2, sheet.max_row, len(DAILY_HEADERS))
    for row in sheet.iter_rows(min_row=2):
        for column in (6, 9, 12):
            row[column - 1].number_format = PERCENT_FORMAT
            if row[column - 1].value is None:
                row[column - 1].value = "未到期"
        for column in range(14, 19):
            row[column - 1].number_format = MONEY_FORMAT
        for column in (2, 3, 4, 5, 7, 8, 10, 11, 13):
            row[column - 1].number_format = COUNT_FORMAT
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:R{max(2, sheet.max_row)}"
    sheet.column_dimensions["A"].width = 16
    for column in range(2, 19):
        sheet.column_dimensions[get_column_letter(column)].width = 20 if column >= 14 else 16


def add_detail(workbook: Workbook, users: list[dict]) -> None:
    sheet = workbook.create_sheet("用户明细")
    sheet.sheet_view.showGridLines = True
    sheet.append(DETAIL_HEADERS)
    style_header(sheet, 1, len(DETAIL_HEADERS))
    for user in users:
        values = metrics(user)
        row_no = sheet.max_row + 1
        sheet.append([
            user["userId"], "", user["countryCode"], user["inviterId"],
            parse_time(user["registerTime"], "registerTime"),
            "已实名" if user["realNameStatus"] else "未实名",
            parse_time(user["firstRealNameTime"], "firstRealNameTime") if user["firstRealNameTime"] else None,
            retention(user["d1"]), retention(user["d3"]), retention(user["d7"]),
            *(values[key] for key in
              ("recharge", "consumption", "income", "withdraw", "reward")),
        ])
        set_text(sheet.cell(row_no, 2), user["nick"])
    style_body(sheet, 2, sheet.max_row, len(DETAIL_HEADERS))
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = ID_FORMAT
        row[3].number_format = ID_FORMAT
        row[4].number_format = "yyyy-mm-dd hh:mm:ss"
        row[6].number_format = "yyyy-mm-dd hh:mm:ss"
        for cell in row[10:15]:
            cell.number_format = MONEY_FORMAT
        for cell in row[7:10]:
            if cell.value == "是":
                cell.font = Font(name="Arial", size=10, color=GREEN)
            elif cell.value == "否":
                cell.font = Font(name="Arial", size=10, color=RED)
    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = f"A1:O{max(2, sheet.max_row)}"
    widths = [14, 24, 12, 14, 22, 14, 22, 15, 15, 15, 20, 20, 20, 24, 24]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def create_workbook(data: dict, output: Path) -> None:
    users = data.get("users")
    if not isinstance(users, list):
        raise ValueError("聚合 JSON 缺少用户列表")
    total, daily = summarize(users)
    workbook = Workbook()
    add_summary(workbook, data, total)
    add_daily(workbook, daily)
    add_detail(workbook, users)
    temporary = output.with_name(f".{output.name}.part.xlsx")
    try:
        workbook.save(temporary)
        # 重开文件检查工作表和行数，避免把不完整文件标记为成功。
        verified = load_workbook(temporary, read_only=True, data_only=False)
        try:
            if verified.sheetnames != ["分析汇总", "注册日汇总", "用户明细"]:
                raise ValueError("工作表缺失或顺序错误")
            if verified["用户明细"].max_row != len(users) + 1:
                raise ValueError("用户明细行数不匹配")
        finally:
            verified.close()
        temporary.replace(output)
    finally:
        if temporary.exists():
            temporary.unlink()


def main() -> None:
    parser = argparse.ArgumentParser(description="查询指定国家注册用户并输出留存与30天金额报表")
    parser.add_argument("--mysql-config", default=str(ROOT / "mysql-config.json"))
    parser.add_argument("--es-config", default=str(ROOT / "es-config.json"))
    parser.add_argument("--analysis-config", default=str(ROOT / "analysis-config.json"))
    parser.add_argument("--resume", action="store_true", help="续跑相同配置的未完成批次")
    args = parser.parse_args()
    config = load_configs(Path(args.mysql_config), Path(args.es_config), Path(args.analysis_config))
    if args.resume:
        paths = find_resume_paths(config)
    else:
        tag = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        paths = output_paths(tag)
    data_path = run_collection(config, paths, resume=args.resume)
    create_workbook(load_json(data_path), paths["report"])
    mark_report_complete(paths)
    print(f"已生成分析报表：{paths['report'].resolve()}")


if __name__ == "__main__":
    main()
