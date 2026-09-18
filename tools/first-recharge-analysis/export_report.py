#!/usr/bin/env python3
"""读取一个或多个首充查询 Excel，生成首充分析报告。"""

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DATE_FORMAT = "%Y-%m-%d"
FONT_NAME = "Arial"

TITLE_FONT = Font(name=FONT_NAME, size=16, bold=True, color="1F1F1F")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, italic=True, color="666666")
BODY_FONT = Font(name=FONT_NAME, size=10, color="333333")
LABEL_FONT = Font(name=FONT_NAME, size=10, bold=True, color="333333")
WHITE_BOLD_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
TOTAL_FILL = PatternFill("solid", fgColor="E2F0D9")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
ERROR_FILL = PatternFill("solid", fgColor="FCE4D6")
LIGHT_BLUE_FILL = PatternFill("solid", fgColor="EAF3F8")
THIN_BLUE = Side(style="thin", color="D6E4F0")
SECTION_BORDER = Border(bottom=Side(style="thin", color="9EADBA"))
BOTTOM_BORDER = Border(bottom=THIN_BLUE)

INTEGER_FORMAT = "#,##0"
USD_FORMAT = '"$"#,##0.00'
PERCENT_FORMAT = "0.00%"
HOURS_FORMAT = "#,##0.000"
DATE_NUMBER_FORMAT = "yyyy-mm-dd"
DATETIME_NUMBER_FORMAT = "yyyy-mm-dd hh:mm:ss"
TEXT_FORMAT = "@"

CHANNEL_LABELS = {"google": "Google", "apple": "Apple"}

HEADER_ALIASES = {
    "user_id": ("user_id", "userid", "用户id"),
    "nick": ("nick", "nickname", "昵称"),
    "channel_code": ("channel_code", "channel", "渠道"),
    "price_id": ("price_id", "priceid", "档位id", "商品id"),
    "tier_coin": ("tier_coin", "config_coin", "档位金币", "档位coin"),
    "standard_amount_usd": (
        "standard_amount_usd",
        "standardamountusd",
        "standard_price_usd",
        "price",
        "标准金额",
    ),
    "actual_money": ("actual_money", "money", "实际金额"),
    "actual_coin": ("actual_coin", "coin", "实际金币"),
    "recharge_time": ("recharge_time", "rechargetime", "充值时间", "create_time"),
    "register_time": ("register_time", "registertime", "注册时间"),
    "register_to_recharge_seconds": (
        "register_to_recharge_seconds",
        "registertorechargeseconds",
        "注册至首充秒数",
    ),
    "register_to_recharge_hours": (
        "register_to_recharge_hours",
        "registertorechargehours",
        "注册至首充小时数",
    ),
    "country_code": ("country_code", "countrycode", "国家代码"),
    "country_name": ("country_name", "countryname", "国家名称"),
    "real_person": ("real_person", "realperson", "真人状态"),
    "proxy_id": ("proxy_id", "proxyid", "代理id"),
    "first_proxy_id": ("first_proxy_id", "firstproxyid", "首个代理id"),
    "first_join_time": ("first_join_time", "firstjointime", "首次加入时间"),
    "successful_first_recharge_count": (
        "successful_first_recharge_count",
        "successfulfirstrechargecount",
        "成功首充记录数",
    ),
    "recharge_sequence": ("recharge_sequence", "rechargesequence", "首充序号"),
}

REQUIRED_COLUMNS = {
    "user_id",
    "channel_code",
    "price_id",
    "tier_coin",
    "standard_amount_usd",
    "recharge_time",
}

LATENCY_BUCKETS = (
    "00-10分钟",
    "10-30分钟",
    "30-60分钟",
    "01-06小时",
    "06-24小时",
    "01-03天",
    "03-07天",
    "07天以上",
    "注册时间缺失",
    "时间异常",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="将一个或多个首充查询 Excel 汇总为首充分析报告。"
    )
    parser.add_argument(
        "--input",
        type=Path,
        nargs="+",
        required=True,
        help="一个或多个查询结果 .xlsx 文件",
    )
    parser.add_argument("--output", type=Path, required=True, help="输出 .xlsx 文件")
    parser.add_argument(
        "--start-date",
        type=parse_cli_date,
        help="可选，统计开始日期，格式 yyyy-MM-dd",
    )
    parser.add_argument(
        "--end-date",
        type=parse_cli_date,
        help="可选，统计结束日期，格式 yyyy-MM-dd",
    )
    args = parser.parse_args()
    if args.start_date and args.end_date and args.end_date < args.start_date:
        parser.error("--end-date 不能早于 --start-date")
    return args


def parse_cli_date(value: str) -> date:
    try:
        return datetime.strptime(value, DATE_FORMAT).date()
    except ValueError as exc:
        raise argparse.ArgumentTypeError(f"日期必须是 {DATE_FORMAT}: {value}") from exc


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"[\s_\-（）()]+", "", str(value)).lower()


def alias_lookup() -> dict[str, str]:
    result: dict[str, str] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        result[normalize_header(canonical)] = canonical
        for alias in aliases:
            result[normalize_header(alias)] = canonical
    return result


ALIASES_TO_CANONICAL = alias_lookup()


def find_header_row(ws) -> tuple[int, dict[str, int]] | None:
    """在前 20 行中寻找包含所有必填列的表头。"""
    for row_number in range(1, min(ws.max_row, 20) + 1):
        values = next(
            ws.iter_rows(
                min_row=row_number,
                max_row=row_number,
                values_only=True,
            )
        )
        mapping: dict[str, int] = {}
        for column_index, value in enumerate(values):
            canonical = ALIASES_TO_CANONICAL.get(normalize_header(value))
            if canonical and canonical not in mapping:
                mapping[canonical] = column_index
        if REQUIRED_COLUMNS.issubset(mapping):
            return row_number, mapping
    return None


def clean_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_identifier(value: Any, field_name: str) -> tuple[str, bool]:
    """返回文本标识符以及是否存在 Excel 数字精度风险。"""
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"{field_name} 为空")
    precision_warning = False
    if isinstance(value, bool):
        raise ValueError(f"{field_name} 不能是布尔值")
    if isinstance(value, int):
        return str(value), len(str(abs(value))) > 15
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            raise ValueError(f"{field_name} 不是有限数字")
        if not value.is_integer():
            raise ValueError(f"{field_name} 必须是整数或文本")
        text = format(value, ".0f")
        precision_warning = len(text.lstrip("-")) > 15
        return text, precision_warning
    text = str(value).strip()
    if text.endswith(".0") and re.fullmatch(r"[+-]?\d+\.0", text):
        text = text[:-2]
    return text, False


def parse_decimal(value: Any, field_name: str, required: bool = False) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise ValueError(f"{field_name} 为空")
        return None
    if isinstance(value, bool):
        raise ValueError(f"{field_name} 不能是布尔值")
    text = str(value).strip().replace(",", "").replace("$", "")
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{field_name} 不是有效数字: {value}") from exc
    if not number.is_finite():
        raise ValueError(f"{field_name} 不是有限数字")
    return number


def parse_integer(value: Any, field_name: str, required: bool = False) -> int | None:
    number = parse_decimal(value, field_name, required=required)
    if number is None:
        return None
    if number != number.to_integral_value():
        raise ValueError(f"{field_name} 必须是整数: {value}")
    return int(number)


def parse_datetime_value(value: Any, field_name: str, required: bool = False) -> datetime | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise ValueError(f"{field_name} 为空")
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value).strip()
    for pattern in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y/%m/%d",
    ):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).replace(tzinfo=None)
    except ValueError as exc:
        raise ValueError(f"{field_name} 不是有效日期时间: {value}") from exc


def display_channel(value: Any) -> tuple[str, str]:
    raw = clean_text(value).lower()
    if not raw:
        raise ValueError("channel_code 为空")
    return raw, CHANNEL_LABELS.get(raw, clean_text(value))


def add_anomaly(
    anomalies: list[dict[str, Any]],
    anomaly_type: str,
    detail: str,
    record: dict[str, Any] | None,
    included: bool,
) -> None:
    record = record or {}
    anomalies.append(
        {
            "type": anomaly_type,
            "detail": detail,
            "user_id": record.get("user_id", ""),
            "recharge_time": record.get("recharge_time"),
            "channel": record.get("channel", ""),
            "price_id": record.get("price_id", ""),
            "standard_amount_usd": record.get("standard_amount_usd"),
            "country_code": record.get("country_code", ""),
            "country_name": record.get("country_name", ""),
            "source_file": record.get("source_file", ""),
            "source_sheet": record.get("source_sheet", ""),
            "source_row": record.get("source_row", ""),
            "included": "是" if included else "否",
        }
    )


def raw_record_for_anomaly(
    raw: dict[str, Any], source_file: str, source_sheet: str, source_row: int
) -> dict[str, Any]:
    return {
        "user_id": clean_text(raw.get("user_id")),
        "recharge_time": raw.get("recharge_time"),
        "channel": clean_text(raw.get("channel_code")),
        "price_id": clean_text(raw.get("price_id")),
        "standard_amount_usd": raw.get("standard_amount_usd"),
        "country_code": clean_text(raw.get("country_code")),
        "country_name": clean_text(raw.get("country_name")),
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": source_row,
    }


def parse_data_row(
    raw: dict[str, Any],
    source_file: str,
    source_sheet: str,
    source_row: int,
    anomalies: list[dict[str, Any]],
) -> dict[str, Any] | None:
    source_record = raw_record_for_anomaly(raw, source_file, source_sheet, source_row)
    try:
        user_id, user_precision_warning = parse_identifier(raw.get("user_id"), "user_id")
        price_id, price_precision_warning = parse_identifier(raw.get("price_id"), "price_id")
        channel_code, channel = display_channel(raw.get("channel_code"))
        tier_coin = parse_integer(raw.get("tier_coin"), "tier_coin", required=True)
        standard_amount = parse_decimal(
            raw.get("standard_amount_usd"),
            "standard_amount_usd",
            required=True,
        )
        recharge_time = parse_datetime_value(
            raw.get("recharge_time"), "recharge_time", required=True
        )
        if standard_amount is not None and standard_amount < 0:
            raise ValueError("standard_amount_usd 不能为负数")
        if tier_coin is not None and tier_coin < 0:
            raise ValueError("tier_coin 不能为负数")
    except ValueError as exc:
        add_anomaly(anomalies, "必填字段错误", str(exc), source_record, included=False)
        return None

    def optional_value(parser, value: Any, field_name: str) -> Any:
        try:
            return parser(value, field_name)
        except ValueError as exc:
            add_anomaly(
                anomalies,
                "可选字段错误",
                str(exc),
                source_record,
                included=True,
            )
            return None

    register_time = optional_value(
        parse_datetime_value, raw.get("register_time"), "register_time"
    )
    actual_money = optional_value(parse_decimal, raw.get("actual_money"), "actual_money")
    actual_coin = optional_value(parse_integer, raw.get("actual_coin"), "actual_coin")
    successful_count = optional_value(
        parse_integer,
        raw.get("successful_first_recharge_count"),
        "successful_first_recharge_count",
    )
    source_sequence = optional_value(
        parse_integer, raw.get("recharge_sequence"), "recharge_sequence"
    )
    first_join_time = optional_value(
        parse_datetime_value, raw.get("first_join_time"), "first_join_time"
    )

    country_code = clean_text(raw.get("country_code")).upper()
    country_name = clean_text(raw.get("country_name"))
    if not country_code:
        country_code = "UNKNOWN"
    if not country_name:
        country_name = "未知国家"

    latency_seconds = None
    if register_time is not None:
        latency_seconds = int((recharge_time - register_time).total_seconds())

    record = {
        "user_id": user_id,
        "nick": clean_text(raw.get("nick")),
        "channel_code": channel_code,
        "channel": channel,
        "price_id": price_id,
        "tier_coin": tier_coin,
        "standard_amount_usd": standard_amount,
        "actual_money": actual_money,
        "actual_coin": actual_coin,
        "recharge_time": recharge_time,
        "register_time": register_time,
        "latency_seconds": latency_seconds,
        "country_code": country_code,
        "country_name": country_name,
        "real_person": raw.get("real_person"),
        "proxy_id": clean_text(raw.get("proxy_id")),
        "first_proxy_id": clean_text(raw.get("first_proxy_id")),
        "first_join_time": first_join_time,
        "successful_count": successful_count,
        "source_sequence": source_sequence,
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": source_row,
    }

    if user_precision_warning:
        add_anomaly(
            anomalies,
            "标识符精度风险",
            "user_id 在 Excel 中以超过 15 位的数字保存，可能已经丢失精度",
            record,
            included=True,
        )
    if price_precision_warning:
        add_anomaly(
            anomalies,
            "标识符精度风险",
            "price_id 在 Excel 中以超过 15 位的数字保存，建议 SQL 使用 CAST(... AS CHAR)",
            record,
            included=True,
        )
    if channel_code not in CHANNEL_LABELS:
        add_anomaly(
            anomalies,
            "未知渠道",
            f"未识别的 channel_code: {channel_code}",
            record,
            included=True,
        )
    if country_code == "UNKNOWN" or country_name == "未知国家":
        add_anomaly(
            anomalies,
            "国家缺失",
            "国家代码或国家名称缺失，已归入未知国家",
            record,
            included=True,
        )
    if register_time is None:
        add_anomaly(
            anomalies,
            "注册时间缺失",
            "无法计算注册至首充耗时",
            record,
            included=True,
        )
    elif latency_seconds is not None and latency_seconds < 0:
        add_anomaly(
            anomalies,
            "时间异常",
            "注册时间晚于首充时间",
            record,
            included=True,
        )
    return record


def read_inputs(
    input_paths: Iterable[Path],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    records: list[dict[str, Any]] = []
    anomalies: list[dict[str, Any]] = []
    input_summary = {
        "files": [],
        "valid_sheets": 0,
        "raw_rows": 0,
        "parsed_rows": 0,
        "skipped_sheets": [],
    }

    for input_path in input_paths:
        resolved = input_path.resolve()
        if not resolved.is_file():
            raise ValueError(f"输入文件不存在: {resolved}")
        if resolved.suffix.lower() != ".xlsx":
            raise ValueError(f"仅支持 .xlsx 输入文件: {resolved.name}")
        input_summary["files"].append(resolved.name)
        workbook = load_workbook(resolved, read_only=True, data_only=True)
        workbook_has_data = False
        try:
            for ws in workbook.worksheets:
                header = find_header_row(ws)
                if header is None:
                    input_summary["skipped_sheets"].append(f"{resolved.name}/{ws.title}")
                    continue
                workbook_has_data = True
                input_summary["valid_sheets"] += 1
                header_row, column_mapping = header
                max_required_column = max(column_mapping.values())
                for source_row, values in enumerate(
                    ws.iter_rows(min_row=header_row + 1, values_only=True),
                    start=header_row + 1,
                ):
                    if not any(value is not None and clean_text(value) for value in values):
                        continue
                    input_summary["raw_rows"] += 1
                    raw = {
                        canonical: values[column_index]
                        if column_index < len(values)
                        else None
                        for canonical, column_index in column_mapping.items()
                    }
                    if len(values) <= max_required_column:
                        add_anomaly(
                            anomalies,
                            "行结构错误",
                            "该行列数少于必填字段所在列",
                            raw_record_for_anomaly(
                                raw, resolved.name, ws.title, source_row
                            ),
                            included=False,
                        )
                        continue
                    record = parse_data_row(
                        raw,
                        resolved.name,
                        ws.title,
                        source_row,
                        anomalies,
                    )
                    if record is not None:
                        records.append(record)
                        input_summary["parsed_rows"] += 1
        finally:
            workbook.close()
        if not workbook_has_data:
            raise ValueError(
                f"{resolved.name} 中未找到包含必填字段的工作表: "
                + ", ".join(sorted(REQUIRED_COLUMNS))
            )
    return records, anomalies, input_summary


def filter_and_deduplicate(
    records: list[dict[str, Any]],
    anomalies: list[dict[str, Any]],
    start_date: date | None,
    end_date: date | None,
) -> tuple[list[dict[str, Any]], int, int]:
    in_range: list[dict[str, Any]] = []
    excluded_by_date = 0
    for record in records:
        business_date = record["recharge_time"].date()
        if start_date and business_date < start_date:
            excluded_by_date += 1
            continue
        if end_date and business_date > end_date:
            excluded_by_date += 1
            continue
        in_range.append(record)

    by_user: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in in_range:
        by_user[record["user_id"]].append(record)

    formal_records: list[dict[str, Any]] = []
    duplicate_count = 0
    for user_id, user_records in by_user.items():
        ordered = sorted(
            user_records,
            key=lambda row: (
                row["recharge_time"],
                row["source_file"],
                row["source_sheet"],
                row["source_row"],
            ),
        )
        first = ordered[0]
        first["global_sequence"] = 1
        formal_records.append(first)
        if first.get("source_sequence") not in (None, 1):
            add_anomaly(
                anomalies,
                "首条记录缺失",
                f"该用户最早输入记录的 recharge_sequence={first['source_sequence']}，请检查导出范围",
                first,
                included=True,
            )
        for sequence, duplicate in enumerate(ordered[1:], start=2):
            duplicate["global_sequence"] = sequence
            duplicate_count += 1
            add_anomaly(
                anomalies,
                "重复首充记录",
                f"同一 user_id 在输入文件中出现 {len(ordered)} 条首充记录，本条未计入正式统计",
                duplicate,
                included=False,
            )

    formal_records.sort(key=lambda row: (row["recharge_time"], row["user_id"]))
    return formal_records, duplicate_count, excluded_by_date


def latency_bucket(seconds: int | None) -> str:
    if seconds is None:
        return "注册时间缺失"
    if seconds < 0:
        return "时间异常"
    if seconds < 10 * 60:
        return "00-10分钟"
    if seconds < 30 * 60:
        return "10-30分钟"
    if seconds < 60 * 60:
        return "30-60分钟"
    if seconds < 6 * 60 * 60:
        return "01-06小时"
    if seconds < 24 * 60 * 60:
        return "06-24小时"
    if seconds < 3 * 24 * 60 * 60:
        return "01-03天"
    if seconds < 7 * 24 * 60 * 60:
        return "03-07天"
    return "07天以上"


def date_range(start_date: date, end_date: date) -> list[date]:
    return [
        start_date + timedelta(days=offset)
        for offset in range((end_date - start_date).days + 1)
    ]


def decimal_median(values: list[int]) -> Decimal | None:
    if not values:
        return None
    ordered = sorted(values)
    middle = len(ordered) // 2
    if len(ordered) % 2:
        return Decimal(ordered[middle])
    return (Decimal(ordered[middle - 1]) + Decimal(ordered[middle])) / Decimal(2)


def aggregate_records(
    records: list[dict[str, Any]], start_date: date, end_date: date
) -> dict[str, Any]:
    total_users = len(records)
    total_amount = sum(
        (row["standard_amount_usd"] for row in records), Decimal(0)
    )
    valid_latency = [
        row["latency_seconds"]
        for row in records
        if row["latency_seconds"] is not None and row["latency_seconds"] >= 0
    ]
    median_seconds = decimal_median(valid_latency)

    channels: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"users": 0, "amount": Decimal(0), "countries": set()}
    )
    tiers: dict[tuple[Any, ...], dict[str, Any]] = defaultdict(
        lambda: {
            "users": 0,
            "amount": Decimal(0),
            "countries": set(),
            "latencies": [],
            "within24": 0,
        }
    )
    countries: dict[tuple[str, str], dict[str, Any]] = defaultdict(
        lambda: {
            "users": 0,
            "amount": Decimal(0),
            "channels": defaultdict(lambda: {"users": 0, "amount": Decimal(0)}),
        }
    )
    country_tiers: dict[tuple[Any, ...], dict[str, Any]] = defaultdict(
        lambda: {"users": 0, "amount": Decimal(0)}
    )
    daily: dict[date, dict[str, Any]] = defaultdict(
        lambda: {
            "users": 0,
            "amount": Decimal(0),
            "channels": defaultdict(lambda: {"users": 0, "amount": Decimal(0)}),
            "countries": set(),
        }
    )
    daily_tiers: dict[tuple[Any, ...], dict[str, Any]] = defaultdict(
        lambda: {"users": 0, "amount": Decimal(0)}
    )
    latency: dict[str, dict[str, Any]] = {
        bucket: {"users": 0, "amount": Decimal(0)} for bucket in LATENCY_BUCKETS
    }

    for row in records:
        amount = row["standard_amount_usd"]
        business_date = row["recharge_time"].date()
        channel = row["channel"]
        country_key = (row["country_code"], row["country_name"])
        country_is_known = (
            row["country_code"] != "UNKNOWN" and row["country_name"] != "未知国家"
        )
        tier_key = (
            channel,
            row["price_id"],
            row["tier_coin"],
            row["standard_amount_usd"],
        )

        channels[channel]["users"] += 1
        channels[channel]["amount"] += amount
        if country_is_known:
            channels[channel]["countries"].add(country_key)

        tiers[tier_key]["users"] += 1
        tiers[tier_key]["amount"] += amount
        if country_is_known:
            tiers[tier_key]["countries"].add(country_key)
        if row["latency_seconds"] is not None and row["latency_seconds"] >= 0:
            tiers[tier_key]["latencies"].append(row["latency_seconds"])
            if row["latency_seconds"] <= 24 * 60 * 60:
                tiers[tier_key]["within24"] += 1

        countries[country_key]["users"] += 1
        countries[country_key]["amount"] += amount
        countries[country_key]["channels"][channel]["users"] += 1
        countries[country_key]["channels"][channel]["amount"] += amount

        country_tier_key = country_key + tier_key
        country_tiers[country_tier_key]["users"] += 1
        country_tiers[country_tier_key]["amount"] += amount

        daily[business_date]["users"] += 1
        daily[business_date]["amount"] += amount
        daily[business_date]["channels"][channel]["users"] += 1
        daily[business_date]["channels"][channel]["amount"] += amount
        if country_is_known:
            daily[business_date]["countries"].add(country_key)

        daily_tier_key = (business_date,) + tier_key
        daily_tiers[daily_tier_key]["users"] += 1
        daily_tiers[daily_tier_key]["amount"] += amount

        bucket = latency_bucket(row["latency_seconds"])
        latency[bucket]["users"] += 1
        latency[bucket]["amount"] += amount

        row["latency_bucket"] = bucket

    daily_rows = []
    for business_date in date_range(start_date, end_date):
        data = daily[business_date]
        google = data["channels"].get("Google", {"users": 0, "amount": Decimal(0)})
        apple = data["channels"].get("Apple", {"users": 0, "amount": Decimal(0)})
        other_users = data["users"] - google["users"] - apple["users"]
        other_amount = data["amount"] - google["amount"] - apple["amount"]
        daily_rows.append(
            {
                "dt": business_date,
                "users": data["users"],
                "amount": data["amount"],
                "average": data["amount"] / data["users"] if data["users"] else Decimal(0),
                "google_users": google["users"],
                "google_amount": google["amount"],
                "apple_users": apple["users"],
                "apple_amount": apple["amount"],
                "other_users": other_users,
                "other_amount": other_amount,
                "countries": len(data["countries"]),
            }
        )

    channel_rows = []
    for channel, data in channels.items():
        channel_rows.append(
            {
                "channel": channel,
                "users": data["users"],
                "amount": data["amount"],
                "user_share": Decimal(data["users"]) / total_users if total_users else Decimal(0),
                "amount_share": data["amount"] / total_amount if total_amount else Decimal(0),
                "countries": len(data["countries"]),
                "average": data["amount"] / data["users"] if data["users"] else Decimal(0),
            }
        )
    channel_rows.sort(key=lambda row: (-row["amount"], row["channel"]))

    tier_rows = []
    for tier_key, data in tiers.items():
        channel, price_id, tier_coin, standard_amount = tier_key
        latency_count = len(data["latencies"])
        tier_rows.append(
            {
                "channel": channel,
                "price_id": price_id,
                "tier_coin": tier_coin,
                "standard_amount": standard_amount,
                "users": data["users"],
                "amount": data["amount"],
                "user_share": Decimal(data["users"]) / total_users if total_users else Decimal(0),
                "amount_share": data["amount"] / total_amount if total_amount else Decimal(0),
                "countries": len(data["countries"]),
                "average_latency_hours": (
                    Decimal(sum(data["latencies"])) / Decimal(latency_count) / Decimal(3600)
                    if latency_count
                    else None
                ),
                "within24_share": (
                    Decimal(data["within24"]) / Decimal(latency_count)
                    if latency_count
                    else None
                ),
            }
        )
    tier_rows.sort(
        key=lambda row: (-row["amount"], row["channel"], row["standard_amount"])
    )

    country_rows = []
    for country_key, data in countries.items():
        country_code, country_name = country_key
        google = data["channels"].get("Google", {"users": 0, "amount": Decimal(0)})
        apple = data["channels"].get("Apple", {"users": 0, "amount": Decimal(0)})
        country_rows.append(
            {
                "country_code": country_code,
                "country_name": country_name,
                "users": data["users"],
                "amount": data["amount"],
                "average": data["amount"] / data["users"] if data["users"] else Decimal(0),
                "user_share": Decimal(data["users"]) / total_users if total_users else Decimal(0),
                "amount_share": data["amount"] / total_amount if total_amount else Decimal(0),
                "google_users": google["users"],
                "google_amount": google["amount"],
                "apple_users": apple["users"],
                "apple_amount": apple["amount"],
                "other_users": data["users"] - google["users"] - apple["users"],
                "other_amount": data["amount"] - google["amount"] - apple["amount"],
            }
        )
    country_rows.sort(key=lambda row: (-row["amount"], row["country_name"]))

    country_tier_rows = []
    for key, data in country_tiers.items():
        country_code, country_name, channel, price_id, tier_coin, standard_amount = key
        country_data = countries[(country_code, country_name)]
        country_tier_rows.append(
            {
                "country_code": country_code,
                "country_name": country_name,
                "channel": channel,
                "price_id": price_id,
                "tier_coin": tier_coin,
                "standard_amount": standard_amount,
                "users": data["users"],
                "amount": data["amount"],
                "country_user_share": Decimal(data["users"]) / country_data["users"],
                "country_amount_share": data["amount"] / country_data["amount"]
                if country_data["amount"]
                else Decimal(0),
                "period_user_share": Decimal(data["users"]) / total_users
                if total_users
                else Decimal(0),
                "period_amount_share": data["amount"] / total_amount
                if total_amount
                else Decimal(0),
            }
        )
    country_tier_rows.sort(
        key=lambda row: (
            -row["amount"],
            row["country_name"],
            row["channel"],
            row["standard_amount"],
        )
    )

    daily_tier_rows = []
    for key, data in daily_tiers.items():
        business_date, channel, price_id, tier_coin, standard_amount = key
        day_data = daily[business_date]
        daily_tier_rows.append(
            {
                "dt": business_date,
                "channel": channel,
                "price_id": price_id,
                "tier_coin": tier_coin,
                "standard_amount": standard_amount,
                "users": data["users"],
                "amount": data["amount"],
                "daily_user_share": Decimal(data["users"]) / day_data["users"],
                "daily_amount_share": data["amount"] / day_data["amount"]
                if day_data["amount"]
                else Decimal(0),
                "period_user_share": Decimal(data["users"]) / total_users
                if total_users
                else Decimal(0),
                "period_amount_share": data["amount"] / total_amount
                if total_amount
                else Decimal(0),
            }
        )
    daily_tier_rows.sort(
        key=lambda row: (row["dt"], row["channel"], row["standard_amount"])
    )

    latency_rows = []
    for bucket in LATENCY_BUCKETS:
        data = latency[bucket]
        latency_rows.append(
            {
                "bucket": bucket,
                "users": data["users"],
                "amount": data["amount"],
                "user_share": Decimal(data["users"]) / total_users if total_users else Decimal(0),
                "amount_share": data["amount"] / total_amount if total_amount else Decimal(0),
                "average": data["amount"] / data["users"] if data["users"] else Decimal(0),
            }
        )

    return {
        "total_users": total_users,
        "total_amount": total_amount,
        "average_amount": total_amount / total_users if total_users else Decimal(0),
        "country_count": sum(
            country_code != "UNKNOWN" and country_name != "未知国家"
            for country_code, country_name in countries
        ),
        "unknown_country_users": sum(
            row["country_code"] == "UNKNOWN" or row["country_name"] == "未知国家"
            for row in records
        ),
        "valid_latency_count": len(valid_latency),
        "registration_coverage": Decimal(len(valid_latency)) / total_users
        if total_users
        else Decimal(0),
        "median_latency_hours": median_seconds / Decimal(3600)
        if median_seconds is not None
        else None,
        "within24_count": sum(seconds <= 24 * 60 * 60 for seconds in valid_latency),
        "within24_share": Decimal(
            sum(seconds <= 24 * 60 * 60 for seconds in valid_latency)
        )
        / Decimal(len(valid_latency))
        if valid_latency
        else None,
        "channel_rows": channel_rows,
        "tier_rows": tier_rows,
        "country_rows": country_rows,
        "country_tier_rows": country_tier_rows,
        "daily_rows": daily_rows,
        "daily_tier_rows": daily_tier_rows,
        "latency_rows": latency_rows,
        "records": records,
    }


def style_title(ws, title: str, subtitle: str, end_column: int) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    ws["A2"].alignment = Alignment(vertical="center")
    for column in range(1, end_column + 1):
        ws.cell(3, column).border = SECTION_BORDER
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 22


def style_header(ws, row: int, start_column: int, end_column: int) -> None:
    for column in range(start_column, end_column + 1):
        cell = ws.cell(row, column)
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            left=Side(style="thin", color="FFFFFF"),
            right=Side(style="thin", color="FFFFFF"),
        )
    ws.row_dimensions[row].height = 30


def style_section(
    ws, row: int, title: str, start_column: int, end_column: int
) -> None:
    for column in range(start_column, end_column + 1):
        cell = ws.cell(row, column)
        cell.fill = SECTION_FILL
        cell.border = SECTION_BORDER
    ws.cell(row, start_column, title)
    ws.cell(row, start_column).font = Font(
        name=FONT_NAME, size=11, bold=True, color="1F4E78"
    )
    ws.row_dimensions[row].height = 24


def style_body_row(ws, row: int, start_column: int, end_column: int) -> None:
    for column in range(start_column, end_column + 1):
        cell = ws.cell(row, column)
        cell.font = BODY_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = BOTTOM_BORDER


def set_widths(ws, widths: list[float]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def finish_table_sheet(
    ws, header_row: int, end_column: int, widths: list[float], freeze_column: int = 0
) -> None:
    ws.freeze_panes = ws.cell(header_row + 1, freeze_column + 1)
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(end_column)}{max(ws.max_row, header_row)}"
    )
    set_widths(ws, widths)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def write_summary(
    ws,
    analysis: dict[str, Any],
    input_summary: dict[str, Any],
    anomalies: list[dict[str, Any]],
    duplicate_count: int,
    excluded_by_date: int,
    start_date: date,
    end_date: date,
) -> None:
    style_title(
        ws,
        "首次充值分析",
        f"统计范围：{start_date:%Y-%m-%d} 至 {end_date:%Y-%m-%d}；金额单位：USD；每位用户仅计最早一条首充",
        11,
    )

    context = [
        ("输入文件", f"{len(input_summary['files'])} 个", "有效工作表", input_summary["valid_sheets"]),
        ("输入数据行", input_summary["raw_rows"], "解析成功行", input_summary["parsed_rows"]),
        ("重复首充行", duplicate_count, "日期范围外行", excluded_by_date),
        ("异常记录", len(anomalies), "生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for row_number, values in enumerate(context, start=4):
        for column, value in zip((1, 2, 5, 6), values):
            ws.cell(row_number, column, value)
            ws.cell(row_number, column).font = LABEL_FONT if column in (1, 5) else BODY_FONT
        for column in range(1, 9):
            ws.cell(row_number, column).border = BOTTOM_BORDER

    style_section(ws, 9, "期间指标", 1, 4)
    kpis = [
        ("首充用户数", analysis["total_users"], INTEGER_FORMAT),
        ("首充总金额", analysis["total_amount"], USD_FORMAT),
        ("人均首充金额", analysis["average_amount"], USD_FORMAT),
        ("涉及国家数", analysis["country_count"], INTEGER_FORMAT),
        ("未知国家用户数", analysis["unknown_country_users"], INTEGER_FORMAT),
        ("注册时间覆盖率", analysis["registration_coverage"], PERCENT_FORMAT),
        ("注册至首充中位时长", analysis["median_latency_hours"], HOURS_FORMAT),
        ("可计算用户24小时内首充占比", analysis["within24_share"], PERCENT_FORMAT),
    ]
    for offset, (label, value, number_format) in enumerate(kpis, start=1):
        row_number = 9 + offset
        ws.cell(row_number, 1, label).font = LABEL_FONT
        ws.cell(row_number, 2, value if value is not None else "n.a.").font = BODY_FONT
        if value is not None:
            ws.cell(row_number, 2).number_format = number_format
        for column in range(1, 4):
            ws.cell(row_number, column).border = BOTTOM_BORDER

    style_section(ws, 9, "渠道汇总", 5, 11)
    channel_headers = ["渠道", "用户数", "首充金额", "用户占比", "金额占比", "涉及国家", "人均金额"]
    for column, header in enumerate(channel_headers, start=5):
        ws.cell(10, column, header)
    style_header(ws, 10, 5, 11)
    for row_number, row in enumerate(analysis["channel_rows"], start=11):
        values = [
            row["channel"],
            row["users"],
            row["amount"],
            row["user_share"],
            row["amount_share"],
            row["countries"],
            row["average"],
        ]
        for column, value in enumerate(values, start=5):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 5, 11)
        ws.cell(row_number, 6).number_format = INTEGER_FORMAT
        ws.cell(row_number, 7).number_format = USD_FORMAT
        ws.cell(row_number, 8).number_format = PERCENT_FORMAT
        ws.cell(row_number, 9).number_format = PERCENT_FORMAT
        ws.cell(row_number, 10).number_format = INTEGER_FORMAT
        ws.cell(row_number, 11).number_format = USD_FORMAT

    ranking_start = 19
    style_section(
        ws,
        ranking_start,
        "渠道档位排名（按首充金额从高到低）",
        1,
        11,
    )
    headers = [
        "排名",
        "渠道",
        "priceId",
        "档位金币",
        "标准金额",
        "用户数",
        "首充金额",
        "用户占比",
        "金额占比",
        "涉及国家",
        "24小时内占比",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(ranking_start + 1, column, header)
    style_header(ws, ranking_start + 1, 1, 11)
    for rank, row in enumerate(analysis["tier_rows"], start=1):
        row_number = ranking_start + 1 + rank
        values = [
            rank,
            row["channel"],
            row["price_id"],
            row["tier_coin"],
            row["standard_amount"],
            row["users"],
            row["amount"],
            row["user_share"],
            row["amount_share"],
            row["countries"],
            row["within24_share"] if row["within24_share"] is not None else "n.a.",
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 1, 11)
        ws.cell(row_number, 3).number_format = TEXT_FORMAT
        ws.cell(row_number, 3).quotePrefix = True
        ws.cell(row_number, 4).number_format = INTEGER_FORMAT
        ws.cell(row_number, 5).number_format = USD_FORMAT
        ws.cell(row_number, 6).number_format = INTEGER_FORMAT
        ws.cell(row_number, 7).number_format = USD_FORMAT
        ws.cell(row_number, 8).number_format = PERCENT_FORMAT
        ws.cell(row_number, 9).number_format = PERCENT_FORMAT
        ws.cell(row_number, 10).number_format = INTEGER_FORMAT
        if row["within24_share"] is not None:
            ws.cell(row_number, 11).number_format = PERCENT_FORMAT

    set_widths(ws, [22, 16, 24, 16, 18, 16, 18, 16, 16, 16, 18])
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def write_daily(ws, analysis: dict[str, Any], start_date: date, end_date: date) -> None:
    style_title(
        ws,
        "每天汇总",
        f"统计范围：{start_date:%Y-%m-%d} 至 {end_date:%Y-%m-%d}；无首充日期按零保留",
        11,
    )
    header_row = 4
    headers = [
        "日期",
        "首充用户数",
        "首充金额（USD）",
        "人均金额（USD）",
        "Google用户数",
        "Google金额（USD）",
        "Apple用户数",
        "Apple金额（USD）",
        "其他渠道用户数",
        "其他渠道金额（USD）",
        "涉及国家数",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(header_row, column, header)
    style_header(ws, header_row, 1, len(headers))
    for row_number, row in enumerate(analysis["daily_rows"], start=header_row + 1):
        values = [
            row["dt"],
            row["users"],
            row["amount"],
            row["average"],
            row["google_users"],
            row["google_amount"],
            row["apple_users"],
            row["apple_amount"],
            row["other_users"],
            row["other_amount"],
            row["countries"],
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 1, len(headers))
        ws.cell(row_number, 1).number_format = DATE_NUMBER_FORMAT
        for column in (2, 5, 7, 9, 11):
            ws.cell(row_number, column).number_format = INTEGER_FORMAT
        for column in (3, 4, 6, 8, 10):
            ws.cell(row_number, column).number_format = USD_FORMAT

    finish_table_sheet(
        ws,
        header_row,
        len(headers),
        [14, 16, 20, 20, 18, 20, 18, 20, 18, 20, 16],
    )

    if analysis["daily_rows"]:
        chart = LineChart()
        chart.title = "每日首充金额"
        chart.style = 13
        chart.y_axis.title = "USD"
        chart.x_axis = DateAxis(crossAx=100)
        chart.x_axis.title = "日期"
        chart.x_axis.number_format = "yyyy-mm-dd"
        chart.x_axis.majorTimeUnit = "days"
        chart.y_axis.numFmt = '$#,##0.00'
        chart.height = 9
        chart.width = 18
        categories = Reference(
            ws,
            min_col=1,
            min_row=header_row + 1,
            max_row=header_row + len(analysis["daily_rows"]),
        )
        for data_column, series_title in (
            (3, "全部渠道（USD）"),
            (6, "Google（USD）"),
            (8, "Apple（USD）"),
        ):
            data = Reference(
                ws,
                min_col=data_column,
                max_col=data_column,
                min_row=header_row + 1,
                max_row=header_row + len(analysis["daily_rows"]),
            )
            chart.add_data(data, titles_from_data=False)
            chart.series[-1].tx = SeriesLabel(v=series_title)
        chart.set_categories(categories)
        chart.legend.position = "t"
        ws.add_chart(chart, "M4")


def write_channel_tier(
    ws, analysis: dict[str, Any], start_date: date, end_date: date
) -> None:
    style_title(
        ws,
        "渠道档位汇总",
        f"统计范围：{start_date:%Y-%m-%d} 至 {end_date:%Y-%m-%d}；按首充金额从高到低排列",
        11,
    )
    header_row = 4
    headers = [
        "渠道",
        "priceId",
        "档位金币",
        "标准金额（USD）",
        "首充用户数",
        "首充金额（USD）",
        "用户占比",
        "金额占比",
        "涉及国家数",
        "平均注册至首充（小时）",
        "24小时内占比",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(header_row, column, header)
    style_header(ws, header_row, 1, len(headers))
    for row_number, row in enumerate(analysis["tier_rows"], start=header_row + 1):
        values = [
            row["channel"],
            row["price_id"],
            row["tier_coin"],
            row["standard_amount"],
            row["users"],
            row["amount"],
            row["user_share"],
            row["amount_share"],
            row["countries"],
            row["average_latency_hours"] if row["average_latency_hours"] is not None else "n.a.",
            row["within24_share"] if row["within24_share"] is not None else "n.a.",
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 1, len(headers))
        ws.cell(row_number, 2).number_format = TEXT_FORMAT
        ws.cell(row_number, 2).quotePrefix = True
        ws.cell(row_number, 3).number_format = INTEGER_FORMAT
        ws.cell(row_number, 4).number_format = USD_FORMAT
        ws.cell(row_number, 5).number_format = INTEGER_FORMAT
        ws.cell(row_number, 6).number_format = USD_FORMAT
        for column in (7, 8, 11):
            if isinstance(ws.cell(row_number, column).value, (int, float, Decimal)):
                ws.cell(row_number, column).number_format = PERCENT_FORMAT
        ws.cell(row_number, 9).number_format = INTEGER_FORMAT
        if row["average_latency_hours"] is not None:
            ws.cell(row_number, 10).number_format = HOURS_FORMAT
    finish_table_sheet(
        ws,
        header_row,
        len(headers),
        [14, 24, 16, 20, 18, 20, 16, 16, 16, 24, 18],
        freeze_column=2,
    )


def write_country_summary(
    ws, analysis: dict[str, Any], start_date: date, end_date: date
) -> None:
    style_title(
        ws,
        "国家汇总",
        f"统计范围：{start_date:%Y-%m-%d} 至 {end_date:%Y-%m-%d}；国家取充值记录中的 country_code",
        13,
    )
    header_row = 4
    headers = [
        "国家代码",
        "国家名称",
        "首充用户数",
        "首充金额（USD）",
        "人均金额（USD）",
        "用户占比",
        "金额占比",
        "Google用户数",
        "Google金额（USD）",
        "Apple用户数",
        "Apple金额（USD）",
        "其他渠道用户数",
        "其他渠道金额（USD）",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(header_row, column, header)
    style_header(ws, header_row, 1, len(headers))
    for row_number, row in enumerate(analysis["country_rows"], start=header_row + 1):
        values = [
            row["country_code"],
            row["country_name"],
            row["users"],
            row["amount"],
            row["average"],
            row["user_share"],
            row["amount_share"],
            row["google_users"],
            row["google_amount"],
            row["apple_users"],
            row["apple_amount"],
            row["other_users"],
            row["other_amount"],
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 1, len(headers))
        for column in (3, 8, 10, 12):
            ws.cell(row_number, column).number_format = INTEGER_FORMAT
        for column in (4, 5, 9, 11, 13):
            ws.cell(row_number, column).number_format = USD_FORMAT
        for column in (6, 7):
            ws.cell(row_number, column).number_format = PERCENT_FORMAT
        if row["country_code"] == "UNKNOWN" or row["country_name"] == "未知国家":
            for cell in ws[row_number][: len(headers)]:
                cell.fill = WARNING_FILL
    finish_table_sheet(
        ws,
        header_row,
        len(headers),
        [14, 24, 18, 20, 20, 16, 16, 18, 20, 18, 20, 18, 20],
        freeze_column=2,
    )


def write_country_tier(
    ws, analysis: dict[str, Any], start_date: date, end_date: date
) -> None:
    style_title(
        ws,
        "国家档位汇总",
        f"统计范围：{start_date:%Y-%m-%d} 至 {end_date:%Y-%m-%d}；展示国家、渠道和首充档位交叉结果",
        12,
    )
    header_row = 4
    headers = [
        "国家代码",
        "国家名称",
        "渠道",
        "priceId",
        "档位金币",
        "标准金额（USD）",
        "首充用户数",
        "首充金额（USD）",
        "国家内用户占比",
        "国家内金额占比",
        "周期用户占比",
        "周期金额占比",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(header_row, column, header)
    style_header(ws, header_row, 1, len(headers))
    for row_number, row in enumerate(analysis["country_tier_rows"], start=header_row + 1):
        values = [
            row["country_code"],
            row["country_name"],
            row["channel"],
            row["price_id"],
            row["tier_coin"],
            row["standard_amount"],
            row["users"],
            row["amount"],
            row["country_user_share"],
            row["country_amount_share"],
            row["period_user_share"],
            row["period_amount_share"],
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 1, len(headers))
        ws.cell(row_number, 4).number_format = TEXT_FORMAT
        ws.cell(row_number, 4).quotePrefix = True
        ws.cell(row_number, 5).number_format = INTEGER_FORMAT
        ws.cell(row_number, 6).number_format = USD_FORMAT
        ws.cell(row_number, 7).number_format = INTEGER_FORMAT
        ws.cell(row_number, 8).number_format = USD_FORMAT
        for column in range(9, 13):
            ws.cell(row_number, column).number_format = PERCENT_FORMAT
        if row["country_code"] == "UNKNOWN" or row["country_name"] == "未知国家":
            for cell in ws[row_number][: len(headers)]:
                cell.fill = WARNING_FILL
    finish_table_sheet(
        ws,
        header_row,
        len(headers),
        [14, 24, 14, 24, 16, 20, 18, 20, 18, 18, 16, 16],
        freeze_column=4,
    )


def write_daily_tier(
    ws, analysis: dict[str, Any], start_date: date, end_date: date
) -> None:
    style_title(
        ws,
        "每天档位汇总",
        f"统计范围：{start_date:%Y-%m-%d} 至 {end_date:%Y-%m-%d}；仅展示当天实际发生首充的档位",
        11,
    )
    header_row = 4
    headers = [
        "日期",
        "渠道",
        "priceId",
        "档位金币",
        "标准金额（USD）",
        "首充用户数",
        "首充金额（USD）",
        "当日用户占比",
        "当日金额占比",
        "周期用户占比",
        "周期金额占比",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(header_row, column, header)
    style_header(ws, header_row, 1, len(headers))
    for row_number, row in enumerate(analysis["daily_tier_rows"], start=header_row + 1):
        values = [
            row["dt"],
            row["channel"],
            row["price_id"],
            row["tier_coin"],
            row["standard_amount"],
            row["users"],
            row["amount"],
            row["daily_user_share"],
            row["daily_amount_share"],
            row["period_user_share"],
            row["period_amount_share"],
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 1, len(headers))
        ws.cell(row_number, 1).number_format = DATE_NUMBER_FORMAT
        ws.cell(row_number, 3).number_format = TEXT_FORMAT
        ws.cell(row_number, 3).quotePrefix = True
        ws.cell(row_number, 4).number_format = INTEGER_FORMAT
        ws.cell(row_number, 5).number_format = USD_FORMAT
        ws.cell(row_number, 6).number_format = INTEGER_FORMAT
        ws.cell(row_number, 7).number_format = USD_FORMAT
        for column in range(8, 12):
            ws.cell(row_number, column).number_format = PERCENT_FORMAT
    finish_table_sheet(
        ws,
        header_row,
        len(headers),
        [14, 14, 24, 16, 20, 18, 20, 18, 18, 16, 16],
        freeze_column=3,
    )


def write_latency(
    ws, analysis: dict[str, Any], start_date: date, end_date: date
) -> None:
    style_title(
        ws,
        "注册至首充",
        f"统计范围：{start_date:%Y-%m-%d} 至 {end_date:%Y-%m-%d}；耗时由用户注册时间和首充时间重新计算",
        6,
    )
    header_row = 4
    headers = [
        "耗时区间",
        "首充用户数",
        "首充金额（USD）",
        "用户占比",
        "金额占比",
        "人均金额（USD）",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(header_row, column, header)
    style_header(ws, header_row, 1, len(headers))
    for row_number, row in enumerate(analysis["latency_rows"], start=header_row + 1):
        values = [
            row["bucket"],
            row["users"],
            row["amount"],
            row["user_share"],
            row["amount_share"],
            row["average"],
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 1, len(headers))
        ws.cell(row_number, 2).number_format = INTEGER_FORMAT
        ws.cell(row_number, 3).number_format = USD_FORMAT
        ws.cell(row_number, 4).number_format = PERCENT_FORMAT
        ws.cell(row_number, 5).number_format = PERCENT_FORMAT
        ws.cell(row_number, 6).number_format = USD_FORMAT
        if row["bucket"] in ("注册时间缺失", "时间异常"):
            for cell in ws[row_number][: len(headers)]:
                cell.fill = WARNING_FILL
    finish_table_sheet(ws, header_row, len(headers), [20, 18, 20, 16, 16, 20])


def write_detail(
    ws, analysis: dict[str, Any], start_date: date, end_date: date
) -> None:
    style_title(
        ws,
        "首充明细",
        f"统计范围：{start_date:%Y-%m-%d} 至 {end_date:%Y-%m-%d}；每位用户仅保留最早一条有效首充",
        22,
    )
    header_row = 4
    headers = [
        "userId",
        "昵称",
        "首充时间",
        "首充日期",
        "渠道",
        "priceId",
        "档位金币",
        "标准金额（USD）",
        "实际金额",
        "实际金币",
        "国家代码",
        "国家名称",
        "注册时间",
        "注册至首充（小时）",
        "耗时区间",
        "realPerson",
        "proxyId",
        "firstProxyId",
        "firstJoinTime",
        "SQL成功记录数",
        "来源文件",
        "来源行",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(header_row, column, header)
    style_header(ws, header_row, 1, len(headers))
    for row_number, row in enumerate(analysis["records"], start=header_row + 1):
        latency_hours = (
            Decimal(row["latency_seconds"]) / Decimal(3600)
            if row["latency_seconds"] is not None
            else None
        )
        values = [
            row["user_id"],
            row["nick"],
            row["recharge_time"],
            row["recharge_time"].date(),
            row["channel"],
            row["price_id"],
            row["tier_coin"],
            row["standard_amount_usd"],
            row["actual_money"],
            row["actual_coin"],
            row["country_code"],
            row["country_name"],
            row["register_time"],
            latency_hours,
            row["latency_bucket"],
            row["real_person"],
            row["proxy_id"],
            row["first_proxy_id"],
            row["first_join_time"],
            row["successful_count"],
            row["source_file"],
            f"{row['source_sheet']}!{row['source_row']}",
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 1, len(headers))
        for column in (1, 6, 17, 18):
            ws.cell(row_number, column).number_format = TEXT_FORMAT
            ws.cell(row_number, column).quotePrefix = True
        ws.cell(row_number, 3).number_format = DATETIME_NUMBER_FORMAT
        ws.cell(row_number, 4).number_format = DATE_NUMBER_FORMAT
        ws.cell(row_number, 7).number_format = INTEGER_FORMAT
        ws.cell(row_number, 8).number_format = USD_FORMAT
        if row["actual_money"] is not None:
            ws.cell(row_number, 9).number_format = USD_FORMAT
        if row["actual_coin"] is not None:
            ws.cell(row_number, 10).number_format = INTEGER_FORMAT
        if row["register_time"] is not None:
            ws.cell(row_number, 13).number_format = DATETIME_NUMBER_FORMAT
        if latency_hours is not None:
            ws.cell(row_number, 14).number_format = HOURS_FORMAT
        if row["first_join_time"] is not None:
            ws.cell(row_number, 19).number_format = DATETIME_NUMBER_FORMAT
        if row["successful_count"] is not None:
            ws.cell(row_number, 20).number_format = INTEGER_FORMAT
        if row["country_code"] == "UNKNOWN" or row["latency_bucket"] in (
            "注册时间缺失",
            "时间异常",
        ):
            for cell in ws[row_number][: len(headers)]:
                cell.fill = WARNING_FILL

    finish_table_sheet(
        ws,
        header_row,
        len(headers),
        [
            16,
            20,
            22,
            14,
            14,
            24,
            16,
            18,
            16,
            16,
            14,
            24,
            22,
            20,
            18,
            14,
            18,
            18,
            22,
            18,
            28,
            18,
        ],
        freeze_column=2,
    )
    if ws.max_row > header_row:
        ws.conditional_formatting.add(
            f"N{header_row + 1}:N{ws.max_row}",
            CellIsRule(
                operator="lessThan",
                formula=["0"],
                fill=ERROR_FILL,
                font=Font(name=FONT_NAME, color="9C0006", bold=True),
            ),
        )


def write_anomalies(
    ws, anomalies: list[dict[str, Any]], start_date: date, end_date: date
) -> None:
    style_title(
        ws,
        "数据异常",
        f"统计范围：{start_date:%Y-%m-%d} 至 {end_date:%Y-%m-%d}；“计入统计”表示该问题未阻止记录参与正式分析",
        13,
    )
    header_row = 4
    headers = [
        "异常类型",
        "异常说明",
        "userId",
        "首充时间",
        "渠道",
        "priceId",
        "标准金额（USD）",
        "国家代码",
        "国家名称",
        "来源文件",
        "来源工作表",
        "来源行",
        "计入统计",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(header_row, column, header)
    style_header(ws, header_row, 1, len(headers))

    if anomalies:
        ordered = sorted(
            anomalies,
            key=lambda row: (
                row["included"] == "是",
                row["type"],
                str(row["source_file"]),
                str(row["source_row"]),
            ),
        )
        for row_number, row in enumerate(ordered, start=header_row + 1):
            values = [
                row["type"],
                row["detail"],
                row["user_id"],
                row["recharge_time"],
                row["channel"],
                row["price_id"],
                row["standard_amount_usd"],
                row["country_code"],
                row["country_name"],
                row["source_file"],
                row["source_sheet"],
                row["source_row"],
                row["included"],
            ]
            for column, value in enumerate(values, start=1):
                ws.cell(row_number, column, value)
            style_body_row(ws, row_number, 1, len(headers))
            for column in (3, 6):
                ws.cell(row_number, column).number_format = TEXT_FORMAT
                ws.cell(row_number, column).quotePrefix = True
            if isinstance(row["recharge_time"], (datetime, date)):
                ws.cell(row_number, 4).number_format = DATETIME_NUMBER_FORMAT
            if isinstance(row["standard_amount_usd"], (int, float, Decimal)):
                ws.cell(row_number, 7).number_format = USD_FORMAT
            fill = WARNING_FILL if row["included"] == "是" else ERROR_FILL
            for cell in ws[row_number][: len(headers)]:
                cell.fill = fill
    else:
        ws.cell(header_row + 1, 1, "无异常")
        style_body_row(ws, header_row + 1, 1, len(headers))

    finish_table_sheet(
        ws,
        header_row,
        len(headers),
        [22, 48, 16, 22, 14, 24, 18, 14, 24, 28, 20, 12, 14],
        freeze_column=2,
    )


def build_workbook(
    analysis: dict[str, Any],
    input_summary: dict[str, Any],
    anomalies: list[dict[str, Any]],
    duplicate_count: int,
    excluded_by_date: int,
    start_date: date,
    end_date: date,
) -> Workbook:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    sheets = {
        "分析汇总": workbook.create_sheet("分析汇总"),
        "每天汇总": workbook.create_sheet("每天汇总"),
        "渠道档位汇总": workbook.create_sheet("渠道档位汇总"),
        "国家汇总": workbook.create_sheet("国家汇总"),
        "国家档位汇总": workbook.create_sheet("国家档位汇总"),
        "每天档位汇总": workbook.create_sheet("每天档位汇总"),
        "注册至首充": workbook.create_sheet("注册至首充"),
        "首充明细": workbook.create_sheet("首充明细"),
        "数据异常": workbook.create_sheet("数据异常"),
    }

    write_summary(
        sheets["分析汇总"],
        analysis,
        input_summary,
        anomalies,
        duplicate_count,
        excluded_by_date,
        start_date,
        end_date,
    )
    write_daily(sheets["每天汇总"], analysis, start_date, end_date)
    write_channel_tier(sheets["渠道档位汇总"], analysis, start_date, end_date)
    write_country_summary(sheets["国家汇总"], analysis, start_date, end_date)
    write_country_tier(sheets["国家档位汇总"], analysis, start_date, end_date)
    write_daily_tier(sheets["每天档位汇总"], analysis, start_date, end_date)
    write_latency(sheets["注册至首充"], analysis, start_date, end_date)
    write_detail(sheets["首充明细"], analysis, start_date, end_date)
    write_anomalies(sheets["数据异常"], anomalies, start_date, end_date)

    sheets["分析汇总"].sheet_properties.tabColor = "1F4E78"
    sheets["每天汇总"].sheet_properties.tabColor = "5B9BD5"
    sheets["数据异常"].sheet_properties.tabColor = "C65911"
    workbook.active = 0
    workbook.properties.title = "首次充值分析"
    workbook.properties.subject = "首充用户、渠道档位、国家和注册至首充耗时分析"
    workbook.properties.creator = "first-recharge-analysis"
    return workbook


def validate_paths(args: argparse.Namespace) -> None:
    output = args.output.resolve()
    if output.suffix.lower() != ".xlsx":
        raise ValueError("--output 必须是 .xlsx 文件")
    input_resolved = {path.resolve() for path in args.input}
    if output in input_resolved:
        raise ValueError("输出文件不能覆盖输入文件")


def run() -> None:
    args = parse_args()
    validate_paths(args)
    records, anomalies, input_summary = read_inputs(args.input)
    formal_records, duplicate_count, excluded_by_date = filter_and_deduplicate(
        records,
        anomalies,
        args.start_date,
        args.end_date,
    )
    if not formal_records:
        raise ValueError("统计范围内没有可用于分析的有效首充记录")

    start_date = args.start_date or min(
        row["recharge_time"].date() for row in formal_records
    )
    end_date = args.end_date or max(
        row["recharge_time"].date() for row in formal_records
    )
    analysis = aggregate_records(formal_records, start_date, end_date)
    workbook = build_workbook(
        analysis,
        input_summary,
        anomalies,
        duplicate_count,
        excluded_by_date,
        start_date,
        end_date,
    )

    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"报告已生成：{output}")
    print(
        f"首充用户={analysis['total_users']:,}，"
        f"首充金额=${analysis['total_amount']:,.2f}，"
        f"重复首充记录={duplicate_count:,}，异常项={len(anomalies):,}"
    )


def main() -> int:
    try:
        run()
        return 0
    except (OSError, ValueError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
