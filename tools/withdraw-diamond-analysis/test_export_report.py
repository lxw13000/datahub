"""自动 ES 导出与 Excel 报告的本地回归测试，不访问真实集群。"""

import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import export_report as report
import fetch_es


def bucket(user_id, status, business_type, tokens):
    return {
        "key": {"userId": user_id, "status": status, "businessType": business_type},
        "doc_count": 1,
        "total_tokens": {"value": float(tokens)},
    }


def response(buckets, after_key=None):
    aggregation = {"buckets": buckets}
    if after_key:
        aggregation["after_key"] = after_key
    return {
        "timed_out": False,
        "_shards": {"total": 1, "successful": 1, "failed": 0},
        "aggregations": {"user_status_business": aggregation},
    }


class FakeES:
    def __init__(self):
        self.queries = []

    def search(self, query):
        self.queries.append(query)
        after = query["aggs"]["user_status_business"]["composite"].get("after")
        if after is None:
            return response([bucket(101, 1, 2, 20_000_000)],
                            {"userId": 101, "status": 1, "businessType": 2})
        return response([bucket(101, -1, 53, 10_000_000)])


class WithdrawAutoReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        (self.root / "withdraw-user-ids.csv").write_text(
            "userId\n101\n102\n101\n", encoding="utf-8"
        )
        self.es_config_path = self.root / "es-config.json"
        self.analysis_config_path = self.root / "analysis-config.json"
        self.es_config = {
            "es_url": "https://test-es:9200",
            "username": "test_user",
            "password": "private_password",
        }
        self.analysis_config = {
            "user_ids_file": "withdraw-user-ids.csv",
            "withdraw_start": "2026-07-20 00:00:00",
            "withdraw_end": "2026-07-31 23:59:59",
            "statistics_start": "2026-07-01 00:00:00",
            "statistics_end": "2026-07-31 23:59:59",
            "exclude_business_types": [74, 503],
        }
        self.es_config_path.write_text(json.dumps(self.es_config), encoding="utf-8")
        self.analysis_config_path.write_text(json.dumps(self.analysis_config), encoding="utf-8")
        self.paths = fetch_es.output_paths(self.root, "test")

    def test_default_outputs_are_next_to_python_files(self):
        paths = fetch_es.output_paths()
        tool_dir = Path(fetch_es.__file__).resolve().parent
        self.assertEqual(paths.data.parent, tool_dir)
        self.assertEqual(paths.manifest.parent, tool_dir)
        self.assertEqual(paths.report.parent, tool_dir)
        self.assertTrue(paths.data.name.startswith("withdraw-analysis-"))
        self.assertEqual(paths.report.suffix, ".xlsx")

    def test_csv_loader_deduplicates_and_skips_blank_rows(self):
        csv_path = self.root / "withdraw-user-ids.csv"
        csv_path.write_text("userId\n\n101\n  \n102\n101\n", encoding="utf-8")
        self.assertEqual(report.load_user_ids(csv_path), {101, 102})

    def test_auto_export_pages_and_keeps_zero_flow_user(self):
        fake = FakeES()
        with patch.object(fetch_es, "ESClient", return_value=fake), \
                patch.object(fetch_es, "output_paths", return_value=self.paths), \
                patch.object(sys, "argv", ["export_report.py", "--es-config", str(self.es_config_path),
                                           "--analysis-config", str(self.analysis_config_path)]):
            self.assertEqual(report.main(), 0)

        data = json.loads(self.paths.data.read_text(encoding="utf-8"))
        self.assertEqual(data["code"], 200)
        self.assertEqual(data["data"], [
            {"userId": 101, "status": 1, "businessType": 2, "tokens": 20_000_000},
            {"userId": 101, "status": -1, "businessType": 53, "tokens": 10_000_000},
        ])
        self.assertEqual(len(fake.queries), 2)
        first = fake.queries[0]
        self.assertEqual(first["query"]["bool"]["filter"][0]["terms"]["user_id"], [101, 102])
        self.assertEqual(first["query"]["bool"]["filter"][2]["range"]["create_time"],
                         {"gte": self.analysis_config["statistics_start"],
                          "lte": self.analysis_config["statistics_end"]})
        self.assertEqual(first["query"]["bool"]["must_not"],
                         [{"terms": {"business_type": [74, 503]}}])
        self.assertEqual(fake.queries[1]["aggs"]["user_status_business"]["composite"]["after"],
                         {"userId": 101, "status": 1, "businessType": 2})

        manifest = self.paths.manifest.read_text(encoding="utf-8")
        self.assertNotIn("private_password", manifest)
        self.assertEqual(json.loads(manifest)["query"]["user_count"], 2)
        self.assertEqual(json.loads(manifest)["status"], "report_complete")
        book = load_workbook(self.paths.report)
        self.assertEqual(book.sheetnames, ["分析汇总", "用户汇总", "加余额明细",
                                           "减余额明细", "业务类型汇总"])
        self.assertEqual([book["分析汇总"][cell].value for cell in ("A7", "B7", "C7", "D7")],
                         [2, 20_000_000, 10_000_000, 10_000_000])
        self.assertEqual([book["分析汇总"][cell].value for cell in ("F7", "G7", "H7")],
                         [2, 1, 1])
        self.assertEqual([book["用户汇总"][cell].value for cell in ("A3", "B3", "D3", "F3")],
                         [102, 0, 0, 0])
        self.assertEqual([book["用户汇总"][cell].value for cell in ("C2", "E2", "G2")],
                         [2, 1, 1])
        self.assertEqual(book["用户汇总"]["C2"].number_format, report.USD_FORMAT)
        self.assertEqual(book["加余额明细"]["E2"].value, 2)
        self.assertEqual(book["减余额明细"]["E2"].value, 1)
        self.assertEqual(book["业务类型汇总"]["F2"].value, 2)
        summary = book["分析汇总"]
        top_row = next(row for row in range(1, summary.max_row + 1)
                       if summary.cell(row, 1).value == "用户排名（最多 20 名）")
        self.assertEqual(summary.cell(top_row + 2, 5).value, 2)
        max_row = next(row for row in range(1, summary.max_row + 1)
                       if summary.cell(row, 1).value == "最大用户占比及剔除影响")
        self.assertEqual(summary.cell(max_row + 2, 10).value, 2)
        self.assertEqual(summary.cell(max_row + 2, 14).value, 0)
        self.assertTrue(all("美元" in book[name][cell].value for name, cell in (
            ("用户汇总", "C1"), ("加余额明细", "E1"),
            ("减余额明细", "E1"), ("业务类型汇总", "F1"))))
        self.assertTrue(all(sheet.sheet_view.showGridLines for sheet in book))
        self.assertEqual(book["用户汇总"]["A2"].border.left.style, "thin")
        self.assertEqual(book["用户汇总"]["A2"].border.left.color.rgb[-6:], "B7C5D3")
        self.assertFalse(any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))
        book.close()

    def test_rejects_unsafe_sum_without_new_json_or_report(self):
        config = fetch_es.load_config(self.es_config_path, self.analysis_config_path)

        class UnsafeES:
            def search(self, _query):
                return response([bucket(101, 1, 2, 2**53)])

        with self.assertRaisesRegex(ValueError, "安全整数"):
            fetch_es.export_json(config, {101, 102}, set(), self.paths, UnsafeES())
        self.assertFalse(self.paths.data.exists())
        self.assertFalse(self.paths.data.with_name(f".{self.paths.data.name}.part").exists())
        self.assertFalse(self.paths.report.exists())

    def test_decimal_conversion_after_integer_aggregation(self):
        aggregated, ids = report.aggregate_records([
            {"userId": 101, "status": 1, "businessType": 2, "tokens": 3_000},
            {"userId": 101, "status": 1, "businessType": 2, "tokens": 3_000},
        ], {101}, set())
        self.assertEqual(aggregated[(101, 1, 2)], 6_000)
        self.assertEqual(report.tokens_to_usd(aggregated[(101, 1, 2)]), 0.001)
        self.assertEqual(report.tokens_to_usd(-5_000), -0.001)
        args = type("Args", (), {"withdraw_start": "2026-01-01 00:00:00",
                                 "withdraw_end": "2026-01-31 23:59:59",
                                 "statistics_start": "2026-01-01 00:00:00",
                                 "statistics_end": "2026-01-31 23:59:59",
                                 "input": Path("input.json"), "user_ids": Path("ids.csv")})()
        book = report.create_workbook(report.build_analysis(aggregated, ids), aggregated, args, set(), 2)
        self.assertEqual(book["用户汇总"]["C2"].value, 0.001)
        self.assertEqual(book["加余额明细"]["E2"].value, 0.001)

    def test_resume_skips_completed_user_batch(self):
        config = fetch_es.load_config(self.es_config_path, self.analysis_config_path)
        ids = set(range(1, 1002))

        class InterruptedES:
            def __init__(self):
                self.calls = 0

            def search(self, query):
                self.calls += 1
                if self.calls == 2:
                    raise ValueError("模拟网络中断")
                return response([bucket(1, 1, 2, 10_000_000)])

        first = InterruptedES()
        with self.assertRaisesRegex(ValueError, "模拟网络中断"):
            fetch_es.export_json(config, ids, set(), self.paths, first)
        manifest = json.loads(self.paths.manifest.read_text(encoding="utf-8"))
        self.assertEqual(manifest["status"], "in_progress")
        self.assertEqual(list(manifest["completed_batches"]), ["0"])
        self.assertFalse(self.paths.data.exists())

        class RemainingES:
            def __init__(self):
                self.queries = []

            def search(self, query):
                self.queries.append(query)
                return response([bucket(1001, -1, 53, 5_000_000)])

        second = RemainingES()
        selected = fetch_es.find_resume_paths(config, ids, set(), self.root)
        self.assertEqual(selected, self.paths)
        fetch_es.export_json(config, ids, set(), selected, second, resume=True)
        self.assertEqual(len(second.queries), 1)
        self.assertEqual(second.queries[0]["query"]["bool"]["filter"][0]["terms"]["user_id"], [1001])
        self.assertEqual(len(json.loads(self.paths.data.read_text(encoding="utf-8"))["data"]), 2)
        with self.assertRaisesRegex(ValueError, "查询口径或用户 ID 名单"):
            fetch_es.export_json(config, ids | {1002}, set(), selected, second, resume=True)
        self.assertEqual(len(second.queries), 1)

    def test_resume_cli_finishes_interrupted_run(self):
        ids_path = self.root / "withdraw-user-ids.csv"
        ids_path.write_text("userId\n" + "\n".join(map(str, range(1, 1002))), encoding="utf-8")

        class InterruptedES:
            def __init__(self):
                self.calls = 0

            def search(self, _query):
                self.calls += 1
                if self.calls == 2:
                    raise ValueError("模拟中断")
                return response([bucket(1, 1, 2, 10_000_000)])

        class RemainingES:
            def __init__(self):
                self.calls = 0

            def search(self, _query):
                self.calls += 1
                return response([bucket(1001, -1, 53, 5_000_000)])

        first, second = InterruptedES(), RemainingES()
        argv = ["export_report.py", "--es-config", str(self.es_config_path),
                "--analysis-config", str(self.analysis_config_path)]
        with patch.object(fetch_es, "ESClient", side_effect=[first, second]), \
                patch.object(fetch_es, "output_paths", return_value=self.paths), \
                patch.object(fetch_es, "find_resume_paths", return_value=self.paths), \
                patch.object(sys, "argv", argv):
            self.assertEqual(report.main(), 2)
            self.assertFalse(self.paths.report.exists())
            with patch.object(sys, "argv", argv + ["--resume"]):
                self.assertEqual(report.main(), 0)
        self.assertEqual(first.calls, 2)
        self.assertEqual(second.calls, 1)
        self.assertTrue(self.paths.report.exists())
        self.assertEqual(json.loads(self.paths.manifest.read_text(encoding="utf-8"))["status"], "report_complete")

    def test_batches_users_and_rejects_partial_es_response(self):
        config = fetch_es.load_config(self.es_config_path, self.analysis_config_path)

        class EmptyES:
            def __init__(self):
                self.batch_sizes = []

            def search(self, query):
                self.batch_sizes.append(
                    len(query["query"]["bool"]["filter"][0]["terms"]["user_id"])
                )
                return response([])

        fake = EmptyES()
        target = fetch_es.export_json(config, set(range(1, 1002)), set(), self.paths, fake)
        self.assertEqual(fake.batch_sizes, [1000, 1])
        self.assertEqual(json.loads(target.read_text(encoding="utf-8"))["data"], [])

        class FailedES:
            def search(self, _query):
                result = response([])
                result["_shards"]["failed"] = 1
                return result

        with self.assertRaisesRegex(ValueError, "分片失败"):
            fetch_es.export_json(config, {101}, set(),
                                 fetch_es.output_paths(self.root, "failed"), FailedES())
        self.assertFalse((self.root / ".withdraw-analysis-failed.json.part").exists())

    def test_es_client_posts_to_diamond_index_with_basic_auth(self):
        config = fetch_es.load_config(self.es_config_path, self.analysis_config_path)
        captured = []

        def fake_urlopen(request, **kwargs):
            captured.append((request, kwargs))
            return io.BytesIO(json.dumps(response([])).encode("utf-8"))

        with patch.object(fetch_es, "urlopen", side_effect=fake_urlopen):
            result = fetch_es.ESClient(config).search({"size": 0})
        self.assertEqual(result["_shards"]["failed"], 0)
        request, options = captured[0]
        self.assertEqual(request.full_url,
                         "https://test-es:9200/sano_wallet_diamond_record/_search")
        self.assertEqual(request.get_method(), "POST")
        self.assertEqual(json.loads(request.data), {"size": 0})
        self.assertTrue(request.get_header("Authorization").startswith("Basic "))
        self.assertIsNotNone(options["context"])

    def test_rejects_invalid_config_and_empty_users(self):
        self.analysis_config["exclude_business_types"] = [True]
        self.analysis_config_path.write_text(json.dumps(self.analysis_config), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "exclude_business_types"):
            fetch_es.load_config(self.es_config_path, self.analysis_config_path)
        self.analysis_config["exclude_business_types"] = []
        self.analysis_config_path.write_text(json.dumps(self.analysis_config), encoding="utf-8")
        (self.root / "withdraw-user-ids.csv").write_text("userId\n", encoding="utf-8")
        with patch.object(fetch_es, "ESClient") as es_client, \
                patch.object(fetch_es, "output_paths", return_value=self.paths), \
                patch.object(sys, "argv", ["export_report.py", "--es-config", str(self.es_config_path),
                                           "--analysis-config", str(self.analysis_config_path)]):
            self.assertEqual(report.main(), 2)
            es_client.assert_not_called()
        self.assertFalse(self.paths.report.exists())


if __name__ == "__main__":
    unittest.main()
