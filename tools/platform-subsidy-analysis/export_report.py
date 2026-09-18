#!/usr/bin/env python3
"""将金币、钻石平台补贴的每日业务汇总 JSON 导出为 Excel 报告。"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


COIN_TOKENS_PER_USD = 10_000
DIAMOND_TOKENS_PER_USD = 10_000_000
DATE_FORMAT = "%Y-%m-%d"
INTEGER_PATTERN = re.compile(r"[+-]?\d+")

# 与两个接口当前使用的 EBusinessType 固定集合保持一致。
COIN_BUSINESS_TYPES = {
    97: "幸运礼物消费榜单奖励",
    99: "游戏消费榜单奖励",
}
DIAMOND_BUSINESS_TYPES = {
    74: "直播时长奖励",
    91: "优质主播时长奖励",
    93: "代理超级主播收益奖励",
    94: "代理超级主播幸运礼物收益奖励",
    95: "超级主播收益奖励",
    96: "超级主播幸运礼物收益奖励",
    98: "幸运礼物收入榜单奖励",
    100: "游戏收入榜单奖励",
    101: "代理收入榜单奖励",
    109: "优质主播奖励",
    110: "PK榜单奖励",
}

FONT_NAME = "Arial"
TITLE_FONT = Font(name=FONT_NAME, size=18, bold=True, color="1F1F1F")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, italic=True, color="666666")
BODY_FONT = Font(name=FONT_NAME, size=10, color="1F1F1F")
LABEL_FONT = Font(name=FONT_NAME, size=10, bold=True, color="1F1F1F")
WHITE_BOLD_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
TOTAL_FILL = PatternFill("solid", fgColor="E2F0D9")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
LIGHT_SIDE = Side(style="thin", color="D9E2F3")
MEDIUM_SIDE = Side(style="medium", color="9EADBA")
BOTTOM_BORDER = Border(bottom=LIGHT_SIDE)
SECTION_BORDER = Border(bottom=MEDIUM_SIDE)
TOKEN_FORMAT = "#,##0"
USD_FORMAT = '"$"#,##0.000;[Red]-"$"#,##0.000'
PERCENT_FORMAT = "0.00%"
DATE_NUMBER_FORMAT = "yyyy-mm-dd"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "将 /walletCoinAnalysis/platformSubsidyByType 和 "
            "/walletDiamondAnalysis/platformSubsidyByType 的每日汇总 JSON 导出为 Excel。"
        )
    )
    parser.add_argument("--coin-input", required=True, type=Path, help="金币接口结果 JSON")
    parser.add_argument(
        "--diamond-input", required=True, type=Path, help="钻石接口结果 JSON"
    )
    parser.add_argument("--output", required=True, type=Path, help="输出 .xlsx 文件")
    parser.add_argument("--start-date", required=True, help="统计开始日期，yyyy-MM-dd")
    parser.add_argument("--end-date", required=True, help="统计结束日期，yyyy-MM-dd")
    return parser.parse_args()


def parse_integer(value: Any, field: str, location: str) -> int:
    """只接受整数或整数字符串，避免 Tokens 经由 float 转换。"""
    if type(value) is int:
        return value
    if isinstance(value, str) and INTEGER_PATTERN.fullmatch(value.strip()):
        return int(value.strip())
    raise ValueError(f"{location} 的 {field} 必须是整数，实际值为 {value!r}")


def reject_json_float(value: str) -> None:
    raise ValueError(f"JSON 中出现小数 {value}；businessType 和 tokens 必须使用整数")


def parse_date(value: Any, field: str, location: str) -> date:
    if not isinstance(value, str):
        raise ValueError(f"{location} 的 {field} 必须是 yyyy-MM-dd 字符串")
    try:
        return datetime.strptime(value, DATE_FORMAT).date()
    except ValueError as error:
        raise ValueError(f"{location} 的 {field} 格式错误，请使用 yyyy-MM-dd") from error


def validate_date_range(start_text: str, end_text: str) -> tuple[date, date]:
    start = parse_date(start_text, "startDate", "命令行")
    end = parse_date(end_text, "endDate", "命令行")
    if end < start:
        raise ValueError("统计结束日期不能早于开始日期")
    return start, end


def dates_between(start: date, end: date) -> list[date]:
    days = (end - start).days + 1
    return [start + timedelta(days=offset) for offset in range(days)]


def unwrap_response(path: Path, label: str) -> list[dict[str, Any]]:
    if not path.is_file():
        raise ValueError(f"{label} JSON 不存在或不是文件：{path}")
    try:
        with path.open("r", encoding="utf-8-sig") as source:
            root = json.load(source, parse_int=int, parse_float=reject_json_float)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
        raise ValueError(f"读取{label} JSON 失败：{error}") from error

    if isinstance(root, list):
        records = root
    elif isinstance(root, dict):
        if root.get("success") is False:
            raise ValueError(f"{label}接口结果 success=false，不能生成报告")
        if "code" in root and parse_integer(root["code"], "code", f"{label} JSON 根节点") != 200:
            raise ValueError(f"{label}接口结果 code={root['code']}，预期为 200")
        records = root.get("data")
        if not isinstance(records, list):
            raise ValueError(f"{label}对象根节点必须包含数组类型的 data 字段")
    else:
        raise ValueError(f"{label} JSON 根节点必须是数组，或包含 data 数组的对象")

    for index, record in enumerate(records, start=1):
        if not isinstance(record, dict):
            raise ValueError(f"{label}第 {index} 条日汇总不是 JSON 对象")
    return records


def load_daily_totals(
    path: Path,
    label: str,
    start: date,
    end: date,
) -> tuple[dict[tuple[date, int], int], int, int]:
    """解析并按 日期 + businessType 二次汇总，兼容重复日期或重复业务项。"""
    daily_records = unwrap_response(path, label)
    totals: dict[tuple[date, int], int] = defaultdict(int)
    stat_count = 0
    for daily_index, daily_record in enumerate(daily_records, start=1):
        location = f"{label}第 {daily_index} 条日汇总"
        if "dt" not in daily_record or "stats" not in daily_record:
            raise ValueError(f"{location}必须包含 dt 和 stats 字段")
        business_date = parse_date(daily_record["dt"], "dt", location)
        if business_date < start or business_date > end:
            raise ValueError(
                f"{location}日期 {business_date:%Y-%m-%d} 超出命令行统计范围"
            )
        stats = daily_record["stats"]
        if not isinstance(stats, list):
            raise ValueError(f"{location}的 stats 必须是数组")
        for stat_index, stat in enumerate(stats, start=1):
            stat_location = f"{location}的第 {stat_index} 个业务项"
            if not isinstance(stat, dict):
                raise ValueError(f"{stat_location}不是 JSON 对象")
            missing = [field for field in ("businessType", "tokens") if field not in stat]
            if missing:
                raise ValueError(f"{stat_location}缺少字段：{', '.join(missing)}")
            business_type = parse_integer(stat["businessType"], "businessType", stat_location)
            tokens = parse_integer(stat["tokens"], "tokens", stat_location)
            if business_type < 0:
                raise ValueError(f"{stat_location}的 businessType 不能小于 0")
            totals[(business_date, business_type)] += tokens
            stat_count += 1
    return dict(totals), len(daily_records), stat_count


def usd_from_tokens(tokens: int, tokens_per_usd: int) -> Decimal:
    """保留精确换算值，工作簿用数字格式显示三位小数。"""
    return Decimal(tokens) / Decimal(tokens_per_usd)


def build_wallet_rows(
    wallet_type: str,
    totals: dict[tuple[date, int], int],
    expected_types: dict[int, str],
    tokens_per_usd: int,
    all_dates: list[date],
) -> list[dict[str, Any]]:
    observed_types = {business_type for _, business_type in totals}
    business_types = list(expected_types)
    business_types.extend(sorted(observed_types - set(expected_types)))
    rows: list[dict[str, Any]] = []
    for business_date in all_dates:
        for business_type in business_types:
            tokens = totals.get((business_date, business_type), 0)
            rows.append(
                {
                    "dt": business_date,
                    "wallet_type": wallet_type,
                    "business_type": business_type,
                    "name": expected_types.get(business_type, "未知业务类型"),
                    "tokens": tokens,
                    "tokens_per_usd": tokens_per_usd,
                    "usd": usd_from_tokens(tokens, tokens_per_usd),
                    "is_expected": business_type in expected_types,
                }
            )
    return rows


def sum_decimal(values: Iterable[Decimal]) -> Decimal:
    return sum(values, Decimal(0))


def build_analysis(
    coin_rows: list[dict[str, Any]],
    diamond_rows: list[dict[str, Any]],
    all_dates: list[date],
) -> dict[str, Any]:
    all_rows = coin_rows + diamond_rows
    period_total = sum_decimal(row["usd"] for row in all_rows)
    wallet_totals: dict[str, dict[str, Any]] = {}
    for wallet_type, rows, rate in (
        ("金币", coin_rows, COIN_TOKENS_PER_USD),
        ("钻石", diamond_rows, DIAMOND_TOKENS_PER_USD),
    ):
        wallet_totals[wallet_type] = {
            "tokens": sum(row["tokens"] for row in rows),
            "usd": sum_decimal(row["usd"] for row in rows),
            "tokens_per_usd": rate,
            "business_count": len({row["business_type"] for row in rows}),
            "nonzero_business_count": len(
                {row["business_type"] for row in rows if row["tokens"] != 0}
            ),
        }

    daily_lookup: dict[date, dict[str, Any]] = {}
    cumulative = Decimal(0)
    for business_date in all_dates:
        coin_day_rows = [row for row in coin_rows if row["dt"] == business_date]
        diamond_day_rows = [row for row in diamond_rows if row["dt"] == business_date]
        coin_tokens = sum(row["tokens"] for row in coin_day_rows)
        diamond_tokens = sum(row["tokens"] for row in diamond_day_rows)
        coin_usd = sum_decimal(row["usd"] for row in coin_day_rows)
        diamond_usd = sum_decimal(row["usd"] for row in diamond_day_rows)
        total_usd = coin_usd + diamond_usd
        cumulative += total_usd
        daily_lookup[business_date] = {
            "dt": business_date,
            "coin_tokens": coin_tokens,
            "coin_usd": coin_usd,
            "diamond_tokens": diamond_tokens,
            "diamond_usd": diamond_usd,
            "total_usd": total_usd,
            "period_share": total_usd / period_total if period_total else Decimal(0),
            "cumulative_usd": cumulative,
        }
    daily_rows = [daily_lookup[business_date] for business_date in all_dates]

    business_rows: list[dict[str, Any]] = []
    for wallet_type, rows in (("金币", coin_rows), ("钻石", diamond_rows)):
        wallet_total = wallet_totals[wallet_type]["usd"]
        business_types = sorted({row["business_type"] for row in rows})
        for business_type in business_types:
            matching = [row for row in rows if row["business_type"] == business_type]
            tokens = sum(row["tokens"] for row in matching)
            usd = sum_decimal(row["usd"] for row in matching)
            max_row = max(matching, key=lambda row: (row["usd"], -row["dt"].toordinal()))
            business_rows.append(
                {
                    "wallet_type": wallet_type,
                    "business_type": business_type,
                    "name": matching[0]["name"],
                    "tokens": tokens,
                    "tokens_per_usd": matching[0]["tokens_per_usd"],
                    "usd": usd,
                    "wallet_share": usd / wallet_total if wallet_total else Decimal(0),
                    "period_share": usd / period_total if period_total else Decimal(0),
                    "active_days": sum(row["tokens"] != 0 for row in matching),
                    "daily_average_usd": usd / Decimal(len(all_dates)),
                    "max_daily_usd": max_row["usd"],
                    "max_date": max_row["dt"],
                    "is_expected": matching[0]["is_expected"],
                }
            )
    business_rows.sort(
        key=lambda row: (-row["usd"], row["wallet_type"], row["business_type"])
    )

    # 每个原子业务行补充其所在日期与钱包的占比，确保四张表口径互相可追溯。
    daily_wallet_totals: dict[tuple[date, str], Decimal] = defaultdict(Decimal)
    for row in all_rows:
        daily_wallet_totals[(row["dt"], row["wallet_type"])] += row["usd"]
    for row in all_rows:
        daily_wallet_total = daily_wallet_totals[(row["dt"], row["wallet_type"])]
        daily_total = daily_lookup[row["dt"]]["total_usd"]
        row["daily_wallet_share"] = row["usd"] / daily_wallet_total if daily_wallet_total else Decimal(0)
        row["daily_total_share"] = row["usd"] / daily_total if daily_total else Decimal(0)
        row["period_share"] = row["usd"] / period_total if period_total else Decimal(0)

    highest_day = max(daily_rows, key=lambda row: (row["total_usd"], -row["dt"].toordinal()))
    return {
        "all_rows": all_rows,
        "wallet_totals": wallet_totals,
        "business_rows": business_rows,
        "daily_rows": daily_rows,
        "period_total": period_total,
        "day_count": len(all_dates),
        "days_with_subsidy": sum(row["total_usd"] != 0 for row in daily_rows),
        "days_without_subsidy": sum(row["total_usd"] == 0 for row in daily_rows),
        "daily_average_usd": period_total / Decimal(len(all_dates)),
        "highest_day": highest_day,
    }


def style_title(ws, title: str, subtitle: str, end_column: int) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    ws["A2"].alignment = Alignment(vertical="center")
    for column in range(1, end_column + 1):
        ws.cell(3, column).border = SECTION_BORDER
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 22


def style_header(ws, row: int, start_column: int, end_column: int) -> None:
    for column in range(start_column, end_column + 1):
        cell = ws.cell(row, column)
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            left=Side(style="thin", color="FFFFFF"),
            right=Side(style="thin", color="FFFFFF"),
        )
    ws.row_dimensions[row].height = 30


def style_section(ws, row: int, title: str, end_column: int) -> None:
    for column in range(1, end_column + 1):
        cell = ws.cell(row, column)
        cell.fill = SECTION_FILL
        cell.border = SECTION_BORDER
    ws.cell(row, 1, title)
    ws.cell(row, 1).font = Font(name=FONT_NAME, size=11, bold=True, color="1F4E78")
    ws.row_dimensions[row].height = 24


def style_body_row(ws, row: int, start_column: int, end_column: int) -> None:
    for column in range(start_column, end_column + 1):
        cell = ws.cell(row, column)
        cell.font = BODY_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = BOTTOM_BORDER


def set_widths(ws, widths: list[float]) -> None:
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width


def format_standard_sheet(ws, header_row: int, widths: list[float]) -> None:
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{max(header_row, ws.max_row)}"
    set_widths(ws, widths)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def write_total_summary(
    ws,
    analysis: dict[str, Any],
    args: argparse.Namespace,
    coin_day_count: int,
    coin_stat_count: int,
    diamond_day_count: int,
    diamond_stat_count: int,
) -> None:
    style_title(
        ws,
        "平台补贴汇总",
        f"统计范围：{args.start_date} 至 {args.end_date}；金额单位：USD",
        11,
    )

    context = [
        ("金币换算", f"{COIN_TOKENS_PER_USD:,} Tokens = 1 USD", "金币输入", args.coin_input.name),
        ("钻石换算", f"{DIAMOND_TOKENS_PER_USD:,} Tokens = 1 USD", "钻石输入", args.diamond_input.name),
        ("输入记录", f"金币 {coin_day_count} 天/{coin_stat_count} 项；钻石 {diamond_day_count} 天/{diamond_stat_count} 项", "生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for row_number, (left_label, left_value, right_label, right_value) in enumerate(context, start=4):
        ws.cell(row_number, 1, left_label).font = LABEL_FONT
        ws.cell(row_number, 2, left_value).font = BODY_FONT
        ws.cell(row_number, 5, right_label).font = LABEL_FONT
        ws.cell(row_number, 6, right_value).font = BODY_FONT
        for column in range(1, 9):
            ws.cell(row_number, column).border = BOTTOM_BORDER

    style_section(ws, 8, "期间指标", 4)
    kpis = [
        ("补贴总额", analysis["period_total"], USD_FORMAT),
        ("统计天数", analysis["day_count"], TOKEN_FORMAT),
        ("日均补贴", analysis["daily_average_usd"], USD_FORMAT),
        ("有补贴天数", analysis["days_with_subsidy"], TOKEN_FORMAT),
        ("无补贴天数", analysis["days_without_subsidy"], TOKEN_FORMAT),
        ("最高补贴日期", analysis["highest_day"]["dt"], DATE_NUMBER_FORMAT),
        ("最高单日补贴", analysis["highest_day"]["total_usd"], USD_FORMAT),
    ]
    for offset, (label, value, number_format) in enumerate(kpis, start=1):
        row_number = 8 + offset
        ws.cell(row_number, 1, label).font = LABEL_FONT
        ws.cell(row_number, 2, value).font = BODY_FONT
        ws.cell(row_number, 2).number_format = number_format
        for column in range(1, 4):
            ws.cell(row_number, column).border = BOTTOM_BORDER

    for column in range(5, 12):
        ws.cell(8, column).fill = SECTION_FILL
        ws.cell(8, column).border = SECTION_BORDER
    ws.cell(8, 5, "钱包汇总")
    ws.cell(8, 5).font = Font(name=FONT_NAME, size=11, bold=True, color="1F4E78")
    wallet_headers = ["钱包", "业务数", "非零业务数", "Tokens", "Tokens/美元", "补贴金额", "总额占比"]
    for column, header in enumerate(wallet_headers, start=5):
        ws.cell(9, column, header)
    style_header(ws, 9, 5, 11)
    for row_number, wallet_type in enumerate(("金币", "钻石"), start=10):
        wallet = analysis["wallet_totals"][wallet_type]
        values = [
            wallet_type,
            wallet["business_count"],
            wallet["nonzero_business_count"],
            wallet["tokens"],
            wallet["tokens_per_usd"],
            wallet["usd"],
            wallet["usd"] / analysis["period_total"] if analysis["period_total"] else Decimal(0),
        ]
        for column, value in enumerate(values, start=5):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 5, 11)
        ws.cell(row_number, 8).number_format = TOKEN_FORMAT
        ws.cell(row_number, 9).number_format = TOKEN_FORMAT
        ws.cell(row_number, 10).number_format = USD_FORMAT
        ws.cell(row_number, 11).number_format = PERCENT_FORMAT
    total_row = 12
    total_values = ["合计", "—", "—", "不同单位不合计", "—", analysis["period_total"], Decimal(1) if analysis["period_total"] else Decimal(0)]
    for column, value in enumerate(total_values, start=5):
        cell = ws.cell(total_row, column, value)
        cell.fill = TOTAL_FILL
        cell.font = LABEL_FONT
        cell.border = BOTTOM_BORDER
    ws.cell(total_row, 10).number_format = USD_FORMAT
    ws.cell(total_row, 11).number_format = PERCENT_FORMAT

    ranking_start = 18
    style_section(ws, ranking_start, "业务补贴金额排名", 7)
    ranking_headers = ["排名", "钱包", "业务类型", "业务名称", "Tokens", "补贴金额", "总额占比"]
    for column, header in enumerate(ranking_headers, start=1):
        ws.cell(ranking_start + 1, column, header)
    style_header(ws, ranking_start + 1, 1, 7)
    for rank, business in enumerate(analysis["business_rows"], start=1):
        row_number = ranking_start + 1 + rank
        values = [rank, business["wallet_type"], business["business_type"], business["name"], business["tokens"], business["usd"], business["period_share"]]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 1, 7)
        ws.cell(row_number, 5).number_format = TOKEN_FORMAT
        ws.cell(row_number, 6).number_format = USD_FORMAT
        ws.cell(row_number, 7).number_format = PERCENT_FORMAT
        if not business["is_expected"]:
            for cell in ws[row_number][:7]:
                cell.fill = WARNING_FILL

    set_widths(ws, [18, 28, 14, 34, 15, 24, 16, 22, 18, 20, 16])
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def write_business_summary(ws, analysis: dict[str, Any], args: argparse.Namespace) -> None:
    style_title(
        ws,
        "按业务汇总",
        f"统计范围：{args.start_date} 至 {args.end_date}；默认按补贴金额从高到低排列",
        12,
    )
    headers = [
        "钱包类型",
        "业务类型",
        "业务名称",
        "Tokens总数",
        "Tokens/美元",
        "补贴金额（USD）",
        "钱包内占比",
        "总补贴占比",
        "有补贴天数",
        "周期日均（USD）",
        "最高单日（USD）",
        "最高日期",
    ]
    header_row = 4
    for column, header in enumerate(headers, start=1):
        ws.cell(header_row, column, header)
    style_header(ws, header_row, 1, len(headers))
    for offset, business in enumerate(analysis["business_rows"], start=1):
        row_number = header_row + offset
        values = [
            business["wallet_type"],
            business["business_type"],
            business["name"],
            business["tokens"],
            business["tokens_per_usd"],
            business["usd"],
            business["wallet_share"],
            business["period_share"],
            business["active_days"],
            business["daily_average_usd"],
            business["max_daily_usd"],
            business["max_date"],
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 1, len(headers))
        for column in (4, 5, 9):
            ws.cell(row_number, column).number_format = TOKEN_FORMAT
        for column in (6, 10, 11):
            ws.cell(row_number, column).number_format = USD_FORMAT
        for column in (7, 8):
            ws.cell(row_number, column).number_format = PERCENT_FORMAT
        ws.cell(row_number, 12).number_format = DATE_NUMBER_FORMAT
        if not business["is_expected"]:
            for cell in ws[row_number]:
                cell.fill = WARNING_FILL
    format_standard_sheet(ws, header_row, [12, 12, 34, 20, 18, 20, 16, 16, 15, 20, 20, 14])


def write_daily_summary(ws, analysis: dict[str, Any], args: argparse.Namespace) -> None:
    style_title(
        ws,
        "每天汇总",
        f"统计范围：{args.start_date} 至 {args.end_date}；日期连续，零补贴日期保留",
        8,
    )
    headers = [
        "日期",
        "金币Tokens",
        "金币补贴（USD）",
        "钻石Tokens",
        "钻石补贴（USD）",
        "当日补贴（USD）",
        "周期占比",
        "累计补贴（USD）",
    ]
    header_row = 4
    for column, header in enumerate(headers, start=1):
        ws.cell(header_row, column, header)
    style_header(ws, header_row, 1, len(headers))
    for offset, daily in enumerate(analysis["daily_rows"], start=1):
        row_number = header_row + offset
        values = [
            daily["dt"],
            daily["coin_tokens"],
            daily["coin_usd"],
            daily["diamond_tokens"],
            daily["diamond_usd"],
            daily["total_usd"],
            daily["period_share"],
            daily["cumulative_usd"],
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 1, len(headers))
        ws.cell(row_number, 1).number_format = DATE_NUMBER_FORMAT
        for column in (2, 4):
            ws.cell(row_number, column).number_format = TOKEN_FORMAT
        for column in (3, 5, 6, 8):
            ws.cell(row_number, column).number_format = USD_FORMAT
        ws.cell(row_number, 7).number_format = PERCENT_FORMAT

    format_standard_sheet(ws, header_row, [14, 20, 20, 20, 20, 20, 16, 20])
    if analysis["daily_rows"]:
        chart = LineChart()
        chart.title = "每日补贴金额"
        chart.style = 13
        chart.y_axis.title = "USD"
        chart.x_axis = DateAxis(crossAx=100)
        chart.x_axis.title = "日期"
        chart.x_axis.number_format = "yyyy-mm-dd"
        chart.x_axis.majorTimeUnit = "days"
        chart.y_axis.numFmt = '$#,##0.000'
        chart.height = 9
        chart.width = 17
        categories = Reference(
            ws,
            min_col=1,
            min_row=header_row + 1,
            max_row=header_row + len(analysis["daily_rows"]),
        )
        for data_column, series_title in (
            (3, "金币补贴（USD）"),
            (5, "钻石补贴（USD）"),
            (6, "合计补贴（USD）"),
        ):
            data = Reference(
                ws,
                min_col=data_column,
                max_col=data_column,
                min_row=header_row + 1,
                max_row=header_row + len(analysis["daily_rows"]),
            )
            chart.add_data(data, titles_from_data=False)
            chart.series[-1].tx = SeriesLabel(v=series_title)
        chart.set_categories(categories)
        chart.legend.position = "t"
        ws.add_chart(chart, "J4")


def write_daily_business(ws, analysis: dict[str, Any], args: argparse.Namespace) -> None:
    style_title(
        ws,
        "每天按业务汇总",
        f"统计范围：{args.start_date} 至 {args.end_date}；固定业务类型的零值记录完整保留",
        10,
    )
    headers = [
        "日期",
        "钱包类型",
        "业务类型",
        "业务名称",
        "Tokens",
        "Tokens/美元",
        "补贴金额（USD）",
        "当日钱包内占比",
        "当日总补贴占比",
        "周期总补贴占比",
    ]
    header_row = 4
    for column, header in enumerate(headers, start=1):
        ws.cell(header_row, column, header)
    style_header(ws, header_row, 1, len(headers))
    rows = sorted(
        analysis["all_rows"],
        key=lambda row: (row["dt"], 0 if row["wallet_type"] == "金币" else 1, row["business_type"]),
    )
    for offset, item in enumerate(rows, start=1):
        row_number = header_row + offset
        values = [
            item["dt"],
            item["wallet_type"],
            item["business_type"],
            item["name"],
            item["tokens"],
            item["tokens_per_usd"],
            item["usd"],
            item["daily_wallet_share"],
            item["daily_total_share"],
            item["period_share"],
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        style_body_row(ws, row_number, 1, len(headers))
        ws.cell(row_number, 1).number_format = DATE_NUMBER_FORMAT
        for column in (5, 6):
            ws.cell(row_number, column).number_format = TOKEN_FORMAT
        ws.cell(row_number, 7).number_format = USD_FORMAT
        for column in (8, 9, 10):
            ws.cell(row_number, column).number_format = PERCENT_FORMAT
        if not item["is_expected"]:
            for cell in ws[row_number]:
                cell.fill = WARNING_FILL
    format_standard_sheet(ws, header_row, [14, 12, 12, 38, 22, 20, 20, 20, 20, 20])


def create_workbook(
    analysis: dict[str, Any],
    args: argparse.Namespace,
    coin_day_count: int,
    coin_stat_count: int,
    diamond_day_count: int,
    diamond_stat_count: int,
) -> Workbook:
    workbook = Workbook()
    total_summary = workbook.active
    total_summary.title = "总汇总"
    write_total_summary(
        total_summary,
        analysis,
        args,
        coin_day_count,
        coin_stat_count,
        diamond_day_count,
        diamond_stat_count,
    )
    write_business_summary(workbook.create_sheet("按业务总汇总"), analysis, args)
    write_daily_summary(workbook.create_sheet("每天汇总"), analysis, args)
    write_daily_business(workbook.create_sheet("每天按业务汇总"), analysis, args)
    workbook.properties.title = "平台补贴汇总"
    workbook.properties.subject = "由 platform-subsidy-analysis 本地工具生成"
    workbook.active = 0
    return workbook


def save_workbook(workbook: Workbook, output: Path) -> None:
    if output.suffix.lower() != ".xlsx":
        raise ValueError("--output 必须使用 .xlsx 扩展名")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.name}.tmp.xlsx")
    try:
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        if temporary.exists():
            temporary.unlink()


def main() -> int:
    args = parse_args()
    try:
        start, end = validate_date_range(args.start_date, args.end_date)
        output_path = args.output.resolve()
        if output_path in {args.coin_input.resolve(), args.diamond_input.resolve()}:
            raise ValueError("--output 不能与任一输入 JSON 指向同一文件")
        coin_totals, coin_day_count, coin_stat_count = load_daily_totals(
            args.coin_input, "金币", start, end
        )
        diamond_totals, diamond_day_count, diamond_stat_count = load_daily_totals(
            args.diamond_input, "钻石", start, end
        )
        all_dates = dates_between(start, end)
        coin_rows = build_wallet_rows(
            "金币",
            coin_totals,
            COIN_BUSINESS_TYPES,
            COIN_TOKENS_PER_USD,
            all_dates,
        )
        diamond_rows = build_wallet_rows(
            "钻石",
            diamond_totals,
            DIAMOND_BUSINESS_TYPES,
            DIAMOND_TOKENS_PER_USD,
            all_dates,
        )
        analysis = build_analysis(coin_rows, diamond_rows, all_dates)
        workbook = create_workbook(
            analysis,
            args,
            coin_day_count,
            coin_stat_count,
            diamond_day_count,
            diamond_stat_count,
        )
        save_workbook(workbook, args.output)
    except ValueError as error:
        print(f"错误：{error}", file=sys.stderr)
        return 2
    except OSError as error:
        print(f"文件操作失败：{error}", file=sys.stderr)
        return 3

    coin_usd = analysis["wallet_totals"]["金币"]["usd"]
    diamond_usd = analysis["wallet_totals"]["钻石"]["usd"]
    print(f"报告已生成：{args.output.resolve()}")
    print(
        f"金币补贴=${coin_usd:,.3f}，钻石补贴=${diamond_usd:,.3f}，"
        f"合计=${analysis['period_total']:,.3f}"
    )
    print(
        f"统计天数={analysis['day_count']}，有补贴天数={analysis['days_with_subsidy']}，"
        f"最高补贴日期={analysis['highest_day']['dt']:%Y-%m-%d}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
