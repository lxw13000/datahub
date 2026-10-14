"""用户消费聚合与报表的离线回归测试；不访问真实 ES。"""

import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import export_report as report
import fetch_es


def bucket(key, tokens, count):
    return {"key": key, "doc_count": count, "total_tokens": {"value": float(tokens)}}


def response(buckets, after=None):
    group = {"buckets": buckets}
    if after is not None:
        group["after_key"] = after
    return {"timed_out": False, "_shards": {"total": 1, "successful": 1, "failed": 0},
            "aggregations": {"groups": group}}


class FakeES:
    def __init__(self):
        self.queries = []

    def search(self, query):
        self.queries.append(query)
        filters = query["query"]["bool"]["filter"]
        cohort = any("user_id" in item.get("terms", {}) for item in filters)
        sources = query["aggs"]["groups"]["composite"]["sources"]
        hourly = any("hour" in item for item in sources)
        if not cohort and not hourly:
            return response([
                bucket({"day": "2026-09-01", "businessType": 1}, 1_000_000, 4),
                bucket({"day": "2026-09-01", "businessType": 11}, 500_000, 2),
                bucket({"day": "2026-09-01", "businessType": 105}, 100_000, 1),
                bucket({"day": "2026-09-02", "businessType": 21}, 300_000, 3),
            ])
        if not cohort and hourly:
            return response([
                bucket({"hour": "2026-09-01T10:00:00"}, 1_600_000, 7),
                bucket({"hour": "2026-09-02T20:00:00"}, 300_000, 3),
            ])
        if cohort and not hourly:
            return response([
                bucket({"userId": 101, "day": "2026-09-01", "businessType": 1}, 200_000, 2),
                bucket({"userId": 101, "day": "2026-09-01", "businessType": 11}, 100_000, 1),
                bucket({"userId": 101, "day": "2026-09-01", "businessType": 105}, 50_000, 1),
                bucket({"userId": 101, "day": "2026-09-02", "businessType": 21}, 100_000, 1),
            ])
        return response([
            bucket({"hour": "2026-09-01T10:00:00"}, 350_000, 4),
            bucket({"hour": "2026-09-02T20:00:00"}, 100_000, 1),
        ])


class UserConsumptionTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        (self.root / "user-ids.csv").write_text("userId\n101\n102\n101\n\n", encoding="utf-8")
        self.es_path = self.root / "es-config.json"
        self.analysis_path = self.root / "analysis-config.json"
        self.es_path.write_text(json.dumps({"es_url": "https://example-es:9200",
                                            "username": "tester", "password": "secret"}), encoding="utf-8")
        self.analysis_path.write_text(json.dumps({"user_ids_file": "user-ids.csv",
                                                  "start": "2026-09-01 00:00:00",
                                                  "end": "2026-09-02 23:59:59",
                                                  "bucket_time_zone": "+00:00"}), encoding="utf-8")
        self.paths = fetch_es.output_paths(self.root, "test")

    def test_csv_and_config(self):
        self.assertEqual(report.load_user_ids(self.root / "user-ids.csv"), {101, 102})
        config = fetch_es.load_config(self.es_path, self.analysis_path)
        self.assertEqual(config.bucket_time_zone, "+00:00")
        (self.root / "user-ids.csv").write_text("userId\n101,extra\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "只有一列"):
            report.load_user_ids(self.root / "user-ids.csv")

    def test_query_whitelist_and_integer_sum(self):
        config = fetch_es.load_config(self.es_path, self.analysis_path)
        query = fetch_es.query_body(config, "daily", [101, 102])
        filters = query["query"]["bool"]["filter"]
        self.assertEqual(filters[0], {"terms": {"business_type": [1, 11, 21, 41, 42, 105, 106]}})
        self.assertEqual(filters[1], {"term": {"status": -1}})
        self.assertEqual(filters[2]["range"]["create_time"],
                         {"gte": "2026-09-01 00:00:00", "lte": "2026-09-02 23:59:59",
                          "time_zone": "+00:00"})
        self.assertEqual(filters[3], {"terms": {"user_id": [101, 102]}})
        self.assertEqual(fetch_es.exact_integer(125.0, "tokens"), 125)
        self.assertEqual(report.money(125), 0.013)
        with self.assertRaisesRegex(ValueError, "安全整数"):
            fetch_es.exact_integer(float(2**53), "tokens")

    def test_end_to_end_report(self):
        fake = FakeES()
        args = ["export_report.py", "--es-config", str(self.es_path),
                "--analysis-config", str(self.analysis_path)]
        with patch.object(fetch_es, "ESClient", return_value=fake), \
                patch.object(fetch_es, "output_paths", return_value=self.paths), \
                patch.object(sys, "argv", args):
            self.assertEqual(report.main(), 0)
        self.assertEqual(len(fake.queries), 4)
        raw = json.loads(self.paths.data.read_text(encoding="utf-8"))
        self.assertEqual(raw["cohortDaily"][0]["tokens"], 200_000)
        self.assertEqual(raw["meta"]["user_count"], 2)
        manifest = self.paths.manifest.read_text(encoding="utf-8")
        self.assertNotIn("secret", manifest)
        self.assertEqual(json.loads(manifest)["status"], "report_complete")

        book = load_workbook(self.paths.report)
        self.assertEqual(book.sheetnames, ["总体对比", "业务类型", "每日趋势", "时段习惯", "用户明细"])
        overview = book["总体对比"]
        self.assertEqual([overview[cell].value for cell in ("A9", "B9", "C9", "D9", "F9")],
                         [2, 1, 1, 450_000, 1_900_000])
        self.assertAlmostEqual(overview["H9"].value, 450_000 / 1_900_000)
        self.assertEqual(overview["E9"].value, 45)
        self.assertEqual(overview["G9"].value, 190)
        self.assertAlmostEqual(overview["D17"].value, 50_000 / 450_000)  # VIP 行
        self.assertEqual(book["业务类型"]["B7"].value, 105)
        self.assertEqual(book["每日趋势"]["A2"].value.date().isoformat(), "2026-09-01")
        self.assertEqual(book["每日趋势"]["E2"].value, 1)
        self.assertEqual(book["时段习惯"]["B12"].value, 350_000)  # 10:00
        self.assertEqual(book["用户明细"]["A3"].value, 102)
        self.assertEqual(book["用户明细"]["B3"].value, 0)
        self.assertEqual(book["用户明细"]["A2"].number_format, "0")
        self.assertEqual(book["用户明细"]["C2"].number_format, report.USD_FORMAT)
        self.assertTrue(all(sheet.sheet_view.showGridLines for sheet in book))
        self.assertEqual(book["用户明细"]["B2"].border.left.style, "thin")
        self.assertFalse(any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))
        book.close()

    def test_resume_skips_completed_jobs_and_rejects_changed_scope(self):
        config = fetch_es.load_config(self.es_path, self.analysis_path)
        ids = {101, 102}

        class Interrupted(FakeES):
            def search(self, query):
                if len(self.queries) == 2:
                    raise ValueError("模拟中断")
                return super().search(query)

        first = Interrupted()
        with self.assertRaisesRegex(ValueError, "模拟中断"):
            fetch_es.export_json(config, ids, self.paths, first)
        manifest = json.loads(self.paths.manifest.read_text(encoding="utf-8"))
        self.assertEqual(set(manifest["completed_jobs"]), {"platform_daily", "platform_hourly"})
        self.assertFalse(self.paths.data.exists())
        self.assertEqual(fetch_es.find_resume_paths(config, ids, self.root), self.paths)
        second = FakeES()
        fetch_es.export_json(config, ids, self.paths, second, resume=True)
        self.assertEqual(len(second.queries), 2)
        with self.assertRaisesRegex(ValueError, "口径或用户名单"):
            fetch_es.export_json(config, ids | {103}, self.paths, second, resume=True)
        self.assertEqual(len(second.queries), 2)

    def test_resume_cli_finishes_report(self):
        class Interrupted(FakeES):
            def search(self, query):
                if len(self.queries) == 2:
                    raise ValueError("模拟中断")
                return super().search(query)

        first, second = Interrupted(), FakeES()
        argv = ["export_report.py", "--es-config", str(self.es_path),
                "--analysis-config", str(self.analysis_path)]
        with patch.object(fetch_es, "ESClient", side_effect=[first, second]), \
                patch.object(fetch_es, "output_paths", return_value=self.paths), \
                patch.object(fetch_es, "find_resume_paths", return_value=self.paths), \
                patch.object(sys, "argv", argv):
            self.assertEqual(report.main(), 2)
            self.assertFalse(self.paths.report.exists())
            with patch.object(sys, "argv", argv + ["--resume"]):
                self.assertEqual(report.main(), 0)
        self.assertEqual(len(second.queries), 2)
        self.assertTrue(self.paths.report.exists())
        self.assertEqual(json.loads(self.paths.manifest.read_text(encoding="utf-8"))["status"], "report_complete")

    def test_reconciliation_rejects_incomplete_hourly_data(self):
        config = fetch_es.load_config(self.es_path, self.analysis_path)
        fetch_es.export_json(config, {101, 102}, self.paths, FakeES())
        data = report.load_data(self.paths.data)
        data["cohortHourly"][0]["tokens"] -= 1
        with self.assertRaisesRegex(ValueError, "不一致"):
            report.build_analysis(data, {101, 102})

    def test_empty_consumption_keeps_all_users(self):
        data = {"platformDaily": [], "platformHourly": [],
                "cohortDaily": [], "cohortHourly": []}
        result = report.build_analysis(data, {101, 102})
        self.assertEqual(result["user_count"], 2)
        self.assertEqual(result["active_user_count"], 0)
        self.assertEqual(result["zero_user_count"], 2)
        config = fetch_es.load_config(self.es_path, self.analysis_path)
        book = report.create_workbook(result, config)
        self.assertEqual(book["用户明细"].max_row, 3)
        self.assertIsNone(book["总体对比"]["H9"].value)  # 平台零消费时占比不可计算

    def test_composite_pagination_and_user_batching(self):
        config = fetch_es.load_config(self.es_path, self.analysis_path)
        ids = set(range(1, 1002))

        class EmptyPagedES:
            def __init__(self):
                self.queries = []

            def search(self, query):
                self.queries.append(query)
                composite = query["aggs"]["groups"]["composite"]
                if len(self.queries) == 1:
                    return response([bucket({"day": "2026-09-01", "businessType": 1}, 0, 1)],
                                    {"day": "2026-09-01", "businessType": 1})
                return response([])

        fake = EmptyPagedES()
        fetch_es.export_json(config, ids, self.paths, fake)
        self.assertEqual(len(fake.queries), 7)  # 平台首查询翻页，加上其余五个任务
        self.assertEqual(fake.queries[1]["aggs"]["groups"]["composite"]["after"],
                         {"day": "2026-09-01", "businessType": 1})
        cohort_queries = [query for query in fake.queries if any(
            "user_id" in item.get("terms", {}) for item in query["query"]["bool"]["filter"])]
        self.assertEqual([len(query["query"]["bool"]["filter"][-1]["terms"]["user_id"])
                          for query in cohort_queries], [1000, 1000, 1, 1])

    def test_es_client_uses_coin_index_and_auth(self):
        config = fetch_es.load_config(self.es_path, self.analysis_path)
        captured = []

        def fake_urlopen(request, **options):
            captured.append((request, options))
            return io.BytesIO(json.dumps(response([])).encode("utf-8"))

        with patch.object(fetch_es, "urlopen", side_effect=fake_urlopen):
            fetch_es.ESClient(config).search({"size": 0})
        request, options = captured[0]
        self.assertEqual(request.full_url, "https://example-es:9200/sano_wallet_coin_record/_search")
        self.assertTrue(request.get_header("Authorization").startswith("Basic "))
        self.assertIsNotNone(options["context"])


if __name__ == "__main__":
    unittest.main()
