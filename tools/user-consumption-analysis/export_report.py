#!/usr/bin/env python3
"""查询金币消费聚合，并生成用户群体与全平台对比分析工作簿。"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


CATEGORIES = ("普通礼物", "幸运礼物", "游戏", "道具", "VIP")
BUSINESS_NAMES = {1: "赠送豪华礼物（普通礼物）", 11: "赠送幸运礼物", 21: "游戏消费",
                  41: "购买装扮道具", 42: "购买VIP", 105: "购买VIP套餐补签卡", 106: "购买VIP套餐"}
CATEGORY_BY_TYPE = {1: "普通礼物", 11: "幸运礼物", 21: "游戏", 41: "道具",
                    42: "VIP", 105: "VIP", 106: "VIP"}
TOKENS_PER_USD = Decimal(10_000)
INTEGER_PATTERN = re.compile(r"[+-]?\d+")
TOKEN_FORMAT = "#,##0"
USD_FORMAT = '"$"#,##0.000'
PERCENT_FORMAT = "0.00%"
USER_ID_FORMAT = "0"
TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="5B9BD5")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
LIGHT_FILL = PatternFill("solid", fgColor="EAF2F8")
HEADER_FONT = Font(name="Arial", color="FFFFFF", bold=True)
BODY_FONT = Font(name="Arial", color="1F1F1F")
LABEL_FONT = Font(name="Arial", color="1F1F1F", bold=True)
SIDE = Side(style="thin", color="B7C5D3")
CELL_BORDER = Border(left=SIDE, right=SIDE, top=SIDE, bottom=SIDE)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="查询 ES 金币消费聚合，导出 JSON 与用户消费分析 Excel")
    parser.add_argument("--es-config", type=Path, required=True, help="ES 连接配置 JSON")
    parser.add_argument("--analysis-config", type=Path, required=True, help="用户 ID 与时间范围配置 JSON")
    parser.add_argument("--resume", action="store_true", help="续跑相同口径的未完成查询")
    return parser.parse_args()


def parse_int(value: Any, field: str, location: str, minimum: int = 0) -> int:
    if type(value) is int:
        result = value
    elif isinstance(value, str) and INTEGER_PATTERN.fullmatch(value.strip()):
        result = int(value.strip())
    else:
        raise ValueError(f"{location} 的 {field} 必须是整数")
    if result < minimum:
        raise ValueError(f"{location} 的 {field} 不能小于 {minimum}")
    return result


def load_user_ids(path: Path) -> set[int]:
    if not path.is_file():
        raise ValueError(f"用户 ID CSV 不存在：{path}")
    ids: set[int] = set()
    first = True
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        for line_no, row in enumerate(csv.reader(source), start=1):
            if not row or all(not cell.strip() for cell in row):
                continue
            if any(cell.strip() for cell in row[1:]):
                raise ValueError(f"用户 ID CSV 第 {line_no} 行必须只有一列")
            value = row[0].strip()
            if first:
                first = False
                if value.lower().replace("_", "").replace("-", "") in {"userid", "uid", "用户id"}:
                    continue
            ids.add(parse_int(value, "userId", f"CSV 第 {line_no} 行", minimum=1))
    if not ids:
        raise ValueError("用户 ID CSV 没有有效用户")
    return ids


def reject_float(value: str) -> None:
    raise ValueError(f"聚合 JSON 出现小数 {value}，tokens 必须保持整数")


def load_data(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise ValueError(f"聚合 JSON 不存在：{path}")
    with path.open("r", encoding="utf-8-sig") as source:
        root = json.load(source, parse_int=int, parse_float=reject_float)
    if not isinstance(root, dict) or not isinstance(root.get("meta"), dict):
        raise ValueError("聚合 JSON 根节点格式错误")
    for key in ("platformDaily", "platformHourly", "cohortDaily", "cohortHourly"):
        if not isinstance(root.get(key), list):
            raise ValueError(f"聚合 JSON 缺少 {key} 数组")
    return root


def money(tokens: int) -> float:
    """先合并整数金币，再在写 Excel 时以十进制保留三位美元。"""
    return float((Decimal(tokens) / TOKENS_PER_USD).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP))


def empty_measure() -> dict[str, int]:
    return {"tokens": 0, "records": 0}


def add_measure(target: dict[str, int], row: dict[str, Any]) -> None:
    target["tokens"] += row["tokens"]
    target["records"] += row["recordCount"]


def ratio(numerator: int, denominator: int) -> float | None:
    return numerator / denominator if denominator else None


def build_analysis(root: dict[str, Any], user_ids: set[int]) -> dict[str, Any]:
    platform_total, cohort_total = empty_measure(), empty_measure()
    platform_hour_total, cohort_hour_total = empty_measure(), empty_measure()
    platform_by_type = defaultdict(empty_measure)
    cohort_by_type = defaultdict(empty_measure)
    platform_by_day = defaultdict(empty_measure)
    cohort_by_day = defaultdict(empty_measure)
    platform_by_day_cat = defaultdict(empty_measure)
    cohort_by_day_cat = defaultdict(empty_measure)
    by_user = {user_id: {"tokens": 0, "records": 0,
                         "days": set(), "by_category": defaultdict(int)} for user_id in user_ids}
    platform_hour_of_day = defaultdict(empty_measure)
    cohort_hour_of_day = defaultdict(empty_measure)
    platform_weekday = defaultdict(empty_measure)
    cohort_weekday = defaultdict(empty_measure)
    active_by_day = defaultdict(set)

    for array_name in ("platformDaily", "cohortDaily", "platformHourly", "cohortHourly"):
        for index, row in enumerate(root[array_name], start=1):
            location = f"{array_name} 第 {index} 桶"
            if not isinstance(row, dict):
                raise ValueError(f"{location} 不是对象")
            clean = {"tokens": parse_int(row.get("tokens"), "tokens", location),
                     "recordCount": parse_int(row.get("recordCount"), "recordCount", location, 1)}
            if array_name.endswith("Daily"):
                day = row.get("day")
                business = parse_int(row.get("businessType"), "businessType", location)
                if not isinstance(day, str) or business not in CATEGORY_BY_TYPE:
                    raise ValueError(f"{location} 日期或业务类型无效")
                date.fromisoformat(day)
                category = CATEGORY_BY_TYPE[business]
                if array_name == "platformDaily":
                    add_measure(platform_total, clean)
                    targets = ((platform_by_type, business), (platform_by_day, day),
                               (platform_by_day_cat, (day, category)))
                else:
                    user_id = parse_int(row.get("userId"), "userId", location, 1)
                    if user_id not in user_ids:
                        raise ValueError(f"{location} 用户不在输入名单")
                    add_measure(cohort_total, clean)
                    targets = ((cohort_by_type, business), (cohort_by_day, day),
                               (cohort_by_day_cat, (day, category)))
                    user = by_user[user_id]
                    user["tokens"] += clean["tokens"]
                    user["records"] += clean["recordCount"]
                    user["days"].add(day)
                    user["by_category"][category] += clean["tokens"]
                    active_by_day[day].add(user_id)
                for destination, key in targets:
                    add_measure(destination[key], clean)
            else:
                timestamp = row.get("hour")
                if not isinstance(timestamp, str):
                    raise ValueError(f"{location} 小时无效")
                hour = datetime.strptime(timestamp, "%Y-%m-%dT%H:%M:%S").hour
                if array_name == "platformHourly":
                    add_measure(platform_hour_total, clean)
                    add_measure(platform_hour_of_day[hour], clean)
                else:
                    add_measure(cohort_hour_total, clean)
                    add_measure(cohort_hour_of_day[hour], clean)

    if platform_total != platform_hour_total or cohort_total != cohort_hour_total:
        raise ValueError("每日与每小时聚合总数不一致，拒绝生成不完整报表")
    if (cohort_total["tokens"] > platform_total["tokens"]
            or cohort_total["records"] > platform_total["records"]):
        raise ValueError("名单消费超过全平台总量，聚合口径不一致")
    for daily, weekday_out in ((platform_by_day, platform_weekday),
                               (cohort_by_day, cohort_weekday)):
        for day, measure in daily.items():
            weekday = date.fromisoformat(day).weekday()
            add_measure(weekday_out[weekday], {"tokens": measure["tokens"],
                                              "recordCount": measure["records"]})

    active_ids = {user_id for user_id, row in by_user.items() if row["records"]}
    for business in BUSINESS_NAMES:
        selected, platform = cohort_by_type[business], platform_by_type[business]
        if selected["tokens"] > platform["tokens"] or selected["records"] > platform["records"]:
            raise ValueError(f"业务类型 {business} 的名单消费超过全平台，聚合口径不一致")
    return {
        "platform_total": platform_total,
        "cohort_total": cohort_total,
        "platform_by_type": platform_by_type,
        "cohort_by_type": cohort_by_type,
        "platform_by_day": platform_by_day,
        "cohort_by_day": cohort_by_day,
        "platform_by_day_cat": platform_by_day_cat,
        "cohort_by_day_cat": cohort_by_day_cat,
        "platform_hour_of_day": platform_hour_of_day,
        "cohort_hour_of_day": cohort_hour_of_day,
        "platform_weekday": platform_weekday,
        "cohort_weekday": cohort_weekday,
        "active_by_day": active_by_day,
        "by_user": by_user,
        "user_count": len(user_ids),
        "active_user_count": len(active_ids),
        "zero_user_count": len(user_ids) - len(active_ids),
    }


def header(ws, row: int, names: list[str]) -> None:
    for column, label in enumerate(names, start=1):
        cell = ws.cell(row, column, label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def body(ws, start: int, end: int, columns: int, *, formats: dict[int, str] | None = None) -> None:
    formats = formats or {}
    for row in ws.iter_rows(min_row=start, max_row=end, min_col=1, max_col=columns):
        for cell in row:
            cell.border = CELL_BORDER
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center", horizontal="right" if isinstance(cell.value, (int, float)) else "left")
            cell.number_format = formats.get(cell.column, TOKEN_FORMAT if isinstance(cell.value, (int, float)) else "General")


def widths(ws, values: list[float]) -> None:
    for column, width in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width


def tabular(ws, names: list[str], rows: list[list[Any]], column_widths: list[float],
            formats: dict[int, str], *, freeze: str = "A2") -> None:
    header(ws, 1, names)
    for row in rows:
        ws.append(row)
    if rows:
        body(ws, 2, ws.max_row, len(names), formats=formats)
    widths(ws, column_widths)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(names))}{max(1, ws.max_row)}"
    ws.sheet_view.showGridLines = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def section(ws, row: int, label: str, columns: int) -> None:
    ws.cell(row, 1, label)
    for column in range(1, columns + 1):
        cell = ws.cell(row, column)
        cell.fill = SECTION_FILL
        cell.border = CELL_BORDER
        cell.font = Font(name="Arial", color="1F4E78", bold=True, size=11)
    ws.row_dimensions[row].height = 24


def write_overview(ws, result: dict[str, Any], config: Any, user_ids_file: Path) -> None:
    ws.title = "总体对比"
    ws.sheet_view.showGridLines = True
    ws["A1"] = "用户消费习惯分析"
    for row in ws["A1:J1"]:
        for cell in row:
            cell.fill = TITLE_FILL
            cell.border = CELL_BORDER
    ws["A1"].font = Font(name="Arial", size=17, bold=True, color="FFFFFF")
    ws.row_dimensions[1].height = 35
    conditions = [("消费开始", config.start, "消费结束", config.end),
                  ("用户 ID 文件", user_ids_file.name, "聚合时间区", config.bucket_time_zone),
                  ("消费口径", "1/11/21/41/42/105/106，status=-1", "金币兑美元", "10,000 : 1")]
    for row_no, values in enumerate(conditions, start=2):
        for col, value in zip((1, 2, 6, 7), values):
            ws.cell(row_no, col, value)
            ws.cell(row_no, col).font = LABEL_FONT if col in (1, 6) else BODY_FONT
        for col in range(1, 11):
            ws.cell(row_no, col).border = CELL_BORDER
        ws.row_dimensions[row_no].height = 23
    ws.merge_cells("B4:E4")
    ws["A5"] = "说明：消费记录数为 ES 文档数；金额先按整数金币汇总，再换算美元保留三位小数。"
    ws["A5"].font = Font(name="Arial", italic=True, color="596579")
    section(ws, 7, "总体消费", 10)
    headline = ["名单用户数", "有消费用户数", "零消费用户数", "名单金币", "名单美元",
                "全平台金币", "全平台美元", "名单金额占比", "名单记录数", "全平台记录数"]
    header(ws, 8, headline)
    cohort, platform = result["cohort_total"], result["platform_total"]
    values = [result["user_count"], result["active_user_count"], result["zero_user_count"],
              cohort["tokens"], money(cohort["tokens"]), platform["tokens"],
              money(platform["tokens"]), ratio(cohort["tokens"], platform["tokens"]),
              cohort["records"], platform["records"]]
    for col, value in enumerate(values, start=1):
        ws.cell(9, col, value)
    body(ws, 9, 9, 10, formats={5: USD_FORMAT, 7: USD_FORMAT, 8: PERCENT_FORMAT})
    section(ws, 11, "五类消费结构及全平台占比", 10)
    header(ws, 12, ["消费类别", "名单金币", "名单美元", "名单内占比", "全平台金币",
                    "全平台美元", "全平台内占比", "名单占平台金额比", "名单记录数", "全平台记录数"])
    for row_no, category in enumerate(CATEGORIES, start=13):
        selected = sum(result["cohort_by_type"][business]["tokens"] for business in BUSINESS_NAMES
                       if CATEGORY_BY_TYPE[business] == category)
        all_tokens = sum(result["platform_by_type"][business]["tokens"] for business in BUSINESS_NAMES
                         if CATEGORY_BY_TYPE[business] == category)
        selected_records = sum(result["cohort_by_type"][business]["records"] for business in BUSINESS_NAMES
                               if CATEGORY_BY_TYPE[business] == category)
        all_records = sum(result["platform_by_type"][business]["records"] for business in BUSINESS_NAMES
                          if CATEGORY_BY_TYPE[business] == category)
        values = [category, selected, money(selected), ratio(selected, cohort["tokens"]),
                  all_tokens, money(all_tokens), ratio(all_tokens, platform["tokens"]),
                  ratio(selected, all_tokens), selected_records, all_records]
        for col, value in enumerate(values, start=1):
            ws.cell(row_no, col, value)
    body(ws, 13, 17, 10, formats={3: USD_FORMAT, 4: PERCENT_FORMAT, 6: USD_FORMAT,
                                   7: PERCENT_FORMAT, 8: PERCENT_FORMAT})
    widths(ws, [22, 20, 18, 17, 20, 18, 18, 21, 17, 18])
    ws.freeze_panes = "A8"


def write_business(ws, result: dict[str, Any]) -> None:
    rows = []
    for business, name in BUSINESS_NAMES.items():
        selected = result["cohort_by_type"][business]
        platform = result["platform_by_type"][business]
        rows.append([CATEGORY_BY_TYPE[business], business, name,
                     selected["tokens"], money(selected["tokens"]), selected["records"],
                     platform["tokens"], money(platform["tokens"]), platform["records"],
                     ratio(selected["tokens"], platform["tokens"]),
                     ratio(selected["records"], platform["records"])])
    tabular(ws, ["类别", "业务类型", "业务名称", "名单金币", "名单美元", "名单记录数",
                 "全平台金币", "全平台美元", "全平台记录数", "金额占平台比", "记录占平台比"],
            rows, [14, 12, 29, 20, 18, 16, 20, 18, 18, 18, 18],
            {5: USD_FORMAT, 8: USD_FORMAT, 10: PERCENT_FORMAT, 11: PERCENT_FORMAT})


def write_daily(ws, result: dict[str, Any]) -> None:
    days = sorted(set(result["platform_by_day"]) | set(result["cohort_by_day"]))
    rows = []
    for day in days:
        selected, platform = result["cohort_by_day"][day], result["platform_by_day"][day]
        rows.append([date.fromisoformat(day), selected["tokens"], money(selected["tokens"]),
                     selected["records"], len(result["active_by_day"][day]),
                     platform["tokens"], money(platform["tokens"]), platform["records"],
                     ratio(selected["tokens"], platform["tokens"])])
    tabular(ws, ["日期", "名单金币", "名单美元", "名单记录数", "名单活跃用户",
                 "全平台金币", "全平台美元", "全平台记录数", "名单金额占比"],
            rows, [15, 19, 18, 17, 18, 20, 18, 19, 19],
            {1: "yyyy-mm-dd", 3: USD_FORMAT, 7: USD_FORMAT, 9: PERCENT_FORMAT})
    # 每日类别明细放在同一工作表下方，保持每天总量与构成可直接对照。
    start = ws.max_row + 3
    section(ws, start, "每日五类消费", 9)
    headers = ["日期", "类别", "名单金币", "名单美元", "名单记录数",
               "全平台金币", "全平台美元", "全平台记录数", "名单金额占比"]
    header(ws, start + 1, headers)
    current = start + 2
    for day in days:
        for category in CATEGORIES:
            selected = result["cohort_by_day_cat"][(day, category)]
            platform = result["platform_by_day_cat"][(day, category)]
            values = [date.fromisoformat(day), category, selected["tokens"], money(selected["tokens"]),
                      selected["records"], platform["tokens"], money(platform["tokens"]),
                      platform["records"], ratio(selected["tokens"], platform["tokens"])]
            for col, value in enumerate(values, start=1):
                ws.cell(current, col, value)
            current += 1
    if current > start + 2:
        body(ws, start + 2, current - 1, 9,
             formats={1: "yyyy-mm-dd", 4: USD_FORMAT, 7: USD_FORMAT, 9: PERCENT_FORMAT})


def write_time(ws, result: dict[str, Any]) -> None:
    hours = []
    for hour in range(24):
        selected, platform = result["cohort_hour_of_day"][hour], result["platform_hour_of_day"][hour]
        hours.append([f"{hour:02d}:00–{hour:02d}:59", selected["tokens"], money(selected["tokens"]),
                      selected["records"], ratio(selected["tokens"], result["cohort_total"]["tokens"]),
                      platform["tokens"], money(platform["tokens"]), platform["records"],
                      ratio(selected["tokens"], platform["tokens"])])
    tabular(ws, ["小时", "名单金币", "名单美元", "名单记录数", "名单内金额占比",
                 "全平台金币", "全平台美元", "全平台记录数", "名单金额占比"],
            hours, [17, 19, 18, 18, 20, 20, 18, 19, 19],
            {3: USD_FORMAT, 5: PERCENT_FORMAT, 7: USD_FORMAT, 9: PERCENT_FORMAT})
    start = ws.max_row + 3
    section(ws, start, "星期消费分布", 9)
    header(ws, start + 1, ["星期", "名单金币", "名单美元", "名单记录数", "名单内金额占比",
                           "全平台金币", "全平台美元", "全平台记录数", "名单金额占比"])
    weekday_names = ("周一", "周二", "周三", "周四", "周五", "周六", "周日")
    for weekday, name in enumerate(weekday_names):
        selected, platform = result["cohort_weekday"][weekday], result["platform_weekday"][weekday]
        values = [name, selected["tokens"], money(selected["tokens"]), selected["records"],
                  ratio(selected["tokens"], result["cohort_total"]["tokens"]),
                  platform["tokens"], money(platform["tokens"]), platform["records"],
                  ratio(selected["tokens"], platform["tokens"])]
        for col, value in enumerate(values, start=1):
            ws.cell(start + 2 + weekday, col, value)
    body(ws, start + 2, start + 8, 9,
         formats={3: USD_FORMAT, 5: PERCENT_FORMAT, 7: USD_FORMAT, 9: PERCENT_FORMAT})


def write_users(ws, result: dict[str, Any]) -> None:
    rows = []
    for user_id, user in sorted(result["by_user"].items(), key=lambda item: (-item[1]["tokens"], item[0])):
        categories = user["by_category"]
        dominant = (min(CATEGORIES, key=lambda category: (-categories[category], CATEGORIES.index(category)))
                    if user["tokens"] else "无消费")
        active_days = len(user["days"])
        average = (float((Decimal(user["tokens"]) / active_days).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP)) if active_days else None)
        rows.append([user_id, user["tokens"], money(user["tokens"]), user["records"], active_days,
                     average,
                     dominant, *(categories[category] for category in CATEGORIES),
                     ratio(user["tokens"], result["cohort_total"]["tokens"])])
    tabular(ws, ["用户ID", "消费金币", "消费美元", "消费记录数", "活跃天数", "活跃日均金币",
                 "主要消费类型", "普通礼物金币", "幸运礼物金币", "游戏金币", "道具金币", "VIP金币",
                 "占名单金额比"],
            rows, [16, 20, 18, 17, 15, 19, 18, 21, 21, 18, 18, 18, 18],
            {1: USER_ID_FORMAT, 3: USD_FORMAT, 6: "#,##0.00", 13: PERCENT_FORMAT},
            freeze="B2")


def create_workbook(result: dict[str, Any], config: Any) -> Workbook:
    book = Workbook()
    write_overview(book.active, result, config, config.user_ids_file)
    write_business(book.create_sheet("业务类型"), result)
    write_daily(book.create_sheet("每日趋势"), result)
    write_time(book.create_sheet("时段习惯"), result)
    write_users(book.create_sheet("用户明细"), result)
    book.properties.title = "用户消费习惯分析"
    book.properties.subject = "金币消费名单与全平台同期对比"
    book.active = 0
    return book


def save_workbook(book: Workbook, path: Path) -> None:
    temporary = path.with_name(f".{path.name}.tmp.xlsx")
    try:
        book.save(temporary)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()


def main() -> int:
    command = parse_args()
    try:
        from fetch_es import (analysis_stream, export_json, find_resume_paths, load_config,
                              mark_report_complete, output_paths)

        config = load_config(command.es_config, command.analysis_config)
        user_ids = load_user_ids(config.user_ids_file)
        paths = (find_resume_paths(config, user_ids) if command.resume else output_paths())
        export_json(config, user_ids, paths, resume=command.resume)
        data = analysis_stream(paths, user_ids)
        result = build_analysis(data, user_ids)
        book = create_workbook(result, config)
        save_workbook(book, paths.report)
        mark_report_complete(paths)
    except (ValueError, OSError, json.JSONDecodeError, csv.Error) as error:
        print(f"错误：{error}", file=sys.stderr)
        return 2
    print(f"报告已生成：{paths.report.resolve()}")
    print(f"名单 {result['user_count']:,} 人，有消费 {result['active_user_count']:,} 人；"
          f"名单消费 {result['cohort_total']['tokens']:,} 金币，"
          f"同期平台消费 {result['platform_total']['tokens']:,} 金币")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
